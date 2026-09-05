r"""
Import courses for كلية الهندسة from Excel file.
Excel file location: C:\Users\rubas\Desktop\كلية الهندسة.xlsx

Sheet names map to programs under Faculty of Engineering (faculty_id = 1):
  - برنامج المستوى العام                   -> program_id = 26 (المستوى العام)
  - برنامج تخطيط وتشييد المدن الذك        -> program_id = 4  (تخطيط وتشييد المدن الذكية)
  -  برنامج هندسه الحاسوب                  -> program_id = 5  (هندسة الحاسوب)
  - برنامج هندسه المواد و اداره ال         -> program_id = 6  (هندسة المواد وإدارة التصنيع)
"""

import openpyxl
import sqlite3
import os
import sys

EXCEL_FILE = r'C:\Users\rubas\Desktop\كلية الهندسة.xlsx'
DB_FILE = os.path.join(os.path.dirname(__file__), 'mnu_system.db')

def get_program_id(sheet_name):
    s = sheet_name.strip()
    if 'العام' in s:
        return 26  # المستوى العام
    if 'تخطيط' in s or 'المدن' in s:
        return 4   # تخطيط وتشييد المدن الذكية
    if 'الحاسوب' in s or 'حاسوب' in s:
        return 5   # هندسة الحاسوب
    if 'المواد' in s or 'التصنيع' in s:
        return 6   # هندسة المواد وإدارة التصنيع
    return None

def parse_val(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s in ("None", "none", "-", "–"):
        return ""
    return s

def parse_float(val):
    if val is None:
        return 0.0
    try:
        return float(str(val).strip())
    except:
        return 0.0

def parse_level(val):
    if not val:
        return 0
    s = str(val).strip()
    if "0" in s:
        return 0
    if "1" in s or "أول" in s or "اول" in s:
        return 1
    if "2" in s or "ثاني" in s or "ثانى" in s:
        return 2
    if "3" in s or "ثالث" in s:
        return 3
    if "4" in s or "رابع" in s:
        return 4
    if "5" in s or "خامس" in s:
        return 5
    return 0

def main():
    if not os.path.exists(EXCEL_FILE):
        print(f"ERROR: Excel file not found at {EXCEL_FILE}")
        return

    if not os.path.exists(DB_FILE):
        print(f"ERROR: Database file not found at {DB_FILE}")
        return

    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # Faculty ID 1 for كلية الهندسة
    fac_row = cursor.execute("SELECT id, name FROM faculties WHERE id = 1 OR name LIKE '%هندسة%'").fetchone()
    if not fac_row:
        print("ERROR: Faculty of Engineering not found!")
        return
    faculty_id = fac_row['id']
    print(f"Target Faculty: {fac_row['name']} (ID: {faculty_id})")

    # Programs for Engineering
    progs = cursor.execute("SELECT id, name FROM programs WHERE faculty_id = ?", (faculty_id,)).fetchall()
    prog_dict = {p['id']: p['name'] for p in progs}
    print("Faculty Programs in DB:", prog_dict)

    wb = openpyxl.load_workbook(EXCEL_FILE)

    total_inserted = 0
    total_skipped = 0

    print("\n--- Starting Engineering Course Registration Process ---")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        prog_id = get_program_id(sheet_name)
        if not prog_id:
            print(f"WARNING: Could not map sheet '{sheet_name}' to any program. Skipping.")
            continue

        prog_name = prog_dict.get(prog_id, 'Unknown Program')
        print(f"\nProcessing Sheet: '{sheet_name}' -> Program: '{prog_name}' (ID: {prog_id})")

        # Dynamic Header Mapping from Row 3 and Row 4
        r3 = [ws.cell(3, col).value for col in range(1, ws.max_column + 1)]
        r4 = [ws.cell(4, col).value for col in range(1, ws.max_column + 1)]
        
        col_map = {}
        current_section = ""
        for idx in range(ws.max_column):
            val3 = str(r3[idx]).strip() if r3[idx] is not None else ""
            val4 = str(r4[idx]).strip() if r4[idx] is not None else ""
            if val3 and val3 not in ("None", ""):
                current_section = val3
            full_header = f"{current_section} - {val4}" if val4 else current_section
            col_map[idx] = full_header

        col_indices = {}
        for idx, header in col_map.items():
            h = header.strip()
            if h == 'كود المقرر':
                col_indices['code'] = idx
            elif h == 'كود المقرر انجليزى':
                col_indices['code_en'] = idx
            elif h == 'اسم المقرر':
                col_indices['name_ar'] = idx
            elif h == 'اسم المقرر انجليزى':
                col_indices['name_en'] = idx
            elif h == 'نوع المقرر':
                col_indices['course_type'] = idx
            elif h == 'المستوي':
                col_indices['level'] = idx
            elif h == 'الفصل الدراسي':
                col_indices['semester'] = idx
            elif 'البرامج الاخرى' in h:
                col_indices['other_programs'] = idx
            elif h == 'متطلب':
                col_indices['requirement'] = idx
            elif 'المتطلب السابق' in h:
                col_indices['prerequisite'] = idx
            elif 'المقررات المتزامنة' in h:
                col_indices['concurrent_courses'] = idx
            elif 'يضاف للمعدل' in h:
                col_indices['added_to_gpa'] = idx
            elif 'نجاح أو رسوب' in h or 'نجاح او رسوب' in h:
                col_indices['pass_fail'] = idx
            elif 'الساعات المعتمدة' in h:
                col_indices['credit_hours'] = idx

        # Fetch existing course codes for this program
        existing_in_prog = cursor.execute("SELECT code FROM courses WHERE program_id = ?", (prog_id,)).fetchall()
        existing_codes = set(r['code'].strip() for r in existing_in_prog if r['code'])

        inserted_sheet = 0
        skipped_sheet = 0

        for r_idx in range(5, ws.max_row + 1):
            seq = ws.cell(r_idx, 1).value
            if seq is None or not str(seq).strip().isdigit():
                continue

            row_vals = [ws.cell(r_idx, c_idx + 1).value for c_idx in range(ws.max_column)]
            
            code_idx = col_indices.get('code', 1)
            code = parse_val(row_vals[code_idx])
            if not code:
                continue

            code_en_idx = col_indices.get('code_en', 2)
            code_en = parse_val(row_vals[code_en_idx]) or code
            
            name_ar_idx = col_indices.get('name_ar', 3)
            raw_name_ar = parse_val(row_vals[name_ar_idx])
            
            name_en_idx = col_indices.get('name_en', 4)
            name_en = parse_val(row_vals[name_en_idx])
            
            name_ar = raw_name_ar
            if code and name_ar.endswith(code):
                name_ar = name_ar[:-len(code)].strip()
            if '\n' in name_ar:
                parts = [p.strip() for p in name_ar.split('\n') if p.strip()]
                if len(parts) > 1 and parts[-1] == code:
                    name_ar = " ".join(parts[:-1])
                else:
                    name_ar = " ".join(parts)

            course_type_idx = col_indices.get('course_type', 5)
            course_type = parse_val(row_vals[course_type_idx])
            
            level_idx = col_indices.get('level', 6)
            level = parse_level(row_vals[level_idx])
            
            semester_idx = col_indices.get('semester', 7)
            semester = parse_val(row_vals[semester_idx])
            
            other_progs_idx = col_indices.get('other_programs')
            other_progs = parse_val(row_vals[other_progs_idx]) if other_progs_idx is not None else ""
            
            req_idx = col_indices.get('requirement', 9)
            requirement = parse_val(row_vals[req_idx])
            
            prereq_idx = col_indices.get('prerequisite', 10)
            prereq = parse_val(row_vals[prereq_idx])
            
            concurrent_idx = col_indices.get('concurrent_courses', 11)
            concurrent = parse_val(row_vals[concurrent_idx])
            
            added_gpa_idx = col_indices.get('added_to_gpa', 12)
            added_to_gpa = parse_val(row_vals[added_gpa_idx])
            
            pass_fail_idx = col_indices.get('pass_fail', 13)
            pass_fail = parse_val(row_vals[pass_fail_idx])
            
            credit_hours_idx = col_indices.get('credit_hours', 14)
            credit_hours = parse_float(row_vals[credit_hours_idx])

            theory_hours = 0.0
            exercise_hours = 0.0
            practical_hours = 0.0
            activity_hours = 0.0
            exam_time_hours = ""
            
            theory_grade = 0.0
            practical_grade = 0.0
            year_work_grade = 0.0
            midterm_grade = ""
            oral_grade = ""
            total_grade = 0.0
            
            success_rate = ""
            fail_rate = ""
            group_code = ""
            elective_count = ""
            summer_reg = ""

            for c_i, header_name in col_map.items():
                val = row_vals[c_i]
                if "الساعات" in header_name:
                    if "محاضرات" in header_name:
                        theory_hours = parse_float(val)
                    elif "تدريب" in header_name:
                        exercise_hours = parse_float(val)
                    elif "عملي" in header_name:
                        practical_hours = parse_float(val)
                    elif "الميداني" in header_name:
                        activity_hours = parse_float(val)
                    elif "الامتحان" in header_name:
                        exam_time_hours = parse_val(val)
                elif "الدرجات" in header_name:
                    if "نهاية الفصل" in header_name:
                        theory_grade = parse_float(val)
                    elif "عملي" in header_name:
                        practical_grade = parse_float(val)
                    elif "أعمال فصل" in header_name:
                        year_work_grade = parse_float(val)
                    elif "منتصف" in header_name:
                        midterm_grade = parse_val(val)
                    elif "شفوي" in header_name:
                        oral_grade = parse_val(val)
                    elif "المجموع" in header_name:
                        total_grade = parse_float(val)
                elif "نسبة النجاح" in header_name:
                    success_rate = parse_val(val)
                elif "الرسوب" in header_name:
                    fail_rate = parse_val(val)
                elif "المجموعة الاختيارية" in header_name:
                    group_code = parse_val(val)
                elif "الوحدات الاختيارية" in header_name or "المقررات الاختيارية" in header_name:
                    elective_count = parse_val(val)
                elif "الصيفى" in header_name:
                    summer_reg = parse_val(val)

            # Check duplicate
            if code in existing_codes:
                skipped_sheet += 1
                continue

            cursor.execute("""
                INSERT INTO courses (
                    code, name_ar, name_en, level, semester, course_type,
                    credit_hours, theory_hours, practical_hours, exercise_hours, activity_hours,
                    total_grade, theory_grade, practical_grade, year_work_grade, midterm_grade, oral_grade,
                    exam_time_hours, requirement, added_to_gpa, pass_fail, summer_registration,
                    other_programs, prerequisite, concurrent_courses, success_rate, fail_rate,
                    elective_group_code, elective_courses_count, program_id, faculty_id,
                    year, is_deleted, is_bundle
                ) VALUES (
                    ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?,
                    ?, 0, 0
                )
            """, (
                code, name_ar, name_en, level, semester, course_type,
                credit_hours, theory_hours, practical_hours, exercise_hours, activity_hours,
                total_grade, theory_grade, practical_grade, year_work_grade, midterm_grade, oral_grade,
                exam_time_hours, requirement, added_to_gpa, pass_fail, summer_reg,
                other_progs, prereq, concurrent, success_rate, fail_rate,
                group_code, elective_count, prog_id, faculty_id,
                '2026/2027'
            ))

            existing_codes.add(code)
            inserted_sheet += 1

        print(f"   -> Inserted: {inserted_sheet} new courses | Skipped (already exist): {skipped_sheet}")
        total_inserted += inserted_sheet
        total_skipped += skipped_sheet

    conn.commit()
    conn.close()

    print("\n==================================================")
    print(f"SUCCESS: Engineering course registration complete!")
    print(f"Total new courses registered: {total_inserted}")
    print(f"Total existing courses skipped: {total_skipped}")
    print("==================================================")

if __name__ == '__main__':
    main()
