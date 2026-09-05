r"""
Import courses for كلية تكنولوجيا العلوم الصحية التطبيقية from Excel file.
Excel file location: C:\Users\rubas\Desktop\كلية تكنولوجيا العلوم الصحية التطبيقية.xlsx

Sheet names map to programs:
  - برنامج عام                           -> program_id = 18 (برنامج عام)
  -  برنامج تكنولوجيا الرعاية التن       -> program_id = 19 (تكنولوجيا الرعاية التنفسية)
  - برنامج تكنولوجيا المختبرات الط        -> program_id = 20 (تكنولوجيا المختبرات الطبية)
  - برنامج تكنولوجيا صناعة تركيبات       -> program_id = 21 (تكنولوجيا صناعة تركيبات الأسنان)
  - برنامج تكنولوجيا علوم الأشعة و        -> program_id = 22 (تكنولوجيا علوم الأشعة والتصوير الطبي)
"""

import openpyxl
import sqlite3
import os
import sys

EXCEL_FILE = r'C:\Users\rubas\Desktop\كلية تكنولوجيا العلوم الصحية التطبيقية.xlsx'
DB_FILE = os.path.join(os.path.dirname(__file__), 'mnu_system.db')

def get_program_id(sheet_name):
    s = sheet_name.strip()
    if 'الرعاية' in s:
        return 19
    if 'المختبرات' in s:
        return 20
    if 'تركيبات' in s:
        return 21
    if 'الأشعة' in s:
        return 22
    if 'عام' in s:
        return 18
    return None

def parse_val(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s in ("None", "none", "-", "–"):
        return ""
    return s

def parse_float(val):
    if val is None:
        return 0.0
    try:
        return float(str(val).strip())
    except:
        return 0.0

def parse_level(val):
    if not val:
        return 1
    s = str(val).strip()
    if "الأول" in s or "الاول" in s or s == "1":
        return 1
    if "الثاني" in s or "الثانى" in s or s == "2":
        return 2
    if "الثالث" in s or s == "3":
        return 3
    if "الرابع" in s or s == "4":
        return 4
    if "الخامس" in s or s == "5":
        return 5
    return 1

def main():
    if not os.path.exists(EXCEL_FILE):
        print(f"ERROR: Excel file not found at {EXCEL_FILE}")
        return

    if not os.path.exists(DB_FILE):
        print(f"ERROR: Database file not found at {DB_FILE}")
        return

    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # Verify Faculty 8 ID
    fac_row = cursor.execute("SELECT id, name FROM faculties WHERE id = 8 OR name LIKE '%تكنولوجيا العلوم الصحية%'").fetchone()
    if not fac_row:
        print("ERROR: Faculty not found!")
        return
    faculty_id = fac_row['id']
    print(f"Target Faculty: {fac_row['name']} (ID: {faculty_id})")

    # Fetch programs
    progs = cursor.execute("SELECT id, name FROM programs WHERE faculty_id = ?", (faculty_id,)).fetchall()
    prog_dict = {p['id']: p['name'] for p in progs}
    print("Faculty Programs in DB:", prog_dict)

    wb = openpyxl.load_workbook(EXCEL_FILE)

    total_inserted = 0
    total_skipped = 0

    print("\n--- Starting Course Registration Process ---")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        prog_id = get_program_id(sheet_name)
        if not prog_id:
            print(f"WARNING: Could not map sheet '{sheet_name}' to any program. Skipping.")
            continue

        prog_name = prog_dict.get(prog_id, 'Unknown Program')
        print(f"\nProcessing Sheet: '{sheet_name}' -> Program: '{prog_name}' (ID: {prog_id})")

        # Map headers from Row 3 and Row 4
        r3 = [ws.cell(3, col).value for col in range(1, ws.max_column + 1)]
        r4 = [ws.cell(4, col).value for col in range(1, ws.max_column + 1)]
        
        col_map = {}
        current_section = ""
        for idx in range(ws.max_column):
            val3 = str(r3[idx]).strip() if r3[idx] is not None else ""
            val4 = str(r4[idx]).strip() if r4[idx] is not None else ""
            if val3 and val3 not in ("None", ""):
                current_section = val3
            full_header = f"{current_section} - {val4}" if val4 else current_section
            col_map[idx] = full_header

        # Fetch existing course codes for this program
        existing_in_prog = cursor.execute("SELECT code FROM courses WHERE program_id = ?", (prog_id,)).fetchall()
        existing_codes = set(r['code'].strip() for r in existing_in_prog if r['code'])

        inserted_sheet = 0
        skipped_sheet = 0

        for r_idx in range(5, ws.max_row + 1):
            seq = ws.cell(r_idx, 1).value
            if seq is None or not str(seq).strip().isdigit():
                continue

            row_vals = [ws.cell(r_idx, c_idx + 1).value for c_idx in range(ws.max_column)]
            
            code = parse_val(row_vals[1])
            if not code:
                continue

            code_en = parse_val(row_vals[2]) or code
            raw_name_ar = parse_val(row_vals[3])
            name_en = parse_val(row_vals[4])
            
            name_ar = raw_name_ar
            if code and name_ar.endswith(code):
                name_ar = name_ar[:-len(code)].strip()
            if '\n' in name_ar:
                parts = [p.strip() for p in name_ar.split('\n') if p.strip()]
                if len(parts) > 1 and parts[-1] == code:
                    name_ar = " ".join(parts[:-1])
                else:
                    name_ar = " ".join(parts)

            course_type = parse_val(row_vals[5])
            level = parse_level(row_vals[6])
            semester = parse_val(row_vals[7])
            other_progs = parse_val(row_vals[8])
            requirement = parse_val(row_vals[9])
            prereq = parse_val(row_vals[10])
            concurrent = parse_val(row_vals[11])
            added_to_gpa = parse_val(row_vals[12])
            pass_fail = parse_val(row_vals[13])
            credit_hours = parse_float(row_vals[14])

            theory_hours = 0.0
            exercise_hours = 0.0
            practical_hours = 0.0
            activity_hours = 0.0
            exam_time_hours = ""
            
            theory_grade = 0.0
            practical_grade = 0.0
            year_work_grade = 0.0
            midterm_grade = ""
            total_grade = 0.0
            
            success_rate = ""
            fail_rate = ""
            group_code = ""
            elective_count = ""
            summer_reg = ""

            for c_i, header_name in col_map.items():
                val = row_vals[c_i]
                if "الساعات" in header_name:
                    if "محاضرات" in header_name:
                        theory_hours = parse_float(val)
                    elif "تدريب" in header_name:
                        exercise_hours = parse_float(val)
                    elif "عملي" in header_name:
                        practical_hours = parse_float(val)
                    elif "الميداني" in header_name:
                        activity_hours = parse_float(val)
                    elif "الامتحان" in header_name:
                        exam_time_hours = parse_val(val)
                elif "الدرجات" in header_name:
                    if "نهاية الفصل" in header_name:
                        theory_grade = parse_float(val)
                    elif "عملي" in header_name:
                        practical_grade = parse_float(val)
                    elif "أعمال فصل" in header_name:
                        year_work_grade = parse_float(val)
                    elif "منتصف الفصل" in header_name:
                        midterm_grade = parse_val(val)
                    elif "المجموع" in header_name:
                        total_grade = parse_float(val)
                elif "نسبة النجاح" in header_name:
                    success_rate = parse_val(val)
                elif "الرسوب" in header_name:
                    fail_rate = parse_val(val)
                elif "المجموعة الاختيارية" in header_name:
                    group_code = parse_val(val)
                elif "الوحدات الاختيارية" in header_name or "المقررات الاختيارية" in header_name:
                    elective_count = parse_val(val)
                elif "الصيفى" in header_name:
                    summer_reg = parse_val(val)

            # Check if course already registered in this program
            if code in existing_codes:
                skipped_sheet += 1
                continue

            # Insert course
            cursor.execute("""
                INSERT INTO courses (
                    code, name_ar, name_en, level, semester, course_type,
                    credit_hours, theory_hours, practical_hours, exercise_hours, activity_hours,
                    total_grade, theory_grade, practical_grade, year_work_grade, midterm_grade,
                    exam_time_hours, requirement, added_to_gpa, pass_fail, summer_registration,
                    other_programs, prerequisite, concurrent_courses, success_rate, fail_rate,
                    elective_group_code, elective_courses_count, program_id, faculty_id,
                    year, is_deleted, is_bundle
                ) VALUES (
                    ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?,
                    ?, 0, 0
                )
            """, (
                code, name_ar, name_en, level, semester, course_type,
                credit_hours, theory_hours, practical_hours, exercise_hours, activity_hours,
                total_grade, theory_grade, practical_grade, year_work_grade, midterm_grade,
                exam_time_hours, requirement, added_to_gpa, pass_fail, summer_reg,
                other_progs, prereq, concurrent, success_rate, fail_rate,
                group_code, elective_count, prog_id, faculty_id,
                '2026/2027'
            ))

            existing_codes.add(code)
            inserted_sheet += 1

        print(f"   -> Inserted: {inserted_sheet} new courses | Skipped (already exist): {skipped_sheet}")
        total_inserted += inserted_sheet
        total_skipped += skipped_sheet

    conn.commit()
    conn.close()

    print("\n==================================================")
    print(f"SUCCESS: Registration complete!")
    print(f"Total new courses registered: {total_inserted}")
    print(f"Total existing courses skipped: {total_skipped}")
    print("==================================================")

if __name__ == '__main__':
    main()
