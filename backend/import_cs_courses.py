"""
Import courses for كلية علوم الحاسوب والذكاء الاصطناعي from Excel file.
Sheet names map to programs: 
  برنامج انترنت الاشياء وتحليل ال  -> program_id = 7
  برنامج ذكاء الآلة                -> program_id = 8
  برنامج علوم البيانات             -> program_id = 9
"""

import openpyxl
import sys
sys.path.insert(0, '.')
from database import SessionLocal
from sqlalchemy import text

EXCEL_FILE = r'C:\Users\rubas\Desktop\كلية علوم الحاسوب والذكاء الاصطناعي.xlsx'

# Mapping: sheet name fragment -> program_id
PROGRAM_MAP = {
    'انترنت الاشياء': 7,
    'ذكاء الآلة': 8,
    'علوم البيانات': 9,
}

# Faculty ID for كلية علوم الحاسوب والذكاء الاصطناعي
FACULTY_ID = None  # Will be fetched from DB

def normalize(val):
    if val is None:
        return None
    s = str(val).strip()
    # Remove embedded newlines and extra spaces
    s = ' '.join(s.split())
    if s in ('', '-', '–', 'None', 'none'):
        return None
    return s

def parse_float(val):
    if val is None:
        return 0.0
    try:
        return float(str(val).strip())
    except:
        return 0.0

def map_requirement(val):
    """Map Arabic requirement text to English enum"""
    if val is None:
        return 'faculty'
    v = str(val).strip().lower()
    if 'جامعة' in v:
        return 'university'
    if 'كلية' in v or 'كلية' in v:
        return 'faculty'
    if 'تخصص' in v or 'برنامج' in v:
        return 'major'
    return 'faculty'

def map_course_type(val):
    """Map Arabic course type to English enum"""
    if val is None:
        return 'mandatory'
    v = str(val).strip()
    if 'اختيار' in v:
        return 'elective'
    return 'mandatory'

def map_level(val):
    """Map Arabic level text to integer"""
    if val is None:
        return None
    v = str(val).strip()
    if 'أول' in v or 'اول' in v or '1' in v:
        return 1
    if 'ثاني' in v or 'ثانى' in v or '2' in v:
        return 2
    if 'ثالث' in v or '3' in v:
        return 3
    if 'رابع' in v or '4' in v:
        return 4
    if 'خامس' in v or '5' in v:
        return 5
    if 'سادس' in v or '6' in v:
        return 6
    if 'سابع' in v or '7' in v:
        return 7
    if 'عام' in v or '0' in v:
        return 0
    return None

def map_semester(val):
    """Map Arabic semester text to integer"""
    if val is None:
        return None
    v = str(val).strip()
    if 'أول' in v or 'اول' in v or '1' in v:
        return 1
    if 'ثاني' in v or 'ثانى' in v or '2' in v:
        return 2
    if 'صيفي' in v or 'صيف' in v:
        return 3
    return None

def read_sheet(ws):
    """Read courses from a sheet. Data rows start at row 5 (0-indexed row 4)."""
    courses = []
    # Row 3 (index 2) = header row with column names
    # Actual column indices:
    # 0=م, 1=كود المقرر, 2=كود انجليزى, 3=اسم المقرر, 4=اسم انجليزى
    # 5=نوع المقرر, 6=المستوى, 7=الفصل, 8=البرامج الأخرى
    # 9=متطلب, 10=المتطلب السابق, 11=المتزامنة
    # 12=يضاف للمعدل التراكمي, 13=نجاح او رسوب
    # 14=الساعات المعتمدة
    # 15=محاضرات(نظري), 16=تدريب, 17=عملي
    # 30=نسبة النجاح, 34=تسجيل صيفي
    
    for row in ws.iter_rows(min_row=5, values_only=True):
        # Skip empty or summary rows
        seq = row[0]
        if seq is None:
            continue
        try:
            int(str(seq).strip())
        except:
            continue  # Not a data row
        
        code = normalize(row[1])
        if not code:
            continue
        
        # Extract name_ar: remove the appended code from the name
        name_raw = normalize(row[3])
        if name_raw and code and name_raw.endswith(code):
            name_raw = name_raw[:-len(code)].strip()
        name_ar = name_raw
        
        name_en = normalize(row[4])
        course_type = map_course_type(row[5])
        level = map_level(row[6])
        semester = map_semester(row[7])
        requirement = map_requirement(row[9])
        added_to_gpa_raw = normalize(row[12])
        added_to_gpa = 'نعم' if added_to_gpa_raw and 'يضاف' in added_to_gpa_raw else 'لا'
        
        # Pass/fail
        pass_fail_raw = normalize(row[13])
        summer_reg_raw = normalize(row[34]) if len(row) > 34 else None
        summer_registration = 'نعم' if summer_reg_raw and 'نعم' in summer_reg_raw else 'لا'
        
        credit_hours = parse_float(row[14])
        theory_hours = parse_float(row[15])
        exercise_hours = parse_float(row[16])  # تدريب
        practical_hours = parse_float(row[17])  # عملي
        
        # success rate
        success_rate = parse_float(row[30]) if len(row) > 30 else None
        total_grade = parse_float(row[29]) if len(row) > 29 else None
        
        courses.append({
            'code': code,
            'name_ar': name_ar,
            'name_en': name_en,
            'course_type': course_type,
            'level': level,
            'semester': semester,
            'requirement': requirement,
            'added_to_gpa': added_to_gpa,
            'summer_registration': summer_registration,
            'credit_hours': credit_hours,
            'theory_hours': theory_hours,
            'exercise_hours': exercise_hours,
            'practical_hours': practical_hours,
            'total_grade': total_grade if total_grade else 100.0,
        })
    
    return courses


def main():
    session = SessionLocal()
    try:
        # Get faculty_id
        faculty_row = session.execute(text(
            "SELECT id FROM faculties WHERE name LIKE '%حاسوب%' OR name LIKE '%ذكاء الاصطناعي%' OR name LIKE '%ذكاء%' LIMIT 1"
        )).fetchone()
        if not faculty_row:
            print("ERROR: Faculty not found!")
            return
        faculty_id = faculty_row[0]
        print(f"Faculty ID: {faculty_id}")
        
        # Load Excel
        wb = openpyxl.load_workbook(EXCEL_FILE)
        
        total_inserted = 0
        total_skipped = 0
        
        for sheet_name in wb.sheetnames:
            # Find matching program
            prog_id = None
            for key, pid in PROGRAM_MAP.items():
                if key in sheet_name:
                    prog_id = pid
                    break
            
            if prog_id is None:
                print(f"WARNING: No program mapping for sheet '{sheet_name}', skipping.")
                continue
            
            ws = wb[sheet_name]
            courses = read_sheet(ws)
            print(f"\nSheet '{sheet_name}' -> program_id={prog_id}: {len(courses)} courses found")
            
            inserted = 0
            skipped = 0
            for c in courses:
                # Check if already exists (same code + program)
                existing = session.execute(text(
                    "SELECT id FROM courses WHERE code = :code AND program_id = :prog_id LIMIT 1"
                ), {'code': c['code'], 'prog_id': prog_id}).fetchone()
                
                if existing:
                    skipped += 1
                    continue
                
                session.execute(text("""
                    INSERT INTO courses (
                        code, name_ar, name_en, course_type, level, semester,
                        requirement, added_to_gpa, summer_registration,
                        credit_hours, theory_hours, exercise_hours, practical_hours,
                        total_grade, faculty_id, program_id
                    ) VALUES (
                        :code, :name_ar, :name_en, :course_type, :level, :semester,
                        :requirement, :added_to_gpa, :summer_registration,
                        :credit_hours, :theory_hours, :exercise_hours, :practical_hours,
                        :total_grade, :faculty_id, :program_id
                    )
                """), {**c, 'faculty_id': faculty_id, 'program_id': prog_id})
                inserted += 1
            
            session.commit()
            total_inserted += inserted
            total_skipped += skipped
            print(f"  Inserted: {inserted}, Skipped (already exist): {skipped}")
        
        print(f"\nDone! Total inserted: {total_inserted}, Total skipped: {total_skipped}")
        
    except Exception as e:
        session.rollback()
        print(f"ERROR: {e}")
        import traceback
        traceback.print_exc()
    finally:
        session.close()


if __name__ == '__main__':
    main()
