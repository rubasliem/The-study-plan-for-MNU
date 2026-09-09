from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, status, Request
from fastapi.security import OAuth2PasswordRequestForm
from typing import Optional, List
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, JSONResponse
from sqlalchemy.orm import Session
from datetime import datetime
from sqlalchemy import inspect, text, func, or_
from database import engine, get_db, SessionLocal
import models, schemas, auth
import json
import re
import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from services.course_service import process_and_replace_courses

# 1. تهيئة التطبيق الأساسي
app = FastAPI(
    title="MNU Study Plan API", 
    version="1.0",
)

# 2. إعداد الـ CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # السماح لكل المواقع (للتطوير فقط)
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# استخدام contextvars بدلاً من threading.local لتخزين بيانات الطلب لكي يعمل بشكل صحيح مع FastAPI
import contextvars

academic_year_var = contextvars.ContextVar("academic_year", default=None)
semester_var = contextvars.ContextVar("semester", default=None)

def get_notif_context():
    """قراءة العام الجامعي والفصل الدراسي من السياق الحالي للطلب"""
    return academic_year_var.get(), semester_var.get()

class ContextMiddleware:
    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] not in ["http", "websocket"]:
            await self.app(scope, receive, send)
            return

        # استخراج الهيدرز من scope
        headers = dict(scope.get("headers", []))
        year = headers.get(b"x-academic-year", b"").decode("utf-8").strip() or None
        semester = headers.get(b"x-semester", b"").decode("utf-8").strip() or None
        
        # تعيين القيم في contextvars للطلب الحالي
        token_year = academic_year_var.set(year)
        token_semester = semester_var.set(semester)
        
        try:
            await self.app(scope, receive, send)
        finally:
            academic_year_var.reset(token_year)
            semester_var.reset(token_semester)

app.add_middleware(ContextMiddleware)


# 3. إنشاء الجداول في قاعدة البيانات
models.Base.metadata.create_all(bind=engine)

def get_user_role_display(user, db: Session = None) -> str:
    if not user:
        return "مستخدم"
    clean_job = ""
    if hasattr(user, 'job_title') and user.job_title and user.job_title.strip():
        clean_job = re.sub(r'\s*\([^\)]*\)', '', user.job_title.strip()).strip()
    if not clean_job:
        role_val = getattr(user, 'role', '')
        role_str = getattr(role_val, 'value', str(role_val))
        role_map = {
            'admin': 'مدير عام',
            'manager': 'مسؤول إدارة',
            'student_affairs': 'مدير شؤون الطلاب',
            'faculty_professor': 'مدير برنامج',
            'reviewer': 'المراجع',
            'faculty_admin': 'مسؤول كلية'
        }
        clean_job = role_map.get(role_str, 'مسؤول كلية')
        
    is_fac_admin = (getattr(user, 'role', None) == models.UserRole.faculty_admin) or ("مسؤول كلية" in clean_job)
    is_prog_mgr = (getattr(user, 'role', None) == models.UserRole.faculty_professor) or ("مدير برنامج" in clean_job)
    is_fac_member = ("عضو هيئة تدريس" in clean_job)
    
    colleges = []
    if getattr(user, 'assigned_faculties', None):
        colleges = [f.name.replace("كلية", "").strip() for f in user.assigned_faculties if f.name]
    if not colleges and getattr(user, 'faculty', None) and user.faculty and user.faculty.name:
        colleges = [user.faculty.name.replace("كلية", "").strip()]
    elif not colleges and getattr(user, 'faculty_id', None) and db:
        fac = db.query(models.Faculty).filter(models.Faculty.id == user.faculty_id).first()
        if fac and fac.name:
            colleges = [fac.name.replace("كلية", "").strip()]
            
    unique_colleges = list(dict.fromkeys([c for c in colleges if c]))
    colleges_str = " - ".join(unique_colleges)
    
    if is_fac_member:
        return f"{clean_job} ({colleges_str})" if colleges_str else clean_job
    elif is_fac_admin:
        if clean_job.endswith("كلية"):
            return f"{clean_job} {colleges_str}" if colleges_str else clean_job
        elif "كلية" in clean_job:
            return f"{clean_job} ({colleges_str})" if colleges_str else clean_job
        else:
            return f"{clean_job} كلية {colleges_str}" if colleges_str else clean_job
    elif is_prog_mgr:
        return f"{clean_job} {colleges_str}" if colleges_str else clean_job
    else:
        return clean_job

def get_user_action_by(user, db: Session = None) -> str:
    if not user:
        return ""
    role_display = get_user_role_display(user, db)
    return f"{user.username} ({role_display})"

def log_activity(
    db: Session,
    username: str,
    action_type: str,
    description: str,
    entity_type: str = "SYSTEM",
    user_id: Optional[int] = None,
    user_role: Optional[str] = None,
    status: str = "success",
    faculty_id: Optional[int] = None,
    academic_year: Optional[str] = None,
    semester: Optional[str] = None,
    ip_address: Optional[str] = None,
    extra_data: Optional[str] = None
):
    """تسجيل حركة أو عملية في جدول سجل العمليات (activity_logs)"""
    try:
        ctx_year, ctx_semester = get_notif_context()
        if not academic_year:
            academic_year = ctx_year
        if not semester:
            semester = ctx_semester
        
        log_entry = models.ActivityLog(
            user_id=user_id,
            username=username or "نظام",
            user_role=user_role or "مستخدم",
            action_type=action_type,
            entity_type=entity_type,
            description=description,
            status=status,
            faculty_id=faculty_id,
            academic_year=academic_year,
            semester=semester,
            ip_address=ip_address,
            extra_data=extra_data,
            created_at=datetime.utcnow()
        )
        db.add(log_entry)
        db.commit()
        return log_entry
    except Exception as e:
        print(f"Error in log_activity: {e}")
        db.rollback()
        return None

def create_notification(db, faculty_id, action_by, action_text, academic_year=None, semester=None):
    """helper لإنشاء إشعار مع العام الجامعي والفصل الدراسي بدقة تامة وتسجيل العملية في سجل النشاطات"""
    import re
    # 1. إذا لم يُمرر العام الجامعي، نستخرجه من نص الحدث إن وجد (الأولوية لمحتوى الحدث الصريح)
    if not academic_year and action_text:
        match_y = re.search(r'(\d{4}\s*[/\\-]\s*\d{4})', action_text)
        if match_y:
            academic_year = match_y.group(1).replace('-', '/').replace(' ', '')

    # 2. إذا لم يُمرر الفصل الدراسي، نستخرجه من نص الحدث إن وجد
    if not semester and action_text:
        if "الفصل الدراسي الأول" in action_text or "فصل أول" in action_text or "الترم الأول" in action_text:
            semester = "الفصل الدراسي الأول"
        elif "الفصل الدراسي الثاني" in action_text or "فصل ثاني" in action_text or "الترم الثاني" in action_text:
            semester = "الفصل الدراسي الثاني"
        elif "الفصل الدراسي الصيفي" in action_text or "فصل صيفي" in action_text or "الترم الصيفي" in action_text:
            semester = "الفصل الدراسي الصيفي"

    # 3. إذا لم يوجد، نأخذه من سياق الطلب
    ctx_year, ctx_semester = get_notif_context()
    if not academic_year:
        academic_year = ctx_year
    if not semester:
        semester = ctx_semester

    if hasattr(action_by, 'username'):
        action_by_str = get_user_action_by(action_by, db)
    else:
        action_by_str = str(action_by) if action_by else ""
        if action_by_str and db:
            m = re.match(r"^([^\(]+)(?:\s*\((.+)\))?$", action_by_str.strip())
            if m:
                raw_uname = m.group(1).strip()
                clean_uname = raw_uname.replace('@gmail.com', '').strip()
                u = db.query(models.User).filter((models.User.username == clean_uname) | (models.User.username == raw_uname)).first()
                if u:
                    action_by_str = get_user_action_by(u, db)
    notif = models.Notification(
        faculty_id=faculty_id,
        action_by=action_by_str,
        action_text=action_text,
        created_at=datetime.utcnow(),
        academic_year=academic_year,
        semester=semester
    )
    db.add(notif)
    
    # مزامنة تلقائية مع سجل العمليات ActivityLog
    try:
        uname = "نظام"
        urole = "نظام"
        if action_by_str:
            m = re.match(r'(.+?)\s*\((.+?)\)', action_by_str)
            if m:
                uname = m.group(1).strip()
                urole = m.group(2).strip()
            else:
                uname = action_by_str.strip()
                urole = uname
                
        atype = "UPDATE"
        txt = action_text or ""
        if any(w in txt for w in ["حذف", "مسح"]):
            atype = "DELETE"
        elif any(w in txt for w in ["إضافة", "اضافة", "جديد"]):
            atype = "CREATE"
        elif any(w in txt for w in ["اعتماد"]):
            atype = "APPROVE"
        elif any(w in txt for w in ["استرجاع"]):
            atype = "RESTORE"
        elif any(w in txt for w in ["طباعة", "تصدير", "تنزيل"]):
            atype = "EXPORT"
        
        etype = "SYSTEM"
        if "هيئة التدريس" in txt or "عضو" in txt:
            etype = "PROFESSORS"
        elif "مقرر" in txt:
            etype = "COURSES"
        elif "الخطة" in txt:
            etype = "STUDY_PLAN"
        elif "مستخدم" in txt or "صلاحية" in txt:
            etype = "USERS"
        elif "توقيع" in txt:
            etype = "SIGNATURES"
        elif "المحذوف" in txt:
            etype = "RECYCLE_BIN"
            
        act_log = models.ActivityLog(
            username=uname,
            user_role=urole,
            action_type=atype,
            entity_type=etype,
            description=txt,
            status="success",
            faculty_id=faculty_id,
            academic_year=academic_year,
            semester=semester,
            created_at=datetime.utcnow()
        )
        db.add(act_log)
    except Exception as e:
        print(f"Error auto-recording ActivityLog from notification: {e}")

    return notif


def migrate_db_add_email():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        columns = [c['name'] for c in inspector.get_columns('professors')]
        if 'email' not in columns:
            db.execute(text("ALTER TABLE professors ADD COLUMN email VARCHAR;"))
            db.commit()
            print("Successfully added 'email' column to 'professors' table.")
    except Exception as e:
        print(f"Error checking/adding email column: {e}")
    finally:
        db.close()

def migrate_db_add_year():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        
        course_cols = [c['name'] for c in inspector.get_columns('courses')]
        if 'year' not in course_cols:
            db.execute(text("ALTER TABLE courses ADD COLUMN year VARCHAR DEFAULT '2026/2027';"))
            db.commit()
            print("Successfully added 'year' column to 'courses' table.")
            
        prof_cols = [c['name'] for c in inspector.get_columns('professors')]
        if 'academic_year' not in prof_cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN academic_year VARCHAR DEFAULT '2026/2027';"))
            db.commit()
            print("Successfully added 'academic_year' column to 'professors' table.")
    except Exception as e:
        print(f"Error checking/adding year columns: {e}")
    finally:
        db.close()

def migrate_db_add_base_course_id():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('study_plan_items')]
        if 'base_course_id' not in cols:
            db.execute(text("ALTER TABLE study_plan_items ADD COLUMN base_course_id INTEGER;"))
            db.commit()
            print("Successfully added 'base_course_id' column to 'study_plan_items' table.")
    except Exception as e:
        print(f"Error checking/adding base_course_id column: {e}")
    finally:
        db.close()

def migrate_db_add_signature_order():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('signatures')]
        if 'order_index' not in cols:
            db.execute(text("ALTER TABLE signatures ADD COLUMN order_index INTEGER DEFAULT 0;"))
            db.commit()
            print("Successfully added 'order_index' column to 'signatures' table.")
    except Exception as e:
        print(f"Error checking/adding order_index column to signatures: {e}")
    finally:
        db.close()
def migrate_db_add_user_permissions():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('users')]
        columns_to_add = [
            "all_faculties_access",
            "perm_view_prof_study_plan",
            "perm_view_prof_data_btn",
            "perm_view_prof_delete_btn",
            "perm_view_prof_view_btn",
            "perm_view_prof_print_btn",
            "perm_view_prof_import_btn",
            "perm_delete_notif_btn",
            "perm_recycle_restore_btn",
            "perm_recycle_delete_btn"
        ]
        
        added_any = False
        for col in columns_to_add:
            if col not in cols:
                db.execute(text(f"ALTER TABLE users ADD COLUMN {col} BOOLEAN DEFAULT FALSE;"))
                added_any = True
                
        if 'job_title' not in cols:
            db.execute(text("ALTER TABLE users ADD COLUMN job_title VARCHAR;"))
            added_any = True

        if added_any:
            db.commit()
            print("Successfully added missing permission columns to 'users' table.")
    except Exception as e:
        print(f"Error checking/adding permission columns to users: {e}")
    finally:
        db.close()

def migrate_db_add_entry_group_id():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('study_plan_items')]
        if 'entry_group_id' not in cols:
            db.execute(text("ALTER TABLE study_plan_items ADD COLUMN entry_group_id VARCHAR;"))
            db.commit()
            print("Successfully added 'entry_group_id' column to 'study_plan_items' table.")
    except Exception as e:
        print(f"Error checking/adding entry_group_id column: {e}")
    finally:
        db.close()

def migrate_db_add_is_bundle():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('courses')]
        if 'is_bundle' not in cols:
            db.execute(text("ALTER TABLE courses ADD COLUMN is_bundle BOOLEAN DEFAULT FALSE;"))
            db.commit()
            print("Successfully added 'is_bundle' column to 'courses' table.")
    except Exception as e:
        print(f"Error checking/adding is_bundle column: {e}")
    finally:
        db.close()

def migrate_db_add_medicine_plan_fields():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('study_plan_items')]
        if 'module_id' not in cols:
            db.execute(text("ALTER TABLE study_plan_items ADD COLUMN module_id INTEGER;"))
            db.commit()
            print("Successfully added 'module_id' column to 'study_plan_items' table.")
        if 'hours_actual_activity' not in cols:
            db.execute(text("ALTER TABLE study_plan_items ADD COLUMN hours_actual_activity FLOAT DEFAULT 0.0;"))
            db.commit()
        if 'course_notes' not in cols:
            db.execute(text("ALTER TABLE study_plan_items ADD COLUMN course_notes VARCHAR;"))
            db.commit()
            print("Successfully added 'course_notes' column to 'study_plan_items' table.")
    except Exception as e:
        print(f"Error checking/adding medicine plan fields: {e}")
    finally:
        db.close()


def migrate_db_add_signature_report_type():
    db = SessionLocal()
    try:
        db.execute(text("ALTER TABLE signatures ADD COLUMN report_type VARCHAR DEFAULT 'الخطة الدراسية';"))
        db.commit()
    except Exception as e:
        db.rollback()
    finally:
        db.close()

def migrate_db_add_user_security_columns():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('users')]
        added_any = False
        if 'security_question' not in cols:
            db.execute(text("ALTER TABLE users ADD COLUMN security_question VARCHAR;"))
            added_any = True
        if 'security_answer' not in cols:
            db.execute(text("ALTER TABLE users ADD COLUMN security_answer VARCHAR;"))
            added_any = True
        
        if added_any:
            db.commit()
            print("Successfully added security columns to 'users' table.")
    except Exception as e:
        print(f"Error checking/adding security columns to users: {e}")
    finally:
        db.close()

def migrate_db_add_academic_years():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        if 'academic_years' in inspector.get_table_names():
            cols = [c['name'] for c in inspector.get_columns('academic_years')]
            if 'semester1_weeks' not in cols:
                db.execute(text("ALTER TABLE academic_years ADD COLUMN semester1_weeks INTEGER DEFAULT 15;"))
            if 'semester2_weeks' not in cols:
                db.execute(text("ALTER TABLE academic_years ADD COLUMN semester2_weeks INTEGER DEFAULT 15;"))
            if 'summer_weeks' not in cols:
                db.execute(text("ALTER TABLE academic_years ADD COLUMN summer_weeks INTEGER DEFAULT 8;"))
            if 'med_semester1_weeks' not in cols:
                db.execute(text("ALTER TABLE academic_years ADD COLUMN med_semester1_weeks INTEGER DEFAULT 15;"))
            if 'med_semester2_weeks' not in cols:
                db.execute(text("ALTER TABLE academic_years ADD COLUMN med_semester2_weeks INTEGER DEFAULT 14;"))
            if 'med_summer_weeks' not in cols:
                db.execute(text("ALTER TABLE academic_years ADD COLUMN med_summer_weeks INTEGER DEFAULT 7;"))
            db.commit()

        count = db.query(models.AcademicYear).count()
        if count == 0:
            default_years = ["2024/2025", "2025/2026", "2026/2027", "2027/2028"]
            for year_name in default_years:
                db.add(models.AcademicYear(name=year_name, semester1_weeks=15, semester2_weeks=15, summer_weeks=8))
            db.commit()
            print("Successfully populated default academic years.")
    except Exception as e:
        print(f"Error checking/adding academic years: {e}")
    finally:
        db.close()

def migrate_db_add_contract_fields():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('professors')]
        if 'contract_type' not in cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN contract_type VARCHAR;"))
        if 'work_days' not in cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN work_days VARCHAR;"))
        if 'mnu_job_title' not in cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN mnu_job_title VARCHAR;"))
        if 'semester1_weeks' not in cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN semester1_weeks INTEGER;"))
        if 'semester2_weeks' not in cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN semester2_weeks INTEGER;"))
        if 'summer_weeks' not in cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN summer_weeks INTEGER;"))
        if 'academic_year_weeks' not in cols:
            db.execute(text("ALTER TABLE professors ADD COLUMN academic_year_weeks VARCHAR;"))
        db.commit()
        print("Successfully added contract and semester weeks fields to 'professors' table.")
    except Exception as e:
        print(f"Error checking/adding contract/semester fields: {e}")
    finally:
        db.close()

def migrate_db_add_hidden_pages():
    db = SessionLocal()
    try:
        inspector = inspect(engine)
        cols = [c['name'] for c in inspector.get_columns('users')]
        if 'hidden_pages' not in cols:
            db.execute(text("ALTER TABLE users ADD COLUMN hidden_pages VARCHAR DEFAULT '[]';"))
            db.commit()
            print("Successfully added hidden_pages column to 'users' table.")
    except Exception as e:
        print(f"Error checking/adding hidden_pages column: {e}")
    finally:
        db.close()

def migrate_db_init_activity_logs():
    db = SessionLocal()
    try:
        models.Base.metadata.create_all(bind=engine)
        count = db.query(models.ActivityLog).count()
        if count == 0:
            import re
            notifs = db.query(models.Notification).order_by(models.Notification.id.asc()).all()
            for n in notifs:
                uname = "نظام"
                urole = "نظام"
                if n.action_by:
                    m = re.match(r'(.+?)\s*\((.+?)\)', n.action_by)
                    if m:
                        uname = m.group(1).strip()
                        urole = m.group(2).strip()
                    else:
                        uname = n.action_by.strip()
                        urole = uname
                
                atype = "UPDATE"
                txt = n.action_text or ""
                if any(w in txt for w in ["حذف", "مسح"]):
                    atype = "DELETE"
                elif any(w in txt for w in ["إضافة", "اضافة", "جديد"]):
                    atype = "CREATE"
                elif any(w in txt for w in ["اعتماد"]):
                    atype = "APPROVE"
                elif any(w in txt for w in ["استرجاع"]):
                    atype = "RESTORE"
                elif any(w in txt for w in ["طباعة", "تصدير", "تنزيل"]):
                    atype = "EXPORT"
                
                etype = "SYSTEM"
                if "هيئة التدريس" in txt or "عضو" in txt:
                    etype = "PROFESSORS"
                elif "مقرر" in txt:
                    etype = "COURSES"
                elif "الخطة" in txt:
                    etype = "STUDY_PLAN"
                elif "مستخدم" in txt or "صلاحية" in txt:
                    etype = "USERS"
                elif "توقيع" in txt:
                    etype = "SIGNATURES"
                elif "المحذوف" in txt:
                    etype = "RECYCLE_BIN"
                
                log = models.ActivityLog(
                    username=uname,
                    user_role=urole,
                    action_type=atype,
                    entity_type=etype,
                    description=txt,
                    status="success",
                    faculty_id=n.faculty_id,
                    academic_year=n.academic_year,
                    semester=n.semester,
                    created_at=n.created_at or datetime.utcnow()
                )
                db.add(log)
            db.commit()
            print(f"Successfully migrated {len(notifs)} historical activity logs to activity_logs table.")
    except Exception as e:
        print(f"Error initializing activity_logs: {e}")
        db.rollback()
    finally:
        db.close()

@app.on_event("startup")
def startup_event():
    migrate_db_add_academic_years()
    migrate_db_add_signature_report_type()
    migrate_db_add_email()
    migrate_db_add_year()
    migrate_db_add_base_course_id()
    migrate_db_add_signature_order()
    migrate_db_add_user_permissions()
    migrate_db_add_entry_group_id()
    migrate_db_add_is_bundle()
    migrate_db_add_medicine_plan_fields()
    migrate_db_add_user_security_columns()
    migrate_db_add_contract_fields()
    migrate_db_add_hidden_pages()
    migrate_db_init_activity_logs()


    db = SessionLocal()
    # إنشاء مستخدم أدمن افتراضي
    if db.query(models.User).filter(models.User.username == "admin").first() is None:
        admin_user = models.User(
            username="admin",
            hashed_password=auth.get_password_hash("admin123"),
            role=models.UserRole.admin,
            faculty_id=None
        )
        db.add(admin_user)
        db.commit()

    if db.query(models.Faculty).count() == 0:
        data = {
            "كلية الهندسة": ["الأمن السيبراني", "الحوسبة السحابية", "الحوسبة عالية الكفاءة", "تخطيط وتشييد المدن الذكية", "هندسة الحاسوب", "هندسة المواد وإدارة التصنيع"],
            "كلية علوم الحاسوب والذكاء الاصطناعي": ["إنترنت الأشياء وتحليل البيانات الضخمة", "ذكاء الآلة", "علوم البيانات"],
            "كلية العلوم الإنسانية": ["اللغة الإنجليزية والترجمة التخصصية"],
            "كلية التمريض": ["علوم التمريض", "تمريض الطوارئ", "تمريض القبالة", "تمريض حديثي الولادة"],
            "كلية الصيدلة": ["فارم دي (الصيدلة الإكلينيكية)"],
            "كلية الطب البيطري": ["الطب البيطري"],
            "كلية طب الأسنان": ["طب وجراحة الفم والأسنان"],
            "كلية تكنولوجيا العلوم الصحية التطبيقية": ["برنامج عام", "تكنولوجيا الرعاية التنفسية", "تكنولوجيا المختبرات الطبية", "تكنولوجيا صناعة تركيبات الأسنان", "تكنولوجيا علوم الأشعة والتصوير الطبي"],
            "كلية العلاج الطبيعي": ["العلاج الطبيعي"],
            "كلية الطب والجراحة": ["الطب والجراحة اللائحة الجديدة", "الطب والجراحة اللائحة القديمة"]
        }
        for faculty_name, programs in data.items():
            faculty = models.Faculty(name=faculty_name)
            db.add(faculty)
            db.commit()
            db.refresh(faculty)
            for prog_name in programs:
                program = models.Program(name=prog_name, faculty_id=faculty.id)
                db.add(program)
        db.commit()
    db.close()

@app.get("/")
def read_root():
    return {"message": "Welcome to MNU Study Plan System Backend!"}


# ==========================================
# مسارات التسجيل والمستخدمين (Auth & Users APIs)
# ==========================================

@app.post("/api/auth/login", response_model=schemas.Token)
def login_for_access_token(request: Request, form_data: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    client_ip = request.client.host if request.client else "127.0.0.1"
    clean_username = form_data.username.strip() if form_data.username else ""
    clean_password = form_data.password.strip() if form_data.password else ""
    # Exact match first
    user = db.query(models.User).filter(models.User.username == clean_username).first()
    if not user:
        # Case-insensitive or underscore/multiple spaces normalized match
        normalized_input = re.sub(r'\s+', ' ', clean_username.replace('_', ' ')).strip().lower()
        all_users = db.query(models.User).all()
        for u in all_users:
            if u.username:
                u_norm = re.sub(r'\s+', ' ', u.username.replace('_', ' ')).strip().lower()
                if u_norm == normalized_input:
                    user = u
                    break
    if not user or not auth.verify_password(clean_password, user.hashed_password):
        log_activity(
            db=db,
            username=form_data.username,
            user_role=get_user_role_display(user) if user else "غير معروف",
            action_type="LOGIN_FAILED",
            entity_type="AUTH",
            description=f"محاولة تسجيل دخول غير ناجحة للمستخدم '{form_data.username}'",
            status="failed",
            ip_address=client_ip
        )
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="اسم المستخدم أو كلمة المرور غير صحيحة",
            headers={"WWW-Authenticate": "Bearer"},
        )
    access_token_expires = auth.timedelta(minutes=auth.ACCESS_TOKEN_EXPIRE_MINUTES)
    access_token = auth.create_access_token(
        data={"sub": user.username, "role": user.role, "faculty_id": user.faculty_id}, expires_delta=access_token_expires
    )
    
    role_disp = get_user_role_display(user)
    log_activity(
        db=db,
        username=user.username,
        user_id=user.id,
        user_role=role_disp,
        action_type="LOGIN",
        entity_type="AUTH",
        description=f"قام المستخدم '{user.username}' بتسجيل الدخول إلى النظام ({role_disp})",
        status="success",
        faculty_id=user.faculty_id,
        ip_address=client_ip
    )
    return {"access_token": access_token, "token_type": "bearer"}

@app.get("/api/users/me", response_model=schemas.UserOut)
def read_users_me(current_user: models.User = Depends(auth.get_current_user)):
    return current_user

@app.post("/api/auth/change-password")
def change_password(req: schemas.ChangePasswordRequest, request: Request, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    client_ip = request.client.host if request.client else "127.0.0.1"
    if not auth.verify_password(req.old_password, current_user.hashed_password):
        log_activity(
            db=db,
            username=current_user.username,
            user_id=current_user.id,
            user_role=get_user_role_display(current_user),
            action_type="SECURITY",
            entity_type="AUTH",
            description=f"فشل تغيير كلمة المرور للمستخدم '{current_user.username}' (كلمة المرور الحالية غير صحيحة)",
            status="failed",
            ip_address=client_ip
        )
        raise HTTPException(status_code=400, detail="كلمة المرور الحالية غير صحيحة")
    current_user.hashed_password = auth.get_password_hash(req.new_password)
    db.commit()
    log_activity(
        db=db,
        username=current_user.username,
        user_id=current_user.id,
        user_role=get_user_role_display(current_user),
        action_type="SECURITY",
        entity_type="AUTH",
        description=f"قام المستخدم '{current_user.username}' بتغيير كلمة المرور بنجاح",
        status="success",
        ip_address=client_ip
    )
    return {"message": "تم تغيير كلمة المرور بنجاح"}

@app.post("/api/auth/security-question")
def set_security_question(req: schemas.SetSecurityQuestionRequest, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    current_user.security_question = req.security_question
    current_user.security_answer = req.security_answer
    db.commit()
    return {"message": "تم ضبط سؤال الأمان بنجاح"}

@app.get("/api/auth/forgot-password/{username}")
def get_security_question(username: str, db: Session = Depends(get_db)):
    clean_name = username.strip()
    user = db.query(models.User).filter(func.lower(models.User.username) == clean_name.lower()).first()
    if not user:
        raise HTTPException(status_code=404, detail="المستخدم غير موجود. تأكد من كتابة اسم المستخدم بشكل صحيح.")
    if not user.security_question:
        raise HTTPException(status_code=400, detail="لم يتم إعداد سؤال أمان لهذا الحساب حتى الآن. يمكنك تسجيل الدخول بكلمة المرور الافتراضية (123456) ثم إعداد سؤال الأمان من الإعدادات أو صفحة الإرشادات.")
    
    import random
    q = (user.security_question or "").strip()
    fake_options = []
    
    if "حيوان" in q:
        fake_options = ["الأسد", "النمر", "الحصان", "الصقر", "الغزال", "الفهد", "الذئب"]
    elif "كلي" in q or "تخرج" in q:
        fake_options = ["كلية الطب", "كلية الهندسة", "كلية الحاسبات والمعلومات", "كلية العلوم", "كلية الصيدلة", "كلية التجارة"]
    elif "عام" in q or "سن" in q or "تاريخ" in q:
        fake_options = ["2018", "2019", "2020", "2021", "2022", "2023", "2024", "2025"]
    elif "مقرر" in q or "ماد" in q:
        fake_options = ["برمجة الحاسب", "الرياضيات", "الفيزياء", "الكيمياء", "نظم التشغيل", "قواعد البيانات"]
    elif "أستاذ" in q or "استاذ" in q or "دكتور" in q or "معلم" in q or "صديق" in q:
        fake_options = ["د. أحمد", "د. محمد", "د. فاطمة", "د. محمود", "د. سارة", "د. طارق", "د. إبراهيم"]
    elif "مدين" in q or "بلد" in q or "محافظ" in q or "مكان" in q:
        fake_options = ["القاهرة", "الإسكندرية", "شبين الكوم", "طنطا", "المنصورة", "الجيزة"]
    elif "لون" in q:
        fake_options = ["الأزرق", "الأخضر", "الأبيض", "الأسود", "الأحمر"]
    elif "أكل" in q or "طعام" in q or "وجب" in q or "فاكه" in q:
        fake_options = ["السمك", "المشويات", "البيتزا", "التفاح", "المانجو"]
    else:
        fake_options = ["الأسد", "كلية الهندسة", "2022", "الرياضيات", "د. محمد", "القاهرة", "الأزرق", "التفاح"]

    ans = user.security_answer.strip() if user.security_answer else ""
    filtered_fake = [opt for opt in fake_options if opt.strip().lower() != ans.lower()]
    
    random.shuffle(filtered_fake)
    options = filtered_fake[:3] + [ans]
    random.shuffle(options)
    
    return {"security_question": user.security_question, "options": options}

@app.post("/api/auth/reset-password")
def reset_password(req: schemas.ResetPasswordRequest, db: Session = Depends(get_db)):
    clean_name = req.username.strip()
    user = db.query(models.User).filter(func.lower(models.User.username) == clean_name.lower()).first()
    if not user:
        raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if not user.security_answer or user.security_answer.strip().lower() != req.security_answer.strip().lower():
        raise HTTPException(status_code=400, detail="إجابة سؤال الأمان غير صحيحة")
    user.hashed_password = auth.get_password_hash(req.new_password)
    db.commit()
    return {"message": "تم إعادة تعيين كلمة المرور بنجاح"}

@app.post("/api/users", response_model=schemas.UserOut)
def create_user(user: schemas.UserCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role not in [models.UserRole.admin, models.UserRole.faculty_professor]:
        raise HTTPException(status_code=403, detail="ليس لديك صلاحية لإضافة مسؤولين")
    if current_user.role == models.UserRole.faculty_professor:
        if user.role != models.UserRole.faculty_admin or user.faculty_id != current_user.faculty_id:
            raise HTTPException(status_code=403, detail="غير مصرح لك بإضافة إلا مسؤول كلية لكليتك فقط")
    db_user = db.query(models.User).filter(models.User.username == user.username).first()
    if db_user:
        raise HTTPException(status_code=400, detail="اسم المستخدم موجود بالفعل")
    hashed_password = auth.get_password_hash(user.password)
    new_user = models.User(
        username=user.username,
        hashed_password=hashed_password,
        role=user.role,
        job_title=user.job_title,
        faculty_id=user.faculty_id,
        all_faculties_access=user.all_faculties_access or (user.role == models.UserRole.admin and not user.assigned_faculty_ids)
    )
    if user.role == models.UserRole.reviewer:
        new_user.perm_review_2 = True
    
    if not new_user.all_faculties_access and user.assigned_faculty_ids:
        facs = db.query(models.Faculty).filter(models.Faculty.id.in_(user.assigned_faculty_ids)).all()
        new_user.assigned_faculties = facs
        if not new_user.faculty_id and facs:
            new_user.faculty_id = facs[0].id
    elif new_user.faculty_id and not user.assigned_faculty_ids and not new_user.all_faculties_access:
        fac = db.query(models.Faculty).filter(models.Faculty.id == new_user.faculty_id).first()
        if fac:
            new_user.assigned_faculties = [fac]
            
    db.add(new_user)
    db.commit()
    db.refresh(new_user)
    
    role_ar = user.job_title if user.job_title else ("مدير عام" if user.role == "admin" else ("مدير شؤون الطلاب" if user.role == "student_affairs" else ("مسؤول كلية" if user.role == "faculty_admin" else ("مدير برنامج" if user.role == "faculty_professor" else "المراجع"))))
    create_notification(db, None, f"{current_user.username.split('@')[0]} ({'مدير عام' if current_user.role == models.UserRole.admin else 'مدير برنامج'})", f"[ADMIN_ONLY] قام بإضافة مستخدم جديد: {user.username} ({role_ar})")
    db.commit()
    
    return new_user

@app.get("/api/users", response_model=List[schemas.UserOut])
def get_users(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role not in [models.UserRole.admin, models.UserRole.faculty_professor]:
        raise HTTPException(status_code=403, detail="ليس لديك صلاحية الوصول للمستخدمين")
    if current_user.role == models.UserRole.faculty_professor:
        if current_user.faculty_id:
            return db.query(models.User).filter(
                models.User.role == models.UserRole.faculty_admin,
                models.User.faculty_id == current_user.faculty_id
            ).all()
        return []
    return db.query(models.User).all()

@app.put("/api/users/{user_id}", response_model=schemas.UserOut)
def update_user(user_id: int, user_data: schemas.UserUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role not in [models.UserRole.admin, models.UserRole.faculty_professor]:
        raise HTTPException(status_code=403, detail="ليس لديك صلاحية لتحديث بيانات المسؤولين")

    db_user = db.query(models.User).filter(models.User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="المستخدم غير موجود")

    if current_user.role == models.UserRole.faculty_professor:
        if db_user.role != models.UserRole.faculty_admin or db_user.faculty_id != current_user.faculty_id:
            raise HTTPException(status_code=403, detail="غير مصرح لك بتعديل صلاحيات مسؤول كليات أخرى")
            
    if user_data.username is not None:
        existing = db.query(models.User).filter(models.User.username == user_data.username, models.User.id != user_id).first()
        if existing:
            raise HTTPException(status_code=400, detail="اسم المستخدم موجود بالفعل")
        db_user.username = user_data.username

        
    if user_data.password is not None and user_data.password.strip() != "":
        db_user.hashed_password = auth.get_password_hash(user_data.password)
        
    if user_data.role is not None:
        db_user.role = user_data.role
        if user_data.role == models.UserRole.admin:
            db_user.faculty_id = None
            
    if user_data.job_title is not None:
        db_user.job_title = user_data.job_title

    if user_data.all_faculties_access is not None:
        db_user.all_faculties_access = user_data.all_faculties_access

    if user_data.assigned_faculty_ids is not None:
        if db_user.all_faculties_access:
            db_user.assigned_faculties = []
        else:
            facs = db.query(models.Faculty).filter(models.Faculty.id.in_(user_data.assigned_faculty_ids)).all()
            db_user.assigned_faculties = facs
            if facs:
                db_user.faculty_id = facs[0].id
    elif user_data.faculty_id is not None and not db_user.all_faculties_access:
        db_user.faculty_id = user_data.faculty_id
        fac = db.query(models.Faculty).filter(models.Faculty.id == user_data.faculty_id).first()
        if fac:
            db_user.assigned_faculties = [fac]
        
    permission_changes = []
    
    if user_data.perm_view_prof_study_plan is not None and db_user.perm_view_prof_study_plan != user_data.perm_view_prof_study_plan:
        db_user.perm_view_prof_study_plan = user_data.perm_view_prof_study_plan
        status = "تفعيل" if user_data.perm_view_prof_study_plan else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'جدول الأساتذة (الخطة الدراسية)'")
        
    if user_data.perm_view_prof_data_btn is not None and db_user.perm_view_prof_data_btn != user_data.perm_view_prof_data_btn:
        db_user.perm_view_prof_data_btn = user_data.perm_view_prof_data_btn
        status = "تفعيل" if user_data.perm_view_prof_data_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر بيانات التدريس (الرئيسية)'")
        
    if user_data.perm_view_prof_delete_btn is not None and db_user.perm_view_prof_delete_btn != user_data.perm_view_prof_delete_btn:
        db_user.perm_view_prof_delete_btn = user_data.perm_view_prof_delete_btn
        status = "تفعيل" if user_data.perm_view_prof_delete_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر الحذف (هيئة التدريس)'")
        
    if user_data.perm_view_prof_view_btn is not None and db_user.perm_view_prof_view_btn != user_data.perm_view_prof_view_btn:
        db_user.perm_view_prof_view_btn = user_data.perm_view_prof_view_btn
        status = "تفعيل" if user_data.perm_view_prof_view_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر الرؤية (هيئة التدريس)'")
        
    if user_data.perm_view_prof_print_btn is not None and db_user.perm_view_prof_print_btn != user_data.perm_view_prof_print_btn:
        db_user.perm_view_prof_print_btn = user_data.perm_view_prof_print_btn
        status = "تفعيل" if user_data.perm_view_prof_print_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر الطباعة (هيئة التدريس)'")
        
    if user_data.perm_view_prof_import_btn is not None and db_user.perm_view_prof_import_btn != user_data.perm_view_prof_import_btn:
        db_user.perm_view_prof_import_btn = user_data.perm_view_prof_import_btn
        status = "تفعيل" if user_data.perm_view_prof_import_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر الاستيراد (هيئة التدريس)'")
        
    if user_data.perm_delete_notif_btn is not None and db_user.perm_delete_notif_btn != user_data.perm_delete_notif_btn:
        db_user.perm_delete_notif_btn = user_data.perm_delete_notif_btn
        status = "تفعيل" if user_data.perm_delete_notif_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر حذف الإشعارات'")
        
    if user_data.perm_recycle_restore_btn is not None and db_user.perm_recycle_restore_btn != user_data.perm_recycle_restore_btn:
        db_user.perm_recycle_restore_btn = user_data.perm_recycle_restore_btn
        status = "تفعيل" if user_data.perm_recycle_restore_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر استرجاع المحذوف'")
        
    if user_data.perm_recycle_delete_btn is not None and db_user.perm_recycle_delete_btn != user_data.perm_recycle_delete_btn:
        db_user.perm_recycle_delete_btn = user_data.perm_recycle_delete_btn
        status = "تفعيل" if user_data.perm_recycle_delete_btn else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'زر الحذف النهائي (المحذوفات)'")
        
    if user_data.perm_review_1 is not None and db_user.perm_review_1 != user_data.perm_review_1:
        db_user.perm_review_1 = user_data.perm_review_1
        status = "تفعيل" if user_data.perm_review_1 else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'مراجعة أولى'")
        
    if user_data.perm_review_2 is not None and db_user.perm_review_2 != user_data.perm_review_2:
        db_user.perm_review_2 = user_data.perm_review_2
        status = "تفعيل" if user_data.perm_review_2 else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'مراجعة ثانية'")
        
    if user_data.perm_approve_plan is not None and db_user.perm_approve_plan != user_data.perm_approve_plan:
        db_user.perm_approve_plan = user_data.perm_approve_plan
        status = "تفعيل" if user_data.perm_approve_plan else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'إعتماد الخطة'")
        
    if user_data.perm_finish_plan is not None and db_user.perm_finish_plan != user_data.perm_finish_plan:
        db_user.perm_finish_plan = user_data.perm_finish_plan
        status = "تفعيل" if user_data.perm_finish_plan else "إلغاء"
        permission_changes.append(f"{status} صلاحية 'إنهاء الخطة'")
        
    if permission_changes and db_user.faculty_id:
        short_target = db_user.username.split('@')[0]
        short_actor = current_user.username.split('@')[0]
        actor_role_ar = "مدير عام" if current_user.role == models.UserRole.admin else "مدير برنامج"
        changes_str = " و ".join(permission_changes)
        action_text = f"قام ب{changes_str} لمسؤول الكلية ({short_target})"
        create_notification(db, db_user.faculty_id, f"{short_actor} ({actor_role_ar})", action_text)
        
    if user_data.hidden_pages is not None:
        db_user.hidden_pages = user_data.hidden_pages

    db.commit()
    db.refresh(db_user)
    return db_user

@app.put("/api/users/{user_id}/hidden-pages", response_model=schemas.UserOut)
def update_user_hidden_pages(user_id: int, payload: schemas.UserHiddenPagesUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role not in [models.UserRole.admin, models.UserRole.faculty_professor]:
        raise HTTPException(status_code=403, detail="ليس لديك صلاحية لتحديث بيانات المسؤولين")
    db_user = db.query(models.User).filter(models.User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    db_user.hidden_pages = json.dumps(payload.hidden_pages, ensure_ascii=False)
    db.commit()
    db.refresh(db_user)
    return db_user

@app.delete("/api/users/{user_id}")

def delete_user(user_id: int, db: Session = Depends(get_db), admin: models.User = Depends(auth.require_admin)):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if user.username == "admin":
        raise HTTPException(status_code=400, detail="لا يمكن حذف مدير النظام الافتراضي")
        
    try:
        # Save user details for notification before deleting
        deleted_username = user.username
        deleted_role = user.role
        role_ar = "مدير عام" if deleted_role == "admin" else ("مدير شؤون الطلاب" if deleted_role == "student_affairs" else ("مسؤول كلية" if deleted_role == "faculty_admin" else ("مدير برنامج" if deleted_role == "faculty_professor" else "المراجع")))
        
        # Remove associations and detach logs to avoid foreign key constraints
        user.assigned_faculties.clear()
        db.query(models.ActivityLog).filter(models.ActivityLog.user_id == user_id).update({"user_id": None}, synchronize_session=False)
        
        db.delete(user)
        create_notification(db, None, f"{admin.username.split('@')[0]} ({'مدير عام' if admin.role == models.UserRole.admin else 'مدير برنامج'})", f"[ADMIN_ONLY] قام بحذف المستخدم: {deleted_username} ({role_ar})")
        db.commit()
    except Exception as e:
        db.rollback()
        import traceback
        error_msg = str(e)
        raise HTTPException(status_code=400, detail=f"Database Error: {error_msg}")
    
    return {"message": "تم الحذف بنجاح"}

# ==========================================
# مسارات سجل العمليات والأنشطة (Activity & Audit Logs APIs)
# ==========================================

@app.get("/api/logs")
def get_activity_logs(
    page: int = 1,
    limit: int = 50,
    search: Optional[str] = None,
    action_type: Optional[str] = None,
    entity_type: Optional[str] = None,
    faculty_id: Optional[int] = None,
    status: Optional[str] = None,
    username: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    query = db.query(models.ActivityLog)
    
    # Filter by user faculty access if not admin
    if current_user.role != models.UserRole.admin and not current_user.all_faculties_access:
        if current_user.faculty_id:
            query = query.filter(or_(models.ActivityLog.faculty_id == current_user.faculty_id, models.ActivityLog.faculty_id == None))
    
    if faculty_id:
        query = query.filter(models.ActivityLog.faculty_id == faculty_id)
    if action_type:
        query = query.filter(models.ActivityLog.action_type == action_type)
    if entity_type:
        query = query.filter(models.ActivityLog.entity_type == entity_type)
    if status:
        query = query.filter(models.ActivityLog.status == status)
    if username:
        query = query.filter(models.ActivityLog.username.ilike(f"%{username}%"))
    if search:
        search_pattern = f"%{search}%"
        query = query.filter(
            or_(
                models.ActivityLog.description.ilike(search_pattern),
                models.ActivityLog.username.ilike(search_pattern),
                models.ActivityLog.user_role.ilike(search_pattern),
                models.ActivityLog.ip_address.ilike(search_pattern)
            )
        )
    if from_date:
        try:
            fd = datetime.strptime(from_date, "%Y-%m-%d")
            query = query.filter(models.ActivityLog.created_at >= fd)
        except Exception:
            pass
    if to_date:
        try:
            td = datetime.strptime(to_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
            query = query.filter(models.ActivityLog.created_at <= td)
        except Exception:
            pass

    total = query.count()
    query = query.order_by(models.ActivityLog.id.desc())
    
    if limit > 0:
        offset = (page - 1) * limit
        items = query.offset(offset).limit(limit).all()
        total_pages = (total + limit - 1) // limit
    else:
        items = query.all()
        total_pages = 1

    result_items = []
    for item in items:
        fac_name = item.faculty.name if item.faculty else None
        result_items.append({
            "id": item.id,
            "user_id": item.user_id,
            "username": item.username,
            "user_role": item.user_role,
            "action_type": item.action_type,
            "entity_type": item.entity_type,
            "description": item.description,
            "status": item.status,
            "faculty_id": item.faculty_id,
            "faculty_name": fac_name,
            "academic_year": item.academic_year,
            "semester": item.semester,
            "ip_address": item.ip_address,
            "created_at": item.created_at.isoformat() if item.created_at else None
        })

    return {
        "items": result_items,
        "total": total,
        "page": page,
        "limit": limit,
        "total_pages": total_pages
    }

@app.get("/api/logs/stats")
def get_activity_log_stats(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    total = db.query(models.ActivityLog).count()
    today_start = datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0)
    today_count = db.query(models.ActivityLog).filter(models.ActivityLog.created_at >= today_start).count()
    login_count = db.query(models.ActivityLog).filter(models.ActivityLog.action_type == "LOGIN").count()
    failed_login_count = db.query(models.ActivityLog).filter(models.ActivityLog.action_type == "LOGIN_FAILED").count()
    create_count = db.query(models.ActivityLog).filter(models.ActivityLog.action_type == "CREATE").count()
    update_count = db.query(models.ActivityLog).filter(models.ActivityLog.action_type == "UPDATE").count()
    delete_count = db.query(models.ActivityLog).filter(models.ActivityLog.action_type == "DELETE").count()

    return {
        "total": total,
        "today": today_count,
        "logins": login_count,
        "failed_logins": failed_login_count,
        "creates": create_count,
        "updates": update_count,
        "deletes": delete_count
    }

@app.delete("/api/logs/clear-all")
def clear_all_activity_logs(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="هذا الإجراء مخصص للمدير العام فقط")
    db.query(models.ActivityLog).delete()
    db.commit()
    return {"message": "تم مسح جميع السجلات بنجاح"}

@app.post("/api/logs/bulk-delete")
def bulk_delete_activity_logs(
    req: schemas.BulkDeleteLogsRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="هذا الإجراء مخصص للمدير العام فقط")
    db.query(models.ActivityLog).filter(models.ActivityLog.id.in_(req.ids)).delete(synchronize_session=False)
    db.commit()
    return {"message": f"تم حذف {len(req.ids)} سجل بنجاح"}

@app.delete("/api/logs/{log_id}")
def delete_activity_log(
    log_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="هذا الإجراء مخصص للمدير العام فقط")
    log = db.query(models.ActivityLog).filter(models.ActivityLog.id == log_id).first()
    if not log:
        raise HTTPException(status_code=404, detail="السجل غير موجود")
    db.delete(log)
    db.commit()
    return {"message": "تم حذف السجل بنجاح"}

# ==========================================
# مسارات الكليات (Faculties APIs)
# ==========================================

@app.post("/api/faculties", response_model=schemas.FacultyOut)
def create_faculty(faculty: schemas.FacultyCreate, db: Session = Depends(get_db)):
    db_faculty = db.query(models.Faculty).filter(models.Faculty.name == faculty.name).first()
    if db_faculty:
        raise HTTPException(status_code=400, detail="هذه الكلية مسجلة بالفعل")
    new_faculty = models.Faculty(name=faculty.name)
    db.add(new_faculty)
    db.commit()
    db.refresh(new_faculty)
    return new_faculty

@app.get("/api/faculties", response_model=list[schemas.FacultyOut])
def get_faculties(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role == models.UserRole.admin or current_user.all_faculties_access:
        return db.query(models.Faculty).all()
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        return db.query(models.Faculty).filter(models.Faculty.id.in_(faculty_ids)).all()
    elif current_user.faculty_id:
        return db.query(models.Faculty).filter(models.Faculty.id == current_user.faculty_id).all()
    elif current_user.role == models.UserRole.student_affairs:
        return db.query(models.Faculty).all()
    else:
        return []

# ==========================================
# مسارات البرامج الدراسية (Programs APIs)
# ==========================================

@app.post("/api/programs", response_model=schemas.ProgramOut)
def create_program(program: schemas.ProgramCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    db_faculty = db.query(models.Faculty).filter(models.Faculty.id == program.faculty_id).first()
    if not db_faculty:
        raise HTTPException(status_code=404, detail="الكلية غير موجودة")
    if current_user.role in [models.UserRole.faculty_admin, models.UserRole.faculty_professor] and current_user.faculty_id != program.faculty_id:
        raise HTTPException(status_code=403, detail="لا يمكنك إضافة برنامج في كلية أخرى")
    new_program = models.Program(name=program.name, faculty_id=program.faculty_id)
    db.add(new_program)
    db.commit()
    db.refresh(new_program)
    return new_program

@app.get("/api/programs", response_model=list[schemas.ProgramOut])
def get_programs(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    query = db.query(models.Program)
    if current_user.role in [models.UserRole.faculty_admin, models.UserRole.faculty_professor]:
        query = query.filter(models.Program.faculty_id == current_user.faculty_id)
    return query.all()

# ==========================================
# مسارات أعضاء هيئة التدريس (Professors APIs)
# ==========================================
@app.get("/api/professors/template")
def download_excel_template(db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.utils import get_column_letter
    from tempfile import NamedTemporaryFile
    from fastapi.responses import FileResponse
    
    def clean_range_name(name):
        forbidden = [" ", "/", "-", "(", ")", "[", "]", "{", "}", "&", "*", "+", ",", ".", "\\", "?", "!", "=", "|", ":", ";", "'", '"', "<", ">", "@", "#", "$", "%", "^"]
        for char in forbidden:
            name = name.replace(char, "_")
        name = name.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
        if name and name[0].isdigit():
            name = "_" + name
        return name

    # Create workbook
    wb = openpyxl.Workbook()
    
    # Setup template sheet
    ws = wb.active
    ws.title = "البيانات"
    ws.views.sheetView[0].showGridLines = True
    ws.sheet_view.rightToLeft = True
    
    # Headers
    headers = [
        "اسم عضو هيئة التدريس",
        "الرقم القومي",
        "الدرجة العلمية",
        "جهة القدوم",
        "رقم الهاتف",
        "البريد الإلكتروني",
        "الكلية التابع لها",
        "نوع التعاقد",
        "طبيعة العمل بجامعة المنوفية الأهلية",
        "حضور الفصل الأول",
        "حضور الفصل الثاني",
        "حضور الفصل الصيفي"
    ]
    ws.append(headers)
    ws.freeze_panes = "A2"
    
    # Style definitions
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    header_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=12)
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))
                         
    # Style headers and set column widths
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 25
    
    # Specific column widths
    ws.column_dimensions['A'].width = 35 # Name
    ws.column_dimensions['B'].width = 28 # National ID
    ws.column_dimensions['C'].width = 22 # Job title
    ws.column_dimensions['D'].width = 25 # Workplace
    ws.column_dimensions['E'].width = 20 # Phone
    ws.column_dimensions['F'].width = 28 # Email
    ws.column_dimensions['G'].width = 30 # Faculty
    ws.column_dimensions['H'].width = 28 # Contract type
    ws.column_dimensions['I'].width = 35 # MNU job title
    ws.column_dimensions['J'].width = 22 # Sem 1
    ws.column_dimensions['K'].width = 22 # Sem 2
    ws.column_dimensions['L'].width = 22 # Summer
    
    # Freeze the first row
    ws.freeze_panes = 'A2'

    # Format data rows
    for row_idx in range(2, 201):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Format الرقم القومي column as whole number
        ws.cell(row=row_idx, column=2).number_format = '0'
        # Set text formatting for Phone just to be safe
        ws.cell(row=row_idx, column=5).number_format = '@'

    # Create helper sheet for validation lists
    ws_lists = wb.create_sheet(title="قوائم المرجعية")
    ws_lists.sheet_state = "hidden"
    ws_lists.sheet_view.rightToLeft = True
    
    # 1. Populate faculties in Column A of ws_lists
    faculties = db.query(models.Faculty).all()
    fac_names = [f.name for f in faculties]
    for idx, name in enumerate(fac_names, 1):
        ws_lists.cell(row=idx, column=1, value=name)
        
    # 2. Write programs of each faculty in subsequent columns of ws_lists
    # and create the Named Range for each faculty
    for fac_idx, fac in enumerate(faculties, 2): # Start from column 2 (B)
        progs = db.query(models.Program).filter(models.Program.faculty_id == fac.id).all()
        prog_names = [p.name for p in progs]
        if not prog_names:
            prog_names = ["لا يوجد برامج حاليا"]
            
        for row_idx, prog_name in enumerate(prog_names, 1):
            ws_lists.cell(row=row_idx, column=fac_idx, value=prog_name)
            
        col_letter = get_column_letter(fac_idx)
        range_name = clean_range_name(fac.name)
        ref_str = f"'قوائم المرجعية'!${col_letter}$1:${col_letter}${len(prog_names)}"
        
        new_range = DefinedName(name=range_name)
        new_range.value = ref_str
        wb.defined_names.add(new_range)
        
    # 3. Populate levels of each program in columns starting after programs columns
    program_levels_map = {
        "المستوى العام": ["المستوى 0", "المستوى 1"],
        "الأمن السيبراني": ["المستوى 2", "المستوى 3", "المستوى 4"],
        "الحوسبة السحابية": ["المستوى 2", "المستوى 3", "المستوى 4"],
        "الحوسبة عالية الكفاءة": ["المستوى 2", "المستوى 3", "المستوى 4"],
        "تخطيط وتشييد المدن الذكية": ["المستوى 0", "المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "هندسة الحاسوب": ["المستوى 0", "المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "هندسة المواد وإدارة التصنيع": ["المستوى 0", "المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "إنترنت الأشياء وتحليل البيانات الضخمة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "ذكاء الآلة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "علوم البيانات": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "اللغة الإنجليزية والترجمة التخصصية": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "علوم التمريض": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تمريض الطوارئ": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تمريض القبالة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تمريض حديثي الولادة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "فارم دي (الصيدلة الإكلينيكية)": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "الطب البيطري": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "طب وجراحة الفم والأسنان": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "برنامج عام": ["المستوى 1", "المستوى 2"],
        "تكنولوجيا الرعاية التنفسية": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تكنولوجيا المختبرات الطبية": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تكنولوجيا صناعة تركيبات الأسنان": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تكنولوجيا علوم الأشعة والتصوير الطبي": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "العلاج الطبيعي": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "الطب والجراحة اللائحة الجديدة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "الطب والجراحة اللائحة القديمة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"]
    }
    
    def get_levels_for_program(prog_name):
        def normalize_text(text):
            if not text:
                return ""
            return text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه").replace(" ", "").strip()
        
        norm_name = normalize_text(prog_name)
        for k, v in program_levels_map.items():
            if normalize_text(k) in norm_name or norm_name in normalize_text(k):
                return v
        return ["المستوى الأول", "المستوى الثاني", "المستوى الثالث", "المستوى الرابع"]

    # Start writing program levels at column: 2 + len(faculties)
    start_level_col = 2 + len(faculties)
    all_db_programs = db.query(models.Program).all()
    
    for prog_idx, prog in enumerate(all_db_programs):
        col_idx = start_level_col + prog_idx
        col_letter = get_column_letter(col_idx)
        
        levels_list = get_levels_for_program(prog.name)
        for r_idx, lvl in enumerate(levels_list, 1):
            ws_lists.cell(row=r_idx, column=col_idx, value=lvl)
            
        range_name = clean_range_name(prog.name)
        ref_str = f"'قوائم المرجعية'!${col_letter}$1:${col_letter}${len(levels_list)}"
        
        new_range = DefinedName(name=range_name)
        new_range.value = ref_str
        wb.defined_names.add(new_range)

    # 3.5. Populate courses of each program in columns of ws_lists
    start_course_col = 2 + len(faculties) + len(all_db_programs)
    for prog_idx, prog in enumerate(all_db_programs):
        col_idx = start_course_col + prog_idx
        col_letter = get_column_letter(col_idx)
        prog_courses = db.query(models.Course).filter(models.Course.program_id == prog.id).all()
        course_names = [c.name_ar for c in prog_courses]
        if not course_names:
            course_names = ["لا يوجد مقررات حاليا"]
            
        for r_idx, c_name in enumerate(course_names, 1):
            ws_lists.cell(row=r_idx, column=col_idx, value=c_name)
            
        range_name = clean_range_name(prog.name) + "_courses"
        ref_str = f"'قوائم المرجعية'!${col_letter}$1:${col_letter}${len(course_names)}"
        
        new_range = DefinedName(name=range_name)
        new_range.value = ref_str
        wb.defined_names.add(new_range)
        
    # 4. Populate jobs in Column 150 of ws_lists
    job_titles = ["معيد", "مدرس مساعد", "مدرس", "أستاذ مساعد", "أستاذ", "أخصائي"]
    for idx, name in enumerate(job_titles, 1):
        ws_lists.cell(row=idx, column=150, value=name)

    # 5. Populate contract options in Column 151 of ws_lists
    contract_options = [
        "تعاقد كلي (5 أيام)",
        "تعاقد جزئي (يوم واحد)",
        "تعاقد جزئي (يومان)",
        "تعاقد جزئي (3 أيام)",
        "تعاقد بالساعة",
        "بدون تعاقد"
    ]
    for idx, name in enumerate(contract_options, 1):
        ws_lists.cell(row=idx, column=151, value=name)

    # 6. Populate Semester 1 & 2 weeks (up to 15 weeks) in Column 152 of ws_lists
    weeks_15_options = [
        "أسبوع واحد",
        "أسبوعان",
        "3 أسابيع",
        "4 أسابيع",
        "5 أسابيع",
        "6 أسابيع",
        "7 أسابيع",
        "8 أسابيع",
        "9 أسابيع",
        "10 أسابيع",
        "11 أسبوع",
        "12 أسبوع",
        "13 أسبوع",
        "14 أسبوع",
        "15 أسبوع"
    ]
    for idx, name in enumerate(weeks_15_options, 1):
        ws_lists.cell(row=idx, column=152, value=name)

    # 7. Populate Summer Semester weeks (up to 8 weeks) in Column 153 of ws_lists
    weeks_8_options = [
        "أسبوع واحد",
        "أسبوعان",
        "3 أسابيع",
        "4 أسابيع",
        "5 أسابيع",
        "6 أسابيع",
        "7 أسابيع",
        "8 أسابيع"
    ]
    for idx, name in enumerate(weeks_8_options, 1):
        ws_lists.cell(row=idx, column=153, value=name)
        
    # Add Data Validations
    # 1. Faculty Dropdown for column G (الكلية التابع لها) - rows 2 to 200
    if fac_names:
        fac_formula = f"='قوائم المرجعية'!$A$1:$A${len(fac_names)}"
        dv_fac = DataValidation(type="list", formula1=fac_formula, allow_blank=True)
        dv_fac.error ='يرجى اختيار كلية من القائمة المتاحة فقط'
        dv_fac.errorTitle = 'اختيار غير صحيح'
        dv_fac.prompt = 'يرجى اختيار الكلية'
        dv_fac.promptTitle = 'الكلية التابع لها'
        ws.add_data_validation(dv_fac)
        dv_fac.add("G2:G200")
        
    # 2. Job Title Dropdown for column C (الدرجة العلمية) - rows 2 to 200
    job_col_letter = get_column_letter(150)
    job_formula = f"='قوائم المرجعية'!${job_col_letter}$1:${job_col_letter}${len(job_titles)}"
    dv_job = DataValidation(type="list", formula1=job_formula, allow_blank=True)
    dv_job.error ='يرجى اختيار وظيفة من القائمة المتاحة فقط'
    dv_job.errorTitle = 'اختيار غير صحيح'
    dv_job.prompt = 'يرجى اختيار الدرجة العلمية'
    dv_job.promptTitle = 'الدرجة العلمية'
    ws.add_data_validation(dv_job)
    dv_job.add("C2:C200")

    # 3. Contract Type Dropdown for column H (نوع التعاقد) - rows 2 to 200
    contract_col_letter = get_column_letter(151)
    contract_formula = f"='قوائم المرجعية'!${contract_col_letter}$1:${contract_col_letter}${len(contract_options)}"
    dv_contract = DataValidation(type="list", formula1=contract_formula, allow_blank=False)
    dv_contract.error = 'يرجى اختيار نوع التعاقد من القائمة المتاحة فقط'
    dv_contract.errorTitle = 'اختيار غير صحيح'
    dv_contract.prompt = 'يرجى اختيار نوع التعاقد (إجباري)'
    dv_contract.promptTitle = 'نوع التعاقد'
    ws.add_data_validation(dv_contract)
    dv_contract.add("H2:H200")

    # 4. National ID validation for column B (الرقم القومي)
    dv_nid = DataValidation(type="whole", operator="between", formula1="20000000000000", formula2="39999999999999", allow_blank=True)
    dv_nid.error = 'يجب إدخال رقم قومي صحيح مكون من 14 رقماً يبدأ بـ 2 أو 3 وبدون فواصل'
    dv_nid.errorTitle = 'إدخال غير صحيح'
    dv_nid.prompt = 'يرجى كتابة الرقم القومي (14 رقماً)'
    dv_nid.promptTitle = 'الرقم القومي'
    ws.add_data_validation(dv_nid)
    dv_nid.add("B2:B200")
    
    # 5. Phone number formatting for column E (رقم الهاتف)
    dv_phone = DataValidation(type="textLength", operator="equal", formula1="11", allow_blank=True)
    dv_phone.error = 'رقم الهاتف يجب أن يتكون من 11 رقماً'
    dv_phone.errorTitle = 'إدخال غير صحيح'
    dv_phone.prompt = 'يرجى كتابة رقم الهاتف (11 رقماً)'
    dv_phone.promptTitle = 'رقم الهاتف'
    ws.add_data_validation(dv_phone)
    dv_phone.add("E2:E200")

    # 6. Semester 1 attendance weeks for column J (حضور الفصل الأول) - up to 15 weeks
    sem1_col_letter = get_column_letter(152)
    sem1_formula = f"='قوائم المرجعية'!${sem1_col_letter}$1:${sem1_col_letter}${len(weeks_15_options)}"
    dv_sem1 = DataValidation(type="list", formula1=sem1_formula, allow_blank=True)
    dv_sem1.error = 'يرجى اختيار عدد أسابيع الحضور من القائمة'
    dv_sem1.errorTitle = 'اختيار غير صحيح'
    dv_sem1.prompt = 'يرجى اختيار عدد أسابيع حضور الفصل الأول'
    dv_sem1.promptTitle = 'حضور الفصل الأول'
    ws.add_data_validation(dv_sem1)
    dv_sem1.add("J2:J200")

    # 7. Semester 2 attendance weeks for column K (حضور الفصل الثاني) - up to 15 weeks
    sem2_col_letter = get_column_letter(152)
    sem2_formula = f"='قوائم المرجعية'!${sem2_col_letter}$1:${sem2_col_letter}${len(weeks_15_options)}"
    dv_sem2 = DataValidation(type="list", formula1=sem2_formula, allow_blank=True)
    dv_sem2.error = 'يرجى اختيار عدد أسابيع الحضور من القائمة'
    dv_sem2.errorTitle = 'اختيار غير صحيح'
    dv_sem2.prompt = 'يرجى اختيار عدد أسابيع حضور الفصل الثاني'
    dv_sem2.promptTitle = 'حضور الفصل الثاني'
    ws.add_data_validation(dv_sem2)
    dv_sem2.add("K2:K200")

    # 8. Summer semester attendance weeks for column L (حضور الفصل الصيفي) - up to 8 weeks
    summer_col_letter = get_column_letter(153)
    summer_formula = f"='قوائم المرجعية'!${summer_col_letter}$1:${summer_col_letter}${len(weeks_8_options)}"
    dv_summer = DataValidation(type="list", formula1=summer_formula, allow_blank=True)
    dv_summer.error = 'يرجى اختيار عدد أسابيع الحضور من القائمة'
    dv_summer.errorTitle = 'اختيار غير صحيح'
    dv_summer.prompt = 'يرجى اختيار عدد أسابيع حضور الفصل الصيفي'
    dv_summer.promptTitle = 'حضور الفصل الصيفي'
    ws.add_data_validation(dv_summer)
    dv_summer.add("L2:L200")
    
    # Save to temp file
    temp_file = NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    wb.close()
    
    return FileResponse(
        temp_file.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="نموذج_استيراد_اعضاء_هيئة_التدريس.xlsx"
    )

def normalize_program_name(name: str) -> str:
    if not name:
        return ""
    name = " ".join(str(name).split()).strip()
    if name.startswith("برنامج "):
        name = name[len("برنامج "):].strip()
    name = name.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    name = name.replace("ى", "ي").replace("ة", "ه")
    
    # Check common spelling variations
    mappings = {
        "فارم دي (الصيدله الاكلينيكيه)": "فارم دي (الصيدلة الإكلينيكية)",
        "فارم دي (الصيدله الإكلينيكيه)": "فارم دي (الصيدلة الإكلينيكية)",
        "فارم دي (الصيدلة الاكلينيكية)": "فارم دي (الصيدلة الإكلينيكية)",
        "الطب البيطرى": "الطب البيطري",
        "العلاج الطبيعى": "العلاج الطبيعي",
        "اللغه الانجليزيه والترجمه التخصصيه": "اللغة الإنجليزية والترجمة التخصصية",
        "اللغة الانجليزية والترجمة التخصصية": "اللغة الإنجليزية والترجمة التخصصية",
        "هندسه الحاسوب": "هندسة الحاسوب",
        "هندسه المواد و اداره التصنيع": "هندسة المواد وإدارة التصنيع",
        "تكنولوجيا علوم الاشعه والتصوير الطبي": "تكنولوجيا علوم الأشعة والتصوير الطبي",
        "تكنولوجيا علوم الاشعه والتصوير الطبى": "تكنولوجيا علوم الأشعة والتصوير الطبي",
        "تكنولوجيا علوم الأشعة والتصوير الطبى": "تكنولوجيا علوم الأشعة والتصوير الطبي",
        "طب وجراحه الفم والاسنان": "طب وجراحة الفم والأسنان",
        "طب وجراحة الفم والاسنان": "طب وجراحة الفم والأسنان",
        "انترنت الاشياء وتحليل البيانات الضخمه": "إنترنت الأشياء وتحليل البيانات الضخمة",
        "انترنت الأشياء وتحليل البيانات الضخمة": "إنترنت الأشياء وتحليل البيانات الضخمة",
        "انترنت الاشياء وتحليل البيانات الضخمة": "إنترنت الأشياء وتحليل البيانات الضخمة",
        "ذكاء الاله": "ذكاء الآلة",
        "ذكاء الآلة": "ذكاء الآلة",
        "علوم البيانات": "علوم البيانات",
    }
    
    clean_k = lambda s: " ".join(s.split()).strip().replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
    cleaned_name = clean_k(name)
    for k, v in mappings.items():
        if clean_k(k) == cleaned_name:
            return v
    return name

def normalize_semester_name(val):
    if not val:
        return "الأول"
    val = str(val).strip()
    val_norm = val.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
    if "اول" in val_norm or "ترم 1" in val_norm or "الاول" in val_norm:
        return "الأول"
    elif "ثاني" in val_norm or "ثانى" in val_norm or "ترم 2" in val_norm or "الثاني" in val_norm or "الثانى" in val_norm:
        return "الثانى"
    elif "صيف" in val_norm or "ترم 3" in val_norm or "الصيفي" in val_norm or "الصيفى" in val_norm:
        return "الصيفي"
    return val


@app.post("/api/professors/import")
async def import_professors_excel(
    file1: UploadFile = File(...),
    file2: Optional[UploadFile] = File(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    # Helper to clean Arabic characters for column match
    def clean_text(text):
        if not isinstance(text, str):
            return ""
        text = " ".join(text.split()).strip()
        text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
        text = text.replace("ى", "ي").replace("ة", "ه")
        return text

    # Helper to clean and format national id
    def clean_national_id(val):
        if pd.isna(val) or val is None:
            return ""
        val_str = str(val).strip()
        if not val_str or val_str.lower() in ["nan", "none", "n/a", "null", "-"]:
            return ""
        if "e" in val_str.lower():
            try:
                val_str = f"{int(float(val_str))}"
            except:
                pass
        elif "." in val_str:
            try:
                val_str = f"{int(float(val_str))}"
            except:
                pass
        val_str = "".join([c for c in val_str if c.isdigit()])
        return val_str

    # Helper function to extract academic title
    def parse_job_title_and_clean_name(name):
        name = name.strip()
        prefixes = [
            ("أ.د/", "أ.د"), ("أ.د.", "أ.د"), ("أ.د", "أ.د"),
            ("أ.م.د/", "أ.م"), ("أ.م.د.", "أ.م"), ("أ.م.د", "أ.م"), ("أ.م/", "أ.م"), ("أ.م.", "أ.م"), ("أ.م", "أ.م"),
            ("م.م/", "م.م"), ("م.م.", "م.م"), ("م.م", "م.م"), ("مدرس مساعد", "م.م"),
            ("م.ع/", "م.ع"), ("م.ع.", "م.ع"), ("م.ع", "م.ع"), ("معيد", "م.ع"),
            ("د/", "د"), ("د.", "د"), ("دكتور", "د"),
            ("أ/", "أ"), ("أستاذ", "أ")
        ]
        extracted_title = "د" # Default title
        clean_name = name
        for prefix, title in prefixes:
            if name.startswith(prefix):
                extracted_title = title
                clean_name = name[len(prefix):].strip()
                if clean_name.startswith("/") or clean_name.startswith("."):
                    clean_name = clean_name[1:].strip()
                break
        clean_name = clean_name.lstrip("/. ").strip()
        return extracted_title, clean_name

    def find_column(columns, keywords, exclude_keywords=[]):
        for col in columns:
            match = False
            for kw in keywords:
                clean_kw = clean_text(kw)
                if clean_kw in col:
                    match = True
                    break
            if match:
                exclude = False
                for ex in exclude_keywords:
                    clean_ex = clean_text(ex)
                    if clean_ex in col:
                        exclude = True
                        break
                if not exclude:
                    return col
        return None

    def parse_weeks_count(val):
        if val is None or pd.isna(val):
            return None
        val_str = str(val).strip()
        if val_str in ["-", "--", "nan", "none", "null", ""]:
            return None
        digits = "".join([c for c in val_str if c.isdigit()])
        if digits:
            try:
                return int(digits)
            except:
                pass
        norm = val_str.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
        if "واحد" in norm or norm == "اسبوع":
            return 1
        elif "اسبوعان" in norm or "اسبوعين" in norm:
            return 2
        elif "ثلاث" in norm:
            return 3
        elif "اربع" in norm:
            return 4
        elif "خمس" in norm:
            return 5
        elif "ست" in norm:
            return 6
        elif "سبع" in norm:
            return 7
        elif "ثمان" in norm:
            return 8
        elif "تسع" in norm:
            return 9
        elif "عشر" in norm:
            return 10
        elif "احد عشر" in norm:
            return 11
        elif "اثنا عشر" in norm:
            return 12
        elif "ثلاثة عشر" in norm:
            return 13
        elif "اربعة عشر" in norm:
            return 14
        elif "خمسة عشر" in norm:
            return 15
        return None

    imported_count = 0
    updated_count = 0

    try:
        contents1 = await file1.read()
        df1 = pd.read_excel(io.BytesIO(contents1), dtype=str)
        df1.columns = [clean_text(col) for col in df1.columns]
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"خطأ في قراءة ملف الإكسيل الأول: {str(e)}")

    if file2 is None:
        # SINGLE FILE TEMPLATE IMPORT
        col_map = {}
        col_map["national_id"] = find_column(df1.columns, ["قومي", "قومى", "national"])
        col_map["name"] = find_column(df1.columns, ["اسم عضو", "اسم دكتور", "الاسم", "الاسم"], ["برنامج", "كلية", "كليه"])
        if not col_map["name"]:
            col_map["name"] = find_column(df1.columns, ["اسم", "name"], ["برنامج", "كلية", "كليه"])
        col_map["job_title"] = find_column(df1.columns, ["درجة", "الدرجة", "وظيفة", "الوظيفة", "وظيفته", "job", "title"])
        col_map["original_workplace"] = find_column(df1.columns, ["قدوم", "القدوم", "جهة", "عمل", "workplace"])
        col_map["phone"] = find_column(df1.columns, ["هاتف", "تليفون", "موبايل", "phone", "mobile"])
        col_map["email"] = find_column(df1.columns, ["بريد", "ايميل", "الإميل", "الإيميل", "email"])
        col_map["faculty"] = find_column(df1.columns, ["كلية", "كليه", "كليات", "faculty"])
        col_map["contract_type"] = find_column(df1.columns, ["نوع التعاقد", "التعاقد", "تعاقد", "contract"])
        col_map["mnu_job_title"] = find_column(df1.columns, ["طبيعة العمل", "منوفية", "اهلية", "أهلية", "mnu_job"])
        col_map["sem1_weeks"] = find_column(df1.columns, ["فصل اول", "فصل أول", "ترم اول", "ترم أول", "الفصل الاول", "الفصل الأول", "حضور الفصل الاول", "حضور الفصل الأول", "اسابيع الفصل الاول", "أسابيع الفصل الأول"])
        col_map["sem2_weeks"] = find_column(df1.columns, ["فصل ثاني", "فصل ثانى", "ترم ثاني", "ترم ثانى", "الفصل الثاني", "الفصل الثانى", "حضور الفصل الثاني", "حضور الفصل الثانى", "اسابيع الفصل الثاني", "أسابيع الفصل الثاني"])
        col_map["summer_weeks"] = find_column(df1.columns, ["صيفي", "صيفى", "صيف", "الفصل الصيفي", "الفصل الصيفى", "حضور الفصل الصيفي", "حضور الفصل الصيفى", "اسابيع الصيفي", "أسابيع الصيفي"])
        col_map["program"] = find_column(df1.columns, ["برنامج", "البرنامج", "program"])
        col_map["course"] = find_column(df1.columns, ["مقرر", "مقررات", "المقررات", "المقرارات", "course"])
        col_map["level"] = find_column(df1.columns, ["مستوى", "المستوى", "level"])
        col_map["semester"] = find_column(df1.columns, ["فصل", "ترم", "الترم", "semester", "الفصل الدراسي"])

        if not col_map.get("national_id"):
            raise HTTPException(status_code=400, detail="لم يتم العثور على عمود الرقم القومي في ملف الإكسيل")

        df1["_national_id_clean"] = df1[col_map["national_id"]].apply(clean_national_id)

        # Filter out rows that are entirely blank across all checked columns
        def is_non_empty_row(row):
            for col in df1.columns:
                v = str(row[col]).strip() if pd.notna(row[col]) else ""
                if v and v.lower() not in ["nan", "none", "n/a", "null", ""]:
                    return True
            return False

        df1 = df1[df1.apply(is_non_empty_row, axis=1)]

        if len(df1) == 0:
            raise HTTPException(status_code=400, detail="ملف الإكسيل المرفوع فارغ ولا يحتوي على أي بيانات. يرجى تعبئة بيانات الأعضاء وحفظ الملف (Ctrl + S) في برنامج Excel والتأكد من اختيار الملف المحفوظ ثم إعادة رفعه.")

        professors_data = {}
        row_idx = 0
        for _, row in df1.iterrows():
            row_idx += 1
            nid = row["_national_id_clean"]
            
            name_col = col_map.get("name")
            name = str(row[name_col]).strip() if name_col and pd.notna(row[name_col]) else ""
            
            job_title_col = col_map.get("job_title")
            job_title_val = str(row[job_title_col]).strip() if job_title_col and pd.notna(row[job_title_col]) else ""
            
            workplace_col = col_map.get("original_workplace")
            workplace = str(row[workplace_col]).strip() if workplace_col and pd.notna(row[workplace_col]) else ""
            
            phone_col = col_map.get("phone")
            phone_val = row[phone_col] if phone_col and pd.notna(row[phone_col]) else ""
            if isinstance(phone_val, float):
                phone_val = f"{int(phone_val)}"
            phone = str(phone_val).strip()
            
            email_col = col_map.get("email")
            email = str(row[email_col]).strip() if email_col and pd.notna(row[email_col]) else ""
            
            faculty_col = col_map.get("faculty")
            faculty_name = str(row[faculty_col]).strip() if faculty_col and pd.notna(row[faculty_col]) else ""

            contract_col = col_map.get("contract_type")
            contract_val = str(row[contract_col]).strip() if contract_col and pd.notna(row[contract_col]) else ""

            mnu_job_col = col_map.get("mnu_job_title")
            mnu_job_val = str(row[mnu_job_col]).strip() if mnu_job_col and pd.notna(row[mnu_job_col]) else ""

            sem1_col = col_map.get("sem1_weeks")
            sem1_val = parse_weeks_count(row[sem1_col]) if sem1_col and pd.notna(row[sem1_col]) else None

            sem2_col = col_map.get("sem2_weeks")
            sem2_val = parse_weeks_count(row[sem2_col]) if sem2_col and pd.notna(row[sem2_col]) else None

            summer_col = col_map.get("summer_weeks")
            summer_val = parse_weeks_count(row[summer_col]) if summer_col and pd.notna(row[summer_col]) else None

            def clean_field(val):
                val_str = str(val).strip()
                if val_str.lower() in ["nan", "none", "n/a", ""]:
                    return ""
                return val_str

            c_name = clean_field(name)
            c_job = clean_field(job_title_val)
            c_workplace = clean_field(workplace)
            c_phone = clean_field(phone)
            c_email = clean_field(email)
            c_contract_raw = clean_field(contract_val)
            c_mnu_job = clean_field(mnu_job_val)

            c_contract_type = ""
            c_work_days = ""
            if c_contract_raw:
                if "ساع" in c_contract_raw or "ساعة" in c_contract_raw:
                    c_contract_type = "بالساعة"
                    c_work_days = "بالساعة"
                elif "كلي" in c_contract_raw or "كلى" in c_contract_raw or "5" in c_contract_raw:
                    c_contract_type = "كلي"
                    c_work_days = "5 أيام في الأسبوع"
                elif "جزئ" in c_contract_raw or "جزئي" in c_contract_raw:
                    c_contract_type = "جزئي"
                    if "يوم واحد" in c_contract_raw or "1" in c_contract_raw or ("يوم" in c_contract_raw and "يومان" not in c_contract_raw and "يومين" not in c_contract_raw and "3" not in c_contract_raw):
                        c_work_days = "يوم واحد"
                    elif "3" in c_contract_raw or "ثلاث" in c_contract_raw:
                        c_work_days = "3 أيام"
                    else:
                        c_work_days = "يومان"
                elif "بدون" in c_contract_raw or "غير محدد" in c_contract_raw:
                    c_contract_type = "بدون تعاقد"
                    c_work_days = "بدون تعاقد"

            key = nid if nid else f"__row_{row_idx}"

            if key not in professors_data:
                professors_data[key] = {
                    "national_id": nid,
                    "name_ar": c_name,
                    "job_title_excel": c_job,
                    "original_workplace": c_workplace,
                    "phone": c_phone,
                    "email": c_email,
                    "contract_type_raw": c_contract_raw,
                    "contract_type": c_contract_type,
                    "work_days": c_work_days,
                    "mnu_job_title": c_mnu_job,
                    "sem1_weeks": sem1_val,
                    "sem2_weeks": sem2_val,
                    "summer_weeks": summer_val,
                    "assignments": []
                }
            else:
                if sem1_val is not None and professors_data[key].get("sem1_weeks") is None:
                    professors_data[key]["sem1_weeks"] = sem1_val
                if sem2_val is not None and professors_data[key].get("sem2_weeks") is None:
                    professors_data[key]["sem2_weeks"] = sem2_val
                if summer_val is not None and professors_data[key].get("summer_weeks") is None:
                    professors_data[key]["summer_weeks"] = summer_val

            if faculty_name:
                professors_data[key]["assignments"].append({
                    "faculty_name": faculty_name
                })

        validation_errors = []
        valid_professors_data = {}
        
        existing_phones = {}
        for p in db.query(models.Professor).filter(models.Professor.phone.isnot(None)).all():
            if p.phone:
                existing_phones[p.phone] = p.national_id
        
        seen_phones_in_excel = {}
        
        for key, data in professors_data.items():
            nid = data["national_id"]
            name_label = data["name_ar"] or f"الصف رقم ({key.replace('__row_', '')})"
            row_errors = []

            if not data["name_ar"]:
                row_errors.append("اسم عضو هيئة التدريس غير محدد")
            elif len(data["name_ar"].split()) < 3:
                row_errors.append("الاسم يجب أن يكون ثلاثياً على الأقل")
                
            if not nid:
                row_errors.append("الرقم القومي غير محدد أو فارغ")
            elif len(nid) != 14 or not (nid.startswith('2') or nid.startswith('3')):
                row_errors.append("الرقم القومي غير صحيح (يجب أن يتكون من 14 رقماً ويبدأ بـ 2 أو 3)")
                
            if not data["job_title_excel"]:
                row_errors.append("الدرجة العلمية غير محددة")
                
            if not data["original_workplace"]:
                row_errors.append("جهة القدوم غير محددة")
                
            if not data["phone"]:
                row_errors.append("رقم الهاتف غير محدد")
            elif not data["phone"].startswith("01") or len(data["phone"]) != 11:
                row_errors.append("رقم الهاتف غير صحيح (يجب أن يكون 11 رقماً ويبدأ بـ 01)")
            else:
                phone = data["phone"]
                if phone in existing_phones and existing_phones[phone] != nid:
                    row_errors.append("رقم الهاتف مسجل مسبقاً لعضو آخر")
                elif phone in seen_phones_in_excel and seen_phones_in_excel[phone] != nid:
                    row_errors.append("رقم الهاتف مكرر لعضو آخر داخل الملف")
                else:
                    seen_phones_in_excel[phone] = nid
                    
            if not data["assignments"]:
                row_errors.append("الكلية التابع لها غير محددة")

            if not data["contract_type"]:
                row_errors.append("نوع التعاقد غير محدد")

            if row_errors:
                validation_errors.append(f"{name_label}: " + " - ".join(row_errors))
            else:
                valid_professors_data[nid] = data

        if not valid_professors_data and validation_errors:
            raise HTTPException(status_code=400, detail="فشل الاستيراد لوجود أخطاء في البيانات:\n" + "\n".join([f"- {e}" for e in validation_errors]))

        if not valid_professors_data and not validation_errors:
            raise HTTPException(status_code=400, detail="لم يتم العثور على أي صفوف صالحة للاستيراد في الملف. يرجى التأكد من كتابة البيانات وحفظ الملف قبل الرفع.")

        updated_names = []
        imported_names = []

        for nid, data in valid_professors_data.items():
            job_title, clean_name = parse_job_title_and_clean_name(data["name_ar"])
            if data["job_title_excel"]:
                title_map = {
                    "أخصائي": "أخصائي",
                    "معيد": "معيد",
                    "مدرس مساعد": "م.م", "م.م": "م.م",
                    "مدرس": "د", "دكتور": "د", "د": "د",
                    "أستاذ مساعد": "أ.م.د", "أ.م.د": "أ.م.د",
                    "أستاذ": "أ.م", "أ.م": "أ.م", "أ.د": "أ.م"
                }
                cleaned_jt = clean_text(data["job_title_excel"])
                for k, v in title_map.items():
                    if clean_text(k) == cleaned_jt:
                        job_title = v
                        break

            sem1_v = data.get("sem1_weeks")
            sem2_v = data.get("sem2_weeks")
            summer_v = data.get("summer_weeks")

            db_prof = db.query(models.Professor).filter(models.Professor.national_id == nid).first()
            if db_prof:
                db_prof.name_ar = clean_name
                db_prof.phone = data["phone"] or db_prof.phone
                db_prof.email = data["email"] or db_prof.email
                db_prof.job_title = job_title or db_prof.job_title
                db_prof.original_workplace = data["original_workplace"] or db_prof.original_workplace
                db_prof.contract_type = data["contract_type"] or db_prof.contract_type
                db_prof.work_days = data["work_days"] or db_prof.work_days
                db_prof.mnu_job_title = data["mnu_job_title"] or db_prof.mnu_job_title
                if sem1_v is not None:
                    db_prof.semester1_weeks = sem1_v
                if sem2_v is not None:
                    db_prof.semester2_weeks = sem2_v
                if summer_v is not None:
                    db_prof.summer_weeks = summer_v

                ay = db_prof.academic_year or "2026/2027"
                ay_map = {}
                if db_prof.academic_year_weeks:
                    try:
                        ay_map = json.loads(db_prof.academic_year_weeks)
                    except:
                        ay_map = {}
                if ay not in ay_map:
                    ay_map[ay] = {}
                if sem1_v is not None:
                    ay_map[ay]["semester1"] = sem1_v
                if sem2_v is not None:
                    ay_map[ay]["semester2"] = sem2_v
                if summer_v is not None:
                    ay_map[ay]["summer"] = summer_v
                db_prof.academic_year_weeks = json.dumps(ay_map, ensure_ascii=False)

                updated_count += 1
                updated_names.append(clean_name)
            else:
                ay = "2026/2027"
                ay_map = {
                    ay: {
                        "semester1": sem1_v,
                        "semester2": sem2_v,
                        "summer": summer_v
                    }
                }
                ay_weeks_val = json.dumps(ay_map, ensure_ascii=False)

                db_prof = models.Professor(
                    name_ar=clean_name,
                    name_en="",
                    national_id=nid,
                    phone=data["phone"],
                    email=data["email"],
                    job_title=job_title,
                    original_workplace=data["original_workplace"],
                    contract_type=data["contract_type"],
                    work_days=data["work_days"],
                    mnu_job_title=data["mnu_job_title"],
                    academic_year=ay,
                    semester1_weeks=sem1_v,
                    semester2_weeks=sem2_v,
                    summer_weeks=summer_v,
                    academic_year_weeks=ay_weeks_val
                )
                db.add(db_prof)
                db.flush()
                imported_count += 1
                imported_names.append(clean_name)

            for assign in data["assignments"]:
                fac_name = assign["faculty_name"]
                cleaned_fac_name = clean_text(fac_name)
                target_faculty = None
                faculties = db.query(models.Faculty).all()
                for f in faculties:
                    if clean_text(f.name) == cleaned_fac_name:
                        target_faculty = f
                        break
                if not target_faculty:
                    target_faculty = models.Faculty(name=fac_name)
                    db.add(target_faculty)
                    db.flush()

                if target_faculty not in db_prof.faculties:
                    db_prof.faculties.append(target_faculty)

        db.commit()

        return {
            "success": True,
            "imported": imported_count,
            "imported_names": imported_names,
            "updated": updated_count,
            "updated_names": updated_names,
            "total": imported_count + updated_count,
            "errors": validation_errors
        }

    else:
        # TWO FILES MERGING IMPORT (COMPATIBILITY)
        try:
            contents2 = await file2.read()
            df2 = pd.read_excel(io.BytesIO(contents2))
            df2.columns = [clean_text(col) for col in df2.columns]
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"خطأ في قراءة ملف الإكسيل الثاني: {str(e)}")

        col_map1 = {}
        col_map1["national_id"] = find_column(df1.columns, ["قومي", "قومى", "national"])
        col_map1["name"] = find_column(df1.columns, ["اسم دكتور", "اسم عضو", "الاسم", "الاسم"], ["برنامج", "كلية", "كليه"])
        if not col_map1["name"]:
            col_map1["name"] = find_column(df1.columns, ["اسم", "name"], ["برنامج", "كلية", "كليه"])
        col_map1["email"] = find_column(df1.columns, ["بريد", "ايميل", "email"])

        col_map2 = {}
        col_map2["national_id"] = find_column(df2.columns, ["قومي", "قومى", "national"])
        col_map2["phone"] = find_column(df2.columns, ["موبايل", "هاتف", "تليفون", "phone", "mobile"])
        col_map2["name"] = find_column(df2.columns, ["اسم دكتور", "اسم عضو", "الاسم", "الاسم"], ["برنامج", "كلية", "كليه"])
        if not col_map2["name"]:
            col_map2["name"] = find_column(df2.columns, ["اسم", "name"], ["برنامج", "كلية", "كليه"])
        col_map2["faculty"] = find_column(df2.columns, ["كلية", "كليه", "كليات", "faculty", "college"])

        if not col_map1.get("national_id"):
            raise HTTPException(status_code=400, detail="لم يتم العثور على عمود الرقم القومي في الملف الأول")
        if not col_map2.get("national_id"):
            raise HTTPException(status_code=400, detail="لم يتم العثور على عمود الرقم القومي في الملف الثاني")

        df1["_national_id_clean"] = df1[col_map1["national_id"]].apply(clean_national_id)
        df2["_national_id_clean"] = df2[col_map2["national_id"]].apply(clean_national_id)

        df1 = df1[df1["_national_id_clean"] != ""]
        df2 = df2[df2["_national_id_clean"] != ""]

        professors_data = {}
        for _, row in df1.iterrows():
            nid = row["_national_id_clean"]
            name_col = col_map1.get("name")
            name = row[name_col].strip() if name_col and pd.notna(row[name_col]) else ""
            email_col = col_map1.get("email")
            email = row[email_col].strip() if email_col and pd.notna(row[email_col]) else ""
            professors_data[nid] = {
                "national_id": nid,
                "name_ar": name,
                "email": email,
                "phone": "",
                "faculty_name": "",
            }

        for _, row in df2.iterrows():
            nid = row["_national_id_clean"]
            phone_col = col_map2.get("phone")
            phone_val = row[phone_col] if phone_col and pd.notna(row[phone_col]) else ""
            if isinstance(phone_val, float):
                phone_val = f"{int(phone_val)}"
            phone = str(phone_val).strip()
            name_col = col_map2.get("name")
            name = row[name_col].strip() if name_col and pd.notna(row[name_col]) else ""
            faculty_col = col_map2.get("faculty")
            faculty_name = row[faculty_col].strip() if faculty_col and pd.notna(row[faculty_col]) else ""

            if nid in professors_data:
                if not professors_data[nid]["name_ar"] and name:
                    professors_data[nid]["name_ar"] = name
                professors_data[nid]["phone"] = phone
                professors_data[nid]["faculty_name"] = faculty_name
            else:
                professors_data[nid] = {
                    "national_id": nid,
                    "name_ar": name,
                    "email": "",
                    "phone": phone,
                    "faculty_name": faculty_name,
                }

        for nid, data in professors_data.items():
            if not data["name_ar"]:
                continue
            job_title, clean_name = parse_job_title_and_clean_name(data["name_ar"])
            db_prof = db.query(models.Professor).filter(models.Professor.national_id == nid).first()
            if db_prof:
                db_prof.name_ar = clean_name
                db_prof.phone = data["phone"] or db_prof.phone
                db_prof.email = data["email"] or db_prof.email
                db_prof.job_title = job_title or db_prof.job_title
                updated_count += 1
            else:
                db_prof = models.Professor(
                    name_ar=clean_name,
                    name_en="",
                    national_id=nid,
                    phone=data["phone"],
                    email=data["email"],
                    job_title=job_title,
                    original_workplace=""
                )
                db.add(db_prof)
                imported_count += 1

            if data["faculty_name"]:
                cleaned_fac_name = clean_text(data["faculty_name"])
                faculties = db.query(models.Faculty).all()
                target_faculty = None
                for f in faculties:
                    if clean_text(f.name) == cleaned_fac_name:
                        target_faculty = f
                        break
                if not target_faculty:
                    target_faculty = models.Faculty(name=data["faculty_name"])
                    db.add(target_faculty)
                    db.flush()
                if target_faculty not in db_prof.faculties:
                    db_prof.faculties.append(target_faculty)
        db.commit()

    return {
        "success": True,
        "imported": imported_count,
        "updated": updated_count,
        "total": imported_count + updated_count
    }

    
@app.post("/api/professors", response_model=schemas.ProfessorOut)
def create_professor(professor: schemas.ProfessorCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    # 1. التحقق أولاً: هل الرقم القومي موجود مسبقاً؟
    db_prof = db.query(models.Professor).filter(models.Professor.national_id == professor.national_id).first()
    if db_prof:
        raise HTTPException(status_code=400, detail="عضو هيئة التدريس بهذا الرقم القومي مسجل بالفعل")
        
    if professor.phone:
        db_prof_phone = db.query(models.Professor).filter(models.Professor.phone == professor.phone).first()
        if db_prof_phone:
            raise HTTPException(status_code=400, detail="رقم الهاتف مسجل مسبقاً لعضو آخر")

    # 2. إنشاء كائن Professor جديد
    ay_weeks_val = None
    if professor.academic_year_weeks is not None:
        ay_weeks_val = json.dumps(professor.academic_year_weeks, ensure_ascii=False) if isinstance(professor.academic_year_weeks, (dict, list)) else str(professor.academic_year_weeks)

    new_professor = models.Professor(
        name_ar=professor.name_ar,
        name_en=professor.name_en,
        national_id=professor.national_id,
        phone=professor.phone,
        email=professor.email,
        job_title=professor.job_title,
        original_workplace=professor.original_workplace,
        contract_type=professor.contract_type,
        work_days=professor.work_days,
        mnu_job_title=professor.mnu_job_title,
        academic_year=professor.academic_year,
        semester1_weeks=professor.semester1_weeks,
        semester2_weeks=professor.semester2_weeks,
        summer_weeks=professor.summer_weeks,
        academic_year_weeks=ay_weeks_val
    )
    
    # 3. ربط الكليات (علاقة many-to-many)
    if professor.faculty_ids:
        faculties = db.query(models.Faculty).filter(models.Faculty.id.in_(professor.faculty_ids)).all()
        new_professor.faculties = faculties
    
    # 4. الحفظ في قاعدة البيانات
    db.add(new_professor)
    db.commit()
    db.refresh(new_professor)
    
    # 5. ربط المقررات (علاقة one-to-many: تحديث حقل professor_id في جدول المقررات)
    if professor.course_ids:
        db.query(models.Course).filter(models.Course.id.in_(professor.course_ids)).update(
            {models.Course.professor_id: new_professor.id, models.Course.year: professor.academic_year},
            synchronize_session=False
        )
        db.commit()
        db.refresh(new_professor)
        
    user_role_str = get_user_role_display(current_user)
    action_text = f"قام بإضافة عضو هيئة تدريس جديد: {new_professor.name_ar}"
    
    if new_professor.faculties:
        for fac in new_professor.faculties:
            create_notification(db, fac.id, f"{current_user.username} ({user_role_str})", action_text)
    else:
        create_notification(db, None, f"{current_user.username} ({user_role_str})", action_text)
    db.commit()
    
    return new_professor

@app.put("/api/professors/{id}", response_model=schemas.ProfessorOut)
def update_professor(id: int, professor: schemas.ProfessorCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    db_prof = db.query(models.Professor).filter(models.Professor.id == id).first()
    if not db_prof:
        raise HTTPException(status_code=404, detail="عضو هيئة التدريس غير موجود")
    
    # التحقق من الرقم القومي إذا تغير
    if db_prof.national_id != professor.national_id:
        existing = db.query(models.Professor).filter(models.Professor.national_id == professor.national_id).first()
        if existing:
            raise HTTPException(status_code=400, detail="عضو هيئة التدريس بهذا الرقم القومي مسجل بالفعل")
            
    # التحقق من رقم الهاتف إذا تغير
    if professor.phone and db_prof.phone != professor.phone:
        existing_phone = db.query(models.Professor).filter(models.Professor.phone == professor.phone).first()
        if existing_phone:
            raise HTTPException(status_code=400, detail="رقم الهاتف مسجل مسبقاً لعضو آخر")

    ay_weeks_val = None
    if professor.academic_year_weeks is not None:
        ay_weeks_val = json.dumps(professor.academic_year_weeks, ensure_ascii=False) if isinstance(professor.academic_year_weeks, (dict, list)) else str(professor.academic_year_weeks)

    db_prof.name_ar = professor.name_ar
    db_prof.name_en = professor.name_en
    db_prof.national_id = professor.national_id
    db_prof.phone = professor.phone
    db_prof.email = professor.email
    db_prof.job_title = professor.job_title
    db_prof.original_workplace = professor.original_workplace
    db_prof.contract_type = professor.contract_type
    db_prof.work_days = professor.work_days
    db_prof.mnu_job_title = professor.mnu_job_title
    db_prof.academic_year = professor.academic_year
    db_prof.semester1_weeks = professor.semester1_weeks
    db_prof.semester2_weeks = professor.semester2_weeks
    db_prof.summer_weeks = professor.summer_weeks
    db_prof.academic_year_weeks = ay_weeks_val
    
    # تحديث الكليات
    if professor.faculty_ids is not None:
        faculties = db.query(models.Faculty).filter(models.Faculty.id.in_(professor.faculty_ids)).all()
        db_prof.faculties = faculties
        
    # تحديث المقررات: فك ارتباط المقررات القديمة أولاً
    db.query(models.Course).filter(models.Course.professor_id == id).update(
        {models.Course.professor_id: None, models.Course.year: "2026/2027"},
        synchronize_session=False
    )
    
    # ربط المقررات الجديدة
    if professor.course_ids:
        db.query(models.Course).filter(models.Course.id.in_(professor.course_ids)).update(
            {models.Course.professor_id: id, models.Course.year: professor.academic_year},
            synchronize_session=False
        )
        
    db.commit()
    db.refresh(db_prof)
    
    user_role_str = get_user_role_display(current_user)
    action_text = f"قام بتعديل بيانات عضو هيئة تدريس: {db_prof.name_ar}"
    if db_prof.faculties:
        for fac in db_prof.faculties:
            create_notification(db, fac.id, f"{current_user.username} ({user_role_str})", action_text)
    else:
        create_notification(db, None, f"{current_user.username} ({user_role_str})", action_text)
    db.commit()
    
    return db_prof

@app.get("/api/professors", response_model=list[schemas.ProfessorOut])
def get_professors(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        profs = db.query(models.Professor).filter(models.Professor.is_deleted == False).all()
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        profs = db.query(models.Professor).filter(
            models.Professor.is_deleted == False,
            (models.Professor.faculties.any(models.Faculty.id.in_(faculty_ids))) | 
            (~models.Professor.faculties.any())
        ).all()
    elif current_user.faculty_id:
        profs = db.query(models.Professor).filter(
            models.Professor.is_deleted == False,
            (models.Professor.faculties.any(models.Faculty.id == current_user.faculty_id)) | 
            (~models.Professor.faculties.any())
        ).all()
    else:
        profs = db.query(models.Professor).filter(models.Professor.is_deleted == False).all()

    prof_ids = [p.id for p in profs]
    prof_courses_map = {pid: [] for pid in prof_ids}
    if prof_ids:
        plan_items = db.query(models.StudyPlanItem).join(models.StudyPlan).filter(
            models.StudyPlan.is_deleted == False,
            models.StudyPlanItem.professor_id.in_(prof_ids)
        ).all()
        
        seen_courses = set()
        for item in plan_items:
            c = item.course or item.base_course
            if c:
                sp = item.study_plan
                sem = sp.semester if sp else c.semester
                yr = sp.academic_year if sp else c.year
                lvl = item.level if item.level is not None else c.level
                key = (item.professor_id, c.id, sem, yr, lvl)
                if key not in seen_courses:
                    seen_courses.add(key)
                    c_dict = {
                        "id": c.id,
                        "code": c.code or "",
                        "name_ar": c.name_ar or "",
                        "name_en": c.name_en or "",
                        "level": lvl if lvl is not None else 1,
                        "semester": sem or "الفصل الدراسي الأول",
                        "year": yr or "2026/2027",
                        "faculty_id": c.faculty_id,
                        "program_id": c.program_id,
                        "theory_hours": c.theory_hours or 0.0,
                        "practical_hours": c.practical_hours or 0.0,
                        "exercise_hours": c.exercise_hours or 0.0,
                        "activity_hours": c.activity_hours or 0.0,
                        "total_grade": c.total_grade or 0.0,
                        "theory_grade": c.theory_grade or 0.0,
                        "practical_grade": c.practical_grade or 0.0,
                        "year_work_grade": c.year_work_grade or 0.0,
                        "exam_time_hours": c.exam_time_hours or "",
                        "duration": c.duration or "",
                        "is_bundle": c.is_bundle or False,
                        "department_name": c.department_name or "",
                        "faculty": c.faculty,
                        "program": c.program,
                        "modules": c.modules or []
                    }
                    prof_courses_map[item.professor_id].append(c_dict)
                    
    prof_out_list = []
    for p in profs:
        p_dict = {
            "id": p.id,
            "name_ar": p.name_ar,
            "name_en": p.name_en or "",
            "national_id": p.national_id or "",
            "phone": p.phone or "",
            "email": p.email or "",
            "job_title": p.job_title or "",
            "original_workplace": p.original_workplace or "",
            "contract_type": p.contract_type or "",
            "work_days": p.work_days or "",
            "mnu_job_title": p.mnu_job_title or "",
            "academic_year": p.academic_year or "2026/2027",
            "semester1_weeks": p.semester1_weeks,
            "semester2_weeks": p.semester2_weeks,
            "summer_weeks": p.summer_weeks,
            "academic_year_weeks": p.academic_year_weeks,
            "faculties": [schemas.FacultyOut.from_orm(f) for f in p.faculties] if p.faculties else [],
            "courses": prof_courses_map.get(p.id, [])
        }
        prof_out_list.append(schemas.ProfessorOut(**p_dict))

    return prof_out_list
# ==========================================
# مسارات المقررات (Courses APIs)
# ==========================================

@app.post("/api/courses", response_model=schemas.CourseOut)
def create_course(course: schemas.CourseCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role in [models.UserRole.faculty_admin, models.UserRole.faculty_professor] and current_user.faculty_id != course.faculty_id:
        raise HTTPException(status_code=403, detail="لا يمكنك إضافة مقرر في كلية أخرى")
    
    course_data = course.dict(exclude={"modules"})
    modules_data = course.dict().get("modules", [])
    
    new_course = models.Course(**course_data)
    db.add(new_course)
    db.flush()

    if new_course.is_bundle and modules_data:
        for mod in modules_data:
            m_obj = models.CourseModule(course_id=new_course.id, **mod)
            db.add(m_obj)

    user_role_str = get_user_role_display(current_user)
    action_text = f"قام بإضافة مقرر جديد: {new_course.name_ar}"
    create_notification(db, new_course.faculty_id, f"{current_user.username} ({user_role_str})", action_text)
    
    db.commit()
    db.refresh(new_course)
    return new_course


@app.get("/api/courses", response_model=list[schemas.CourseOut])
def get_courses(faculty_id: Optional[int] = None, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    query = db.query(models.Course).filter(models.Course.is_deleted == False)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id is not None:
            query = query.filter(models.Course.faculty_id == faculty_id)
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id is not None:
            if faculty_id in faculty_ids:
                query = query.filter(models.Course.faculty_id == faculty_id)
            else:
                return []
        else:
            query = query.filter(models.Course.faculty_id.in_(faculty_ids))
    elif current_user.faculty_id:
        query = query.filter(models.Course.faculty_id == current_user.faculty_id)
    elif faculty_id is not None:
        query = query.filter(models.Course.faculty_id == faculty_id)
    return query.all()

@app.get("/api/courses/template")
def download_courses_template(db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.utils import get_column_letter
    from tempfile import NamedTemporaryFile
    from fastapi.responses import FileResponse
    
    def clean_range_name(name):
        forbidden = [" ", "/", "-", "(", ")", "[", "]", "{", "}", "&", "*", "+", ",", ".", "\\\\", "?", "!", "=", "|", ":", ";", "'", '"', "<", ">", "@", "#", "$", "%", "^"]
        for char in forbidden:
            name = name.replace(char, "_")
        name = name.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
        if name and name[0].isdigit():
            name = "_" + name
        return name

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "المقررات"
    ws.views.sheetView[0].showGridLines = True
    ws.sheet_view.rightToLeft = True
    
    headers_row1 = [
        "اسم الكلية", "اسم البرنامج", "كود المقرر", "اسم المقرر بالعربية", "اسم المقرر بالإنجليزية",
        "نوع المقرر", "المستوى", "الفصل الدراسي", "البرامج الاخرى المسجل بها", "متطلب",
        "المتطلب السابق", "المقررات المتزامنة", "يضاف للمعدل التراكمي", "مادة نجاح أو رسوب", "الساعات المعتمدة",
        "الساعات", "", "", "", "", "", "",
        "الدرجات", "", "", "", "", "", "", "", "", "", "",
        "نسبة النجاح", "نسبة الرسوب النظري", "كود المجموعة الاختيارية", "عدد المقررات او الوحدات الاختيارية",
        "تسجيل المقرر في الصيفي"
    ]

    headers_row2 = [
        "", "", "", "", "",
        "", "", "", "", "",
        "", "", "", "", "",
        "محاضرات", "تدريب", "عملي", "ساعات التدريب الميداني", "ساعات الامتحان", "الساعات الدراسية", "المجموع",
        "أعمال الفصل", "نهاية الفصل", "منتصف الفصل", "تحريري خلال الفصل", "شفوي", "عملي", "كلينك", "تقييم نهائي", "منتصف الفصل ٢", "Attendance, Activity & Attitude", "المجموع",
        "", "", "", "", ""
    ]
    
    ws.append(headers_row1)
    ws.append(headers_row2)
    
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:C2")
    ws.merge_cells("D1:D2")
    ws.merge_cells("E1:E2")
    ws.merge_cells("F1:F2")
    ws.merge_cells("G1:G2")
    ws.merge_cells("H1:H2")
    ws.merge_cells("I1:I2")
    ws.merge_cells("J1:J2")
    ws.merge_cells("K1:K2")
    ws.merge_cells("L1:L2")
    ws.merge_cells("M1:M2")
    ws.merge_cells("N1:N2")
    ws.merge_cells("O1:O2")
    
    ws.merge_cells("P1:V1") # الساعات
    ws.merge_cells("W1:AG1") # الدرجات
    
    ws.merge_cells("AH1:AH2")
    ws.merge_cells("AI1:AI2")
    ws.merge_cells("AJ1:AJ2")
    ws.merge_cells("AK1:AK2")
    ws.merge_cells("AL1:AL2")
    
    from openpyxl.styles import Border, Side
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    header_fill = openpyxl.styles.PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    header_font = openpyxl.styles.Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    for row in range(1, 3):
        for col_idx in range(1, len(headers_row1) + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            
    data_font = openpyxl.styles.Font(name="Calibri", size=12)
    for row_idx in range(3, 201):
        for col_idx in range(1, len(headers_row1) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            
    ws.freeze_panes = "A3"
        
    # Helper Sheet
    ws_lists = wb.create_sheet(title="قوائم المرجعية")
    ws_lists.sheet_state = "hidden"
    ws_lists.sheet_view.rightToLeft = True
    
    # 1. Faculties in A
    faculties = db.query(models.Faculty).all()
    fac_names = [f.name for f in faculties]
    for idx, name in enumerate(fac_names, 1):
        ws_lists.cell(row=idx, column=1, value=name)
        
    # 2. Programs for each Faculty
    for fac_idx, fac in enumerate(faculties, 2):
        progs = db.query(models.Program).filter(models.Program.faculty_id == fac.id).all()
        prog_names = [p.name for p in progs]
        if not prog_names:
            prog_names = ["لا يوجد برامج حاليا"]
        for row_idx, prog_name in enumerate(prog_names, 1):
            ws_lists.cell(row=row_idx, column=fac_idx, value=prog_name)
            
        col_letter = get_column_letter(fac_idx)
        range_name = clean_range_name(fac.name)
        ref_str = f"'قوائم المرجعية'!${col_letter}$1:${col_letter}${len(prog_names)}"
        
        new_range = DefinedName(name=range_name)
        new_range.value = ref_str
        wb.defined_names.add(new_range)
        
    # 3. Levels for each Program
    program_levels_map = {
        "المستوى العام": ["المستوى 0", "المستوى 1"],
        "الأمن السيبراني": ["المستوى 2", "المستوى 3", "المستوى 4"],
        "الحوسبة السحابية": ["المستوى 2", "المستوى 3", "المستوى 4"],
        "الحوسبة عالية الكفاءة": ["المستوى 2", "المستوى 3", "المستوى 4"],
        "تخطيط وتشييد المدن الذكية": ["المستوى 0", "المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "هندسة الحاسوب": ["المستوى 0", "المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "هندسة المواد وإدارة التصنيع": ["المستوى 0", "المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "إنترنت الأشياء وتحليل البيانات الضخمة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "ذكاء الآلة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "علوم البيانات": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "اللغة الإنجليزية والترجمة التخصصية": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "علوم التمريض": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تمريض الطوارئ": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تمريض القبالة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تمريض حديثي الولادة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "فارم دي (الصيدلة الإكلينيكية)": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "الطب البيطري": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "طب وجراحة الفم والأسنان": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "برنامج عام": ["المستوى 1", "المستوى 2"],
        "تكنولوجيا الرعاية التنفسية": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تكنولوجيا المختبرات الطبية": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تكنولوجيا صناعة تركيبات الأسنان": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "تكنولوجيا علوم الأشعة والتصوير الطبي": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"],
        "العلاج الطبيعي": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "الطب والجراحة اللائحة الجديدة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"],
        "الطب والجراحة اللائحة القديمة": ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4", "المستوى 5"]
    }
    
    def get_levels_for_program(prog_name):
        def normalize_text(text):
            if not text:
                return ""
            return text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه").replace(" ", "").strip()
        norm_name = normalize_text(prog_name)
        for k, v in program_levels_map.items():
            if normalize_text(k) in norm_name or norm_name in normalize_text(k):
                return v
        return ["المستوى 1", "المستوى 2", "المستوى 3", "المستوى 4"]

    start_level_col = 2 + len(faculties)
    all_db_programs = db.query(models.Program).all()
    for prog_idx, prog in enumerate(all_db_programs):
        col_idx = start_level_col + prog_idx
        col_letter = get_column_letter(col_idx)
        levels_list = get_levels_for_program(prog.name)
        for r_idx, lvl in enumerate(levels_list, 1):
            ws_lists.cell(row=r_idx, column=col_idx, value=lvl)
            
        range_name = clean_range_name(prog.name) + "_levels"
        ref_str = f"'قوائم المرجعية'!${col_letter}$1:${col_letter}${len(levels_list)}"
        new_range = DefinedName(name=range_name)
        new_range.value = ref_str
        wb.defined_names.add(new_range)

    # 5. Courses list for each Program
    start_course_col = 2 + len(faculties) + len(all_db_programs)
    
    # Find the general program in 'كلية تكنولوجيا العلوم الصحية التطبيقية'
    health_sci_fac = db.query(models.Faculty).filter(models.Faculty.name == "كلية تكنولوجيا العلوم الصحية التطبيقية").first()
    health_general_prog = None
    if health_sci_fac:
        health_general_prog = db.query(models.Program).filter(
            models.Program.faculty_id == health_sci_fac.id,
            models.Program.name == "برنامج عام"
        ).first()
        
    health_general_courses = []
    if health_general_prog:
        health_general_courses_query = db.query(models.Course.name_ar).filter(models.Course.program_id == health_general_prog.id).distinct().all()
        health_general_courses = [c[0] for c in health_general_courses_query if c[0]]

    for prog_idx, prog in enumerate(all_db_programs):
        col_idx = start_course_col + prog_idx
        col_letter = get_column_letter(col_idx)
        
        # Get courses for this program
        prog_courses_query = db.query(models.Course.name_ar).filter(models.Course.program_id == prog.id).distinct().all()
        prog_courses = [c[0] for c in prog_courses_query if c[0]]
        
        # If this program belongs to Health Sciences faculty, add health general program courses
        if health_sci_fac and prog.faculty_id == health_sci_fac.id and prog.id != health_general_prog.id:
            prog_courses = list(set(prog_courses + health_general_courses))
            
        if not prog_courses:
            prog_courses = ["لا يوجد مقررات حاليا"]
            
        for r_idx, c_name in enumerate(prog_courses, 1):
            ws_lists.cell(row=r_idx, column=col_idx, value=c_name)
            
        range_name = clean_range_name(prog.name) + "_courses"
        ref_str = f"'قوائم المرجعية'!${col_letter}$1:${col_letter}${len(prog_courses)}"
        new_range = DefinedName(name=range_name)
        new_range.value = ref_str
        wb.defined_names.add(new_range)

    # 4. Semesters in Column 150 of ws_lists
    semesters = ["الفصل الدراسي الأول", "الفصل الدراسي الثاني", "الفصل الدراسي الصيفي"]
    for idx, sem in enumerate(semesters, 1):
        ws_lists.cell(row=idx, column=150, value=sem)

    # Query unique fail rates from database for AI validation
    db_fail_rates = db.query(models.Course.fail_rate).distinct().all()
    fail_rates_list = sorted(list(set([r[0] for r in db_fail_rates if r[0] and r[0].strip() and r[0] != "-"])))
    if not fail_rates_list:
        fail_rates_list = [
            "نهاية الفصل:0",
            "نهاية الفصل:12",
            "نهاية الفصل:13.5",
            "نهاية الفصل:18",
            "نهاية الفصل:24",
            "نهاية الفصل:30",
            "نهاية الفصل:36",
            "نهاية الفصل:40",
            "نهاية الفصل:50",
            "منتصف الفصل:20"
        ]
    # Build list under 255 characters limit
    joined_str = ",".join(fail_rates_list)
    if len(joined_str) > 240:
        while fail_rates_list and len(",".join(fail_rates_list)) > 240:
            fail_rates_list.pop()
        joined_str = ",".join(fail_rates_list)
    fail_rates_str = f'"{joined_str}"'

    # Apply Excel cell validations
    dv_faculty = DataValidation(type="list", formula1=f"='قوائم المرجعية'!$A$1:$A${len(fac_names)}", allow_blank=True)
    dv_faculty.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_faculty.errorTitle = 'خطأ في الإدخال'
    dv_faculty.showErrorMessage = True
    ws.add_data_validation(dv_faculty)
    dv_faculty.add("A3:A200")
    
    dv_program = DataValidation(type="list", formula1='=INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(A3," ","_"),"أ","ا"),"إ","ا"),"ة","ه"),"ى","ي"))', allow_blank=True)
    dv_program.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_program.errorTitle = 'خطأ في الإدخال'
    dv_program.showErrorMessage = True
    ws.add_data_validation(dv_program)
    dv_program.add("B3:B200")
    
    # البرامج الاخرى المسجل بها (Column I) - تظهر جميع برامج الكلية ماعدا المختار في العمود B
    for row_idx in range(3, 201):
        col_letter = get_column_letter(200 + row_idx)
        
        # نكتب صيغة الإزاحة لـ 15 خلية في العمود المساعد
        for r in range(1, 16):
            ws_lists.cell(row=r, column=200 + row_idx, value=(
                f"=IFERROR(IF(ISERROR(MATCH('المقررات'!$B${row_idx}, INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('المقررات'!$A${row_idx},\" \",\"_\"),\"أ\",\"ا\"),\"إ\",\"ا\"),\"ة\",\"ه\"),\"ى\",\"ي\")), 0)), "
                f"INDEX(INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('المقررات'!$A${row_idx},\" \",\"_\"),\"أ\",\"ا\"),\"إ\",\"ا\"),\"ة\",\"ه\"),\"ى\",\"ي\")), {r}), "
                f"IF({r} < MATCH('المقررات'!$B${row_idx}, INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('المقررات'!$A${row_idx},\" \",\"_\"),\"أ\",\"ا\"),\"إ\",\"ا\"),\"ة\",\"ه\"),\"ى\",\"ي\")), 0), "
                f"INDEX(INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('المقررات'!$A${row_idx},\" \",\"_\"),\"أ\",\"ا\"),\"إ\",\"ا\"),\"ة\",\"ه\"),\"ى\",\"ي\")), {r}), "
                f"INDEX(INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('المقررات'!$A${row_idx},\" \",\"_\"),\"أ\",\"ا\"),\"إ\",\"ا\"),\"ة\",\"ه\"),\"ى\",\"ي\")), {r}+1))), \"\")"
            ))
        
        dv_row = DataValidation(type="list", formula1=f"='قوائم المرجعية'!${col_letter}$1:${col_letter}$15", allow_blank=True)
        dv_row.showErrorMessage = False
        dv_row.prompt = 'اختر البرنامج من القائمة، أو اكتب عدة برامج مفصولة بـ /'
        dv_row.promptTitle = 'البرامج المسجلة'
        dv_row.showInputMessage = True
        ws.add_data_validation(dv_row)
        dv_row.add(f"I{row_idx}")
    
    dv_course_type = DataValidation(type="list", formula1='"اجبارى,اختيارى,اختياري حر,تدريب,مشروع"', allow_blank=True)
    dv_course_type.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_course_type.errorTitle = 'خطأ في الإدخال'
    dv_course_type.showErrorMessage = True
    ws.add_data_validation(dv_course_type)
    dv_course_type.add("F3:F200")
    
    dv_level = DataValidation(type="list", formula1='=INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(B3," ","_"),"أ","ا"),"إ","ا"),"ة","ه"),"ى","ي") & "_levels")', allow_blank=True)
    dv_level.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_level.errorTitle = 'خطأ في الإدخال'
    dv_level.showErrorMessage = True
    ws.add_data_validation(dv_level)
    dv_level.add("G3:G200")
    
    # متطلب سابق والمقررات المتزامنة (Columns K and L) - يسمح بكتابة أكثر من مقرر يدوياً
    dv_courses = DataValidation(type="list", formula1='=INDIRECT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(B3," ","_"),"أ","ا"),"إ","ا"),"ة","ه"),"ى","ي") & "_courses")', allow_blank=True)
    dv_courses.showErrorMessage = False  # عدم منع المستخدم من كتابة نصوص مخصصة أو عدة مقررات
    dv_courses.prompt = 'اختر المقرر من القائمة، أو اكتب عدة مقررات مفصولة بـ /'
    dv_courses.promptTitle = 'متطلبات المقرر'
    dv_courses.showInputMessage = True
    ws.add_data_validation(dv_courses)
    dv_courses.add("K3:K200")
    dv_courses.add("L3:L200")
    
    dv_semester = DataValidation(type="list", formula1='"الفصل الدراسي الأول,الفصل الدراسي الثاني,الفصل الدراسي الصيفي"', allow_blank=True)
    dv_semester.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_semester.errorTitle = 'خطأ في الإدخال'
    dv_semester.showErrorMessage = True
    ws.add_data_validation(dv_semester)
    dv_semester.add("H3:H200")

    # متطلب (Column J)
    dv_requirement = DataValidation(type="list", formula1='"تخصص,جامعة,علوم أساسية,كلية,متطلب التدريب والتعلم الذاتي"', allow_blank=True)
    dv_requirement.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_requirement.errorTitle = 'خطأ في الإدخال'
    dv_requirement.showErrorMessage = True
    ws.add_data_validation(dv_requirement)
    dv_requirement.add("J3:J200")

    # يضاف للمعدل التراكمي (Column M)
    dv_added_to_gpa = DataValidation(type="list", formula1='"يضاف للمعدل التراكمي,لا يضاف للمعدل التراكمي,يضاف للساعات فقط"', allow_blank=True)
    dv_added_to_gpa.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_added_to_gpa.errorTitle = 'خطأ في الإدخال'
    dv_added_to_gpa.showErrorMessage = True
    ws.add_data_validation(dv_added_to_gpa)
    dv_added_to_gpa.add("M3:M200")
    
    # نسبة الرسوب النظري (Column AI)
    dv_fail_rate = DataValidation(type="list", formula1=fail_rates_str, allow_blank=True)
    dv_fail_rate.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_fail_rate.errorTitle = 'خطأ في الإدخال'
    dv_fail_rate.showErrorMessage = True
    ws.add_data_validation(dv_fail_rate)
    dv_fail_rate.add("AI3:AI200")

    # Yes/No for N3:N200 and AL3:AL200
    dv_yes_no = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=True)
    dv_yes_no.error = 'يرجى اختيار نعم أو لا'
    dv_yes_no.errorTitle = 'خطأ في الإدخال'
    dv_yes_no.showErrorMessage = True
    ws.add_data_validation(dv_yes_no)
    dv_yes_no.add("N3:N200")
    dv_yes_no.add("AL3:AL200")
    
    # Pre-populate Column 33 (AG) with SUM formulas for rows 3 to 200
    for row_idx in range(3, 201):
        ws.cell(row=row_idx, column=33, value=f"=SUM(W{row_idx}:AF{row_idx})")
        
    # Pre-populate Column 22 (V) with SUM formulas for hours (P:T) for rows 3 to 200
    for row_idx in range(3, 201):
        ws.cell(row=row_idx, column=22, value=f"=SUM(P{row_idx}:T{row_idx})")
        
    # Numeric DataValidation (type="decimal") for hours and grades columns
    dv_decimal = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv_decimal.error = 'المدخل يجب أن يكون رقماً (عدد صحيح أو عشري) أكبر من أو يساوي 0 وبدون أي نصوص أو مسافات'
    dv_decimal.errorTitle = 'خطأ في إدخال الرقم'
    dv_decimal.showErrorMessage = True
    ws.add_data_validation(dv_decimal)
    
    # Apply to all hours columns: O (15) to U (21)
    for col_idx in range(15, 22):
        col_letter = get_column_letter(col_idx)
        dv_decimal.add(f"{col_letter}3:{col_letter}200")
        
    # Apply to all grades columns: W (23) to AF (32)
    for col_idx in range(23, 33):
        col_letter = get_column_letter(col_idx)
        dv_decimal.add(f"{col_letter}3:{col_letter}200")

    # Apply number format '0.##' to all hours and grades cells (Rows 3 to 200)
    for row_idx in range(3, 201):
        for col_idx in list(range(15, 23)) + list(range(23, 34)):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.number_format = '0.##'
            
    temp_file = NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    wb.close()
    
    return FileResponse(
        temp_file.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="نموذج_استيراد_المقررات_الدراسية.xlsx"
    )
@app.get("/api/courses/template/medicine")
def download_medicine_courses_template(db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from tempfile import NamedTemporaryFile
    from fastapi.responses import FileResponse
    from openpyxl.utils import get_column_letter
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "المقررات"
    ws.views.sheetView[0].showGridLines = True
    ws.sheet_view.rightToLeft = True
    
    ws_ref = wb.create_sheet("قوائم المرجعية")
    ws_ref.sheet_state = 'hidden'
    departments = db.query(models.CourseModule.department_name).distinct().all()
    dept_names = [d[0] for d in departments if d[0]]
    if not dept_names:
        dept_names = ["Anatomy", "Histology", "Physiology", "Biochemistry", "Pathology", "Pharmacology", "Microbiology", "Parasitology", "Clinical Pharmacology", "Forensic Medicine", "Community Medicine", "Internal Medicine", "General Surgery", "Pediatrics", "Obstetrics and Gynecology"]
    
    for idx, d_name in enumerate(dept_names, start=1):
        ws_ref.cell(row=idx, column=1, value=d_name)
        
    headers_row1 = [
        "كود المقرر",
        "المقرر / الحزمة الدراسية (عربي)",
        "المقرر / الحزمة الدراسية (إنجليزي)",
        "الأقسام العلمية",
        "أقسام أخرى (إن وجدت)",
        "الساعات المعتمدة للأقسام",
        "توزيع الساعات المعتمدة", "", "",
        "توزيع الساعات التدريسية", "", "",
        "توزيع الدرجات", "", "",
        "الدرجات للأقسام العلمية",
        "الدرجة الكاملة للمقرر",
        "الساعات المعتمدة للمقرر",
        "عدد الأسابيع",
        "زمن الامتحان النهائي",
        "المستوى",
        "الفصل الدراسي"
    ]
    
    headers_row2 = [
        "", "", "", "", "", "",
        "نظري", "عملي", "أنشطة",
        "نظري", "عملي", "أنشطة",
        "نظري", "عملي", "أعمال سنة",
        "", "", "", "", "", "", ""
    ]
    
    ws.append(headers_row1)
    ws.append(headers_row2)
    
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:C2")
    ws.merge_cells("D1:D2")
    ws.merge_cells("E1:E2")
    ws.merge_cells("F1:F2")
    ws.merge_cells("G1:I1") 
    ws.merge_cells("J1:L1") 
    ws.merge_cells("M1:O1") 
    ws.merge_cells("P1:P2") 
    ws.merge_cells("Q1:Q2") 
    ws.merge_cells("R1:R2") 
    ws.merge_cells("S1:S2") 
    ws.merge_cells("T1:T2") 
    ws.merge_cells("U1:U2") 
    ws.merge_cells("V1:V2")
    
    header_fill = openpyxl.styles.PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = openpyxl.styles.Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    thin_border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin'), 
        right=openpyxl.styles.Side(style='thin'), 
        top=openpyxl.styles.Side(style='thin'), 
        bottom=openpyxl.styles.Side(style='thin')
    )
    
    for row in range(1, 3):
        for col_idx in range(1, 23):
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)

    wrap_alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=3, max_row=200, min_col=1, max_col=22):
        for cell in row:
            cell.border = thin_border
            cell.alignment = wrap_alignment
            
    ws.freeze_panes = "A3"
            
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 60
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 18
    for col_letter in ['G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O']:
        ws.column_dimensions[col_letter].width = 15
    ws.column_dimensions['P'].width = 18
    ws.column_dimensions['Q'].width = 18
    ws.column_dimensions['R'].width = 20
    ws.column_dimensions['S'].width = 20
    ws.column_dimensions['T'].width = 18
    ws.column_dimensions['U'].width = 20
    ws.column_dimensions['V'].width = 20

    dv_sem = DataValidation(type="list", formula1='"الفصل الدراسي الأول,الفصل الدراسي الثاني,الفصل الدراسي الصيفي"', allow_blank=True)
    ws.add_data_validation(dv_sem)
    dv_sem.add('V3:V1000')

    dv_lvl = DataValidation(type="list", formula1='"المستوى الأول,المستوى الثاني,المستوى الثالث,المستوى الرابع,المستوى الخامس"', allow_blank=True)
    ws.add_data_validation(dv_lvl)
    dv_lvl.add('U3:U1000')

    dv_dept = DataValidation(type="list", formula1=f"='قوائم المرجعية'!$A$1:$A${max(len(dept_names), 1)}", allow_blank=True)
    dv_dept.error = 'القيمة المدخلة غير صحيحة، يرجى الاختيار من القائمة'
    dv_dept.errorTitle = 'خطأ في الإدخال'
    dv_dept.showErrorMessage = True
    ws.add_data_validation(dv_dept)
    dv_dept.add('D3:D1000')

    # Helper column for forward filling course code (column Z)
    ws.column_dimensions['Z'].hidden = True
    
    for row_idx in range(3, 201):
        ws.cell(row=row_idx, column=26, value=f'=IF(A{row_idx}<>"",A{row_idx},Z{row_idx-1})')
        # الساعات المعتمدة للأقسام = مجموع توزيع الساعات المعتمدة (نظري + عملي + أنشطة)
        ws.cell(row=row_idx, column=6, value=f'=IF(SUM(G{row_idx}:I{row_idx})=0,"",SUM(G{row_idx}:I{row_idx}))')
        ws.cell(row=row_idx, column=10, value=f'=IF(G{row_idx}="","",G{row_idx}*15)')
        ws.cell(row=row_idx, column=11, value=f'=IF(H{row_idx}="","",H{row_idx}*30)')
        ws.cell(row=row_idx, column=12, value=f'=IF(I{row_idx}="","",I{row_idx}*60)')
        ws.cell(row=row_idx, column=16, value=f'=IF(SUM(M{row_idx}:O{row_idx})=0,"",SUM(M{row_idx}:O{row_idx}))')
        ws.cell(row=row_idx, column=17, value=f'=IF(A{row_idx}="","",SUMIF($Z$3:$Z$200, A{row_idx}, $P$3:$P$200))')
        ws.cell(row=row_idx, column=18, value=f'=IF(A{row_idx}="","",SUMIF($Z$3:$Z$200, A{row_idx}, $F$3:$F$200))')

    temp_file = NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    wb.close()
    
    return FileResponse(
        temp_file.name, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
        filename="نموذج_استيراد_المقررات_الدراسية_طب.xlsx"
    )

@app.post("/api/courses/import")
def import_courses_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    import pandas as pd
    try:
        # مسح جميع المقررات القديمة قبل استيراد المقررات الجديدة
        db.query(models.Course).delete()
        db.commit()

        contents = file.file.read()
        df = pd.read_excel(io.BytesIO(contents), dtype=str)
        
        # Deduplicate column names to prevent pandas from returning a Series for duplicated columns
        new_cols = []
        seen = {}
        for c in df.columns:
            c_str = str(c)
            if c_str in seen:
                seen[c_str] += 1
                new_cols.append(f"{c_str}_{seen[c_str]}")
            else:
                seen[c_str] = 0
                new_cols.append(c_str)
        df.columns = new_cols
        
        def clean_col(c):
            val = str(c).strip().replace(" ", "").replace("_", "").lower()
            val = val.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
            return val
            
        col_map = {}
        for col in df.columns:
            cleaned = clean_col(col)
            if "كود" in cleaned or "code" in cleaned:
                col_map["code"] = col
            elif "عربي" in cleaned or "ar" in cleaned or "الاسم" in cleaned:
                if "انجليزي" not in cleaned and "en" not in cleaned:
                    col_map["name_ar"] = col
            elif "انجليزي" in cleaned or "en" in cleaned or "english" in cleaned:
                col_map["name_en"] = col
            elif "كلية" in cleaned or "faculty" in cleaned:
                col_map["faculty"] = col
            elif "برنامج" in cleaned or "program" in cleaned:
                col_map["program"] = col
            elif "مستوى" in cleaned or "level" in cleaned:
                col_map["level"] = col
            elif "فصل" in cleaned or "ترم" in cleaned or "semester" in cleaned:
                col_map["semester"] = col
            elif "نظري" in cleaned and "ساعات" in cleaned:
                col_map["theory_hours"] = col
            elif "عملي" in cleaned and "ساعات" in cleaned:
                col_map["practical_hours"] = col
            elif "تمارين" in cleaned and "ساعات" in cleaned:
                col_map["exercise_hours"] = col
            elif "انشطة" in cleaned and "ساعات" in cleaned:
                col_map["activity_hours"] = col
            elif "درجة" in cleaned and "كلية" in cleaned:
                col_map["total_grade"] = col
            elif "درجة" in cleaned and "نظري" in cleaned:
                col_map["theory_grade"] = col
            elif "درجة" in cleaned and "عملي" in cleaned:
                col_map["practical_grade"] = col
            elif "اعمال" in cleaned or "سنة" in cleaned:
                col_map["year_work_grade"] = col
            elif "زمن" in cleaned or "امتحان" in cleaned:
                col_map["exam_time_hours"] = col
            elif "وصف" in cleaned or "description" in cleaned:
                col_map["description"] = col

        req = ["code", "name_ar", "faculty", "program", "level"]
        missing = [r for r in req if r not in col_map]
        if missing:
            raise HTTPException(status_code=400, detail=f"الملف لا يحتوي على الأعمدة المطلوبة: {', '.join(missing)}")

        imported = 0
        updated = 0

        for _, row in df.iterrows():
            code_val = str(row[col_map["code"]]).strip() if pd.notna(row[col_map["code"]]) else ""
            name_ar_val = str(row[col_map["name_ar"]]).strip() if pd.notna(row[col_map["name_ar"]]) else ""
            if not code_val or not name_ar_val:
                continue

            name_en_val = str(row[col_map["name_en"]]).strip() if "name_en" in col_map and pd.notna(row[col_map["name_en"]]) else ""
            fac_name = str(row[col_map["faculty"]]).strip() if pd.notna(row[col_map["faculty"]]) else ""
            prog_name = str(row[col_map["program"]]).strip() if pd.notna(row[col_map["program"]]) else ""
            if prog_name:
                prog_name = normalize_program_name(prog_name)
            lvl_val = str(row[col_map["level"]]).strip() if pd.notna(row[col_map["level"]]) else ""
            sem_val = str(row[col_map["semester"]]).strip() if "semester" in col_map and pd.notna(row[col_map["semester"]]) else "الفصل الدراسي الأول"

            level_int = 1
            try:
                digits = "".join([c for c in lvl_val if c.isdigit()])
                if digits:
                    level_int = int(digits)
                else:
                    norm = lvl_val.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
                    if "اول" in norm:
                        level_int = 1
                    elif "ثاني" in norm or "ثانى" in norm:
                        level_int = 2
                    elif "ثالث" in norm:
                        level_int = 3
                    elif "رابع" in norm:
                        level_int = 4
                    elif "خامس" in norm:
                        level_int = 5
                    elif "عام" in norm:
                        level_int = 0
            except:
                pass

            target_faculty = db.query(models.Faculty).filter(models.Faculty.name == fac_name).first()
            if not target_faculty:
                target_faculty = models.Faculty(name=fac_name)
                db.add(target_faculty)
                db.flush()

            target_program = db.query(models.Program).filter(
                models.Program.name == prog_name,
                models.Program.faculty_id == target_faculty.id
            ).first()
            if not target_program:
                target_program = models.Program(name=prog_name, faculty_id=target_faculty.id)
                db.add(target_program)
                db.flush()

            theory_hours = float(row[col_map["theory_hours"]]) if "theory_hours" in col_map and pd.notna(row[col_map["theory_hours"]]) else 0.0
            practical_hours = float(row[col_map["practical_hours"]]) if "practical_hours" in col_map and pd.notna(row[col_map["practical_hours"]]) else 0.0
            exercise_hours = float(row[col_map["exercise_hours"]]) if "exercise_hours" in col_map and pd.notna(row[col_map["exercise_hours"]]) else 0.0
            activity_hours = float(row[col_map["activity_hours"]]) if "activity_hours" in col_map and pd.notna(row[col_map["activity_hours"]]) else 0.0
            
            total_grade = float(row[col_map["total_grade"]]) if "total_grade" in col_map and pd.notna(row[col_map["total_grade"]]) else 0.0
            theory_grade = float(row[col_map["theory_grade"]]) if "theory_grade" in col_map and pd.notna(row[col_map["theory_grade"]]) else 0.0
            practical_grade = float(row[col_map["practical_grade"]]) if "practical_grade" in col_map and pd.notna(row[col_map["practical_grade"]]) else 0.0
            year_work_grade = float(row[col_map["year_work_grade"]]) if "year_work_grade" in col_map and pd.notna(row[col_map["year_work_grade"]]) else 0.0
            exam_time_hours = float(row[col_map["exam_time_hours"]]) if "exam_time_hours" in col_map and pd.notna(row[col_map["exam_time_hours"]]) else 0.0
            description_val = str(row[col_map["description"]]).strip() if "description" in col_map and pd.notna(row[col_map["description"]]) else None

            db_course = db.query(models.Course).filter(
                models.Course.code == code_val,
                models.Course.program_id == target_program.id
            ).first()

            if db_course:
                db_course.name_ar = name_ar_val
                db_course.name_en = name_en_val
                db_course.level = level_int
                db_course.semester = sem_val
                db_course.theory_hours = theory_hours
                db_course.practical_hours = practical_hours
                db_course.exercise_hours = exercise_hours
                db_course.activity_hours = activity_hours
                db_course.total_grade = total_grade
                db_course.theory_grade = theory_grade
                db_course.practical_grade = practical_grade
                db_course.year_work_grade = year_work_grade
                db_course.exam_time_hours = exam_time_hours
                db_course.description = description_val
                updated += 1
            else:
                db_course = models.Course(
                    code=code_val,
                    name_ar=name_ar_val,
                    name_en=name_en_val,
                    level=level_int,
                    semester=sem_val,
                    theory_hours=theory_hours,
                    practical_hours=practical_hours,
                    exercise_hours=exercise_hours,
                    activity_hours=activity_hours,
                    total_grade=total_grade,
                    theory_grade=theory_grade,
                    practical_grade=practical_grade,
                    year_work_grade=year_work_grade,
                    exam_time_hours=exam_time_hours,
                    faculty_id=target_faculty.id,
                    program_id=target_program.id,
                    description=description_val
                )
                db.add(db_course)
                imported += 1

        db.commit()
        return {
            "success": True,
            "message": "تمت عملية استيراد المقررات بنجاح",
            "imported": imported,
            "updated": updated,
            "total": imported + updated
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"حدث خطأ أثناء معالجة ملف الاستيراد: {str(e)}")

@app.post("/api/courses/import-all")
async def import_all_courses(file: UploadFile = File(...), db: Session = Depends(get_db)):
    result = await process_and_replace_courses(file, db)
    return result

@app.get("/api/courses/fail_rates")
def get_fail_rates(db: Session = Depends(get_db)):
    db_fail_rates = db.query(models.Course.fail_rate).distinct().all()
    fail_rates_list = sorted(list(set([r[0] for r in db_fail_rates if r[0] and r[0].strip() and r[0] != "-"])))
    if not fail_rates_list:
        fail_rates_list = [
            "نهاية الفصل:0",
            "نهاية الفصل:12",
            "نهاية الفصل:13.5",
            "نهاية الفصل:18",
            "نهاية الفصل:24",
            "نهاية الفصل:30",
            "نهاية الفصل:36",
            "نهاية الفصل:40",
            "نهاية الفصل:50",
            "منتصف الفصل:20"
        ]
    return fail_rates_list

@app.put("/api/courses/{id}", response_model=schemas.CourseOut)
def update_course(id: int, course: schemas.CourseCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    db_course = db.query(models.Course).filter(models.Course.id == id).first()
    if not db_course:
        raise HTTPException(status_code=404, detail="المقرر غير موجود")
    
    course_data = course.dict(exclude={"modules"})
    modules_data = course.dict().get("modules", [])
    
    for key, val in course_data.items():
        setattr(db_course, key, val)
        
    # Update modules safely without violating foreign key constraints
    existing_modules = db.query(models.CourseModule).filter(models.CourseModule.course_id == id).all()
    existing_map = {m.department_name.strip().lower(): m for m in existing_modules}
    
    if db_course.is_bundle and modules_data:
        kept_module_ids = set()
        for mod in modules_data:
            dept_key = mod.get("department_name", "").strip().lower()
            if dept_key in existing_map:
                # Update existing module
                m_obj = existing_map[dept_key]
                for k, v in mod.items():
                    setattr(m_obj, k, v)
                kept_module_ids.add(m_obj.id)
            else:
                # Add new module
                m_obj = models.CourseModule(course_id=id, **mod)
                db.add(m_obj)
                db.flush()
                kept_module_ids.add(m_obj.id)
                
        # Delete modules that were removed by the user
        for old_m in existing_modules:
            if old_m.id not in kept_module_ids:
                db.query(models.StudyPlanItem).filter(models.StudyPlanItem.module_id == old_m.id).update(
                    {"module_id": None}, synchronize_session=False
                )
                db.delete(old_m)
    elif not db_course.is_bundle:
        for old_m in existing_modules:
            db.query(models.StudyPlanItem).filter(models.StudyPlanItem.module_id == old_m.id).update(
                {"module_id": None}, synchronize_session=False
            )
            db.delete(old_m)

    user_role_str = get_user_role_display(current_user)
    course_title = db_course.name_ar or db_course.name_en or db_course.code or "مقرر"
    action_text = f"قام بتعديل المقرر: {course_title}"
    create_notification(db, db_course.faculty_id, f"{current_user.username} ({user_role_str})", action_text)
    
    db.commit()
    db.refresh(db_course)
    return db_course


@app.delete("/api/courses/{id}")
def delete_course(id: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    db_course = db.query(models.Course).filter(models.Course.id == id).first()
    if not db_course:
        raise HTTPException(status_code=404, detail="المقرر غير موجود")
        
    user_role_str = get_user_role_display(current_user)
    action_text = f"قام بحذف المقرر: {db_course.name_ar}"
    create_notification(db, db_course.faculty_id, f"{current_user.username} ({user_role_str})", action_text)
    
    db_course.is_deleted = True
    db_course.deleted_at = datetime.utcnow()
    db.commit()
    return {"message": "تم حذف المقرر بنجاح"}

@app.delete("/api/professors/{id}")
def delete_professor(id: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    prof = db.query(models.Professor).filter(models.Professor.id == id).first()
    if not prof:
        raise HTTPException(status_code=404, detail="غير موجود")
        
    user_role_str = get_user_role_display(current_user)
    action_text = f"قام بحذف عضو هيئة تدريس: {prof.name_ar}"
    if prof.faculties:
        for fac in prof.faculties:
            create_notification(db, fac.id, f"{current_user.username} ({user_role_str})", action_text)
    else:
        create_notification(db, None, f"{current_user.username} ({user_role_str})", action_text)
    
    prof.is_deleted = True
    prof.deleted_at = datetime.utcnow()
    db.commit()
    return {"message": "تم الحذف"}

@app.delete("/api/professors/{id}/assignments/{assignment_id}")
def delete_professor_assignment(id: int, assignment_id: int, db: Session = Depends(get_db)):
    """حذف مقرر مكلف به دكتور (إفراغ الحقل professor_id في خطة الدراسة)"""
    item = db.query(models.StudyPlanItem).filter(
        models.StudyPlanItem.id == assignment_id,
        models.StudyPlanItem.professor_id == id
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="التكليف غير موجود")
    item.professor_id = None
    db.commit()
    return {"message": "تم حذف التكليف بنجاح"}

from pydantic import BaseModel

class ExportAssignmentsRequest(BaseModel):
    professor_ids: List[int]
    academic_year: str
    semester: Optional[str] = None
    level: Optional[str] = None

@app.post("/api/professors/export-assignments")
def export_professors_assignments(request: ExportAssignmentsRequest, db: Session = Depends(get_db)):
    from collections import defaultdict
    
    query = db.query(models.StudyPlanItem).join(models.StudyPlan).outerjoin(models.Course, models.StudyPlanItem.course_id == models.Course.id).filter(
        models.StudyPlanItem.professor_id.in_(request.professor_ids),
        models.StudyPlan.is_deleted == False
    )
    if request.academic_year and request.academic_year != "الكل":
        query = query.filter(models.StudyPlan.academic_year == request.academic_year)
    
    if request.semester and request.semester != "الكل":
        if request.semester == "أول" or "اول" in request.semester:
            query = query.filter(or_(
                models.StudyPlan.semester.ilike("%أول%"),
                models.StudyPlan.semester.ilike("%اول%"),
                models.StudyPlan.semester.ilike("%1%"),
                models.Course.semester.ilike("%أول%"),
                models.Course.semester.ilike("%اول%"),
                models.Course.semester.ilike("%1%")
            ))
        elif request.semester == "ثاني" or "ثان" in request.semester:
            query = query.filter(or_(
                models.StudyPlan.semester.ilike("%ثان%"),
                models.StudyPlan.semester.ilike("%2%"),
                models.Course.semester.ilike("%ثان%"),
                models.Course.semester.ilike("%2%")
            ))
        elif request.semester == "صيفي" or "صيف" in request.semester:
            query = query.filter(or_(
                models.StudyPlan.semester.ilike("%صيف%"),
                models.StudyPlan.semester.ilike("%3%"),
                models.Course.semester.ilike("%صيف%"),
                models.Course.semester.ilike("%3%")
            ))
            
    if request.level and request.level != "الكل":
        lvl_str = str(request.level).strip()
        lvl_num = None
        if lvl_str == "عام" or lvl_str == "0":
            lvl_num = 0
        elif "أول" in lvl_str or "اول" in lvl_str or lvl_str == "1":
            lvl_num = 1
        elif "ثان" in lvl_str or lvl_str == "2":
            lvl_num = 2
        elif "ثالث" in lvl_str or lvl_str == "3":
            lvl_num = 3
        elif "رابع" in lvl_str or lvl_str == "4":
            lvl_num = 4
        elif "خامس" in lvl_str or lvl_str == "5":
            lvl_num = 5

        if lvl_num is not None:
            query = query.filter(or_(
                models.StudyPlanItem.level == lvl_num,
                models.Course.level == lvl_num,
                models.StudyPlanItem.level.ilike(f"%{lvl_str}%")
            ))
        else:
            query = query.filter(or_(
                models.StudyPlanItem.level.ilike(f"%{lvl_str}%"),
                models.Course.level.ilike(f"%{lvl_str}%")
            ))
            
    items = query.all()
    
    assignments_by_prof = defaultdict(list)
    for item in items:
        sp = item.study_plan
        faculty = sp.faculty
        course = item.course or item.base_course or (item.module.course if item.module else None)
        total_hours = round(
            (item.hours_actual_theory or 0.0) +
            (item.hours_actual_practical or 0.0) +
            (item.hours_actual_exercise or 0.0) +
            (item.hours_actual_activity or 0.0),
            2
        )
        
        course_names = []
        base_course_name = ""
        if course:
            base_course_name = (course.name_ar or course.name_en or course.code or "").strip()
            if base_course_name and item.module and item.module.department_name:
                m_dept = item.module.department_name.strip()
                if m_dept and m_dept.lower() not in base_course_name.lower():
                    base_course_name = f"{base_course_name} ({m_dept})"
            elif not base_course_name and item.module and item.module.department_name:
                base_course_name = item.module.department_name.strip()
        elif item.module and item.module.department_name:
            base_course_name = item.module.department_name.strip()

        if base_course_name:
            course_names.append(base_course_name)

        program_names = []
        if item.entry_group_id:
            raw_entry = str(item.entry_group_id).strip()
            if raw_entry.startswith("{"):
                try:
                    import json
                    parsed = json.loads(raw_entry)
                    pids = parsed.get("prog_ids") or []
                    for pid in pids:
                        prg = db.query(models.Program).filter(models.Program.id == int(pid)).first()
                        if prg and prg.name and prg.name not in program_names:
                            program_names.append(prg.name)
                    sc_list = parsed.get("shared_courses") or []
                    for sc in sc_list:
                        sc_name = (sc.get("name_ar") or sc.get("name_en") or "").strip()
                        if sc_name and sc_name not in course_names:
                            course_names.append(sc_name)
                except:
                    pass
            elif "," in raw_entry:
                try:
                    pids = [int(x.strip()) for x in raw_entry.split(",") if x.strip()]
                    for pid in pids:
                        prg = db.query(models.Program).filter(models.Program.id == pid).first()
                        if prg and prg.name and prg.name not in program_names:
                            program_names.append(prg.name)
                except:
                    pass

        if not program_names:
            if item.program and item.program.name:
                program_names.append(item.program.name)
            elif course and course.program and course.program.name:
                program_names.append(course.program.name)

        final_prog_name = " - ".join(program_names) if program_names else "-"
        final_course_name = " - ".join(course_names) if course_names else (course.name_ar or (course.name_en if course else ""))

        assignments_by_prof[item.professor_id].append({
            "faculty_name": faculty.name if faculty else "",
            "program_name": final_prog_name,
            "course_name": final_course_name,
            "level": item.level if item.level is not None else (course.level if course else ""),
            "course_semester": sp.semester if sp else "",
            "academic_year": sp.academic_year if sp else "",
            "hours": total_hours
        })
    return assignments_by_prof

# ==========================================
# مسارات الخطة الدراسية (StudyPlan APIs)
# ==========================================

def get_semester_filter(semester_str: str):
    if not semester_str:
        return None
    if "الصيفي" in semester_str:
        return models.StudyPlan.semester.like("%الصيفي%")
    elif "الأول" in semester_str:
        return models.StudyPlan.semester.like("%الأول%")
    elif "الثاني" in semester_str:
        return models.StudyPlan.semester.like("%الثاني%")
    return models.StudyPlan.semester == semester_str

@app.get("/api/study-plans", response_model=list[schemas.StudyPlanOut])
def get_study_plans(
    faculty_id: Optional[int] = None,
    semester: Optional[str] = None,
    year: Optional[str] = None,
    academic_year: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    query = db.query(models.StudyPlan).filter(models.StudyPlan.is_deleted == False)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            query = query.filter(models.StudyPlan.faculty_id == faculty_id)
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id is not None:
            if faculty_id in faculty_ids:
                query = query.filter(models.StudyPlan.faculty_id == faculty_id)
            else:
                return []
        else:
            query = query.filter(models.StudyPlan.faculty_id.in_(faculty_ids))
    elif current_user.faculty_id:
        query = query.filter(models.StudyPlan.faculty_id == current_user.faculty_id)
    elif faculty_id:
        query = query.filter(models.StudyPlan.faculty_id == faculty_id)
    
    if semester:
        query = query.filter(get_semester_filter(semester))
    target_year = year or academic_year
    if target_year:
        query = query.filter(models.StudyPlan.academic_year == target_year)
    return query.all()

@app.post("/api/study-plans", response_model=schemas.StudyPlanOut)
def create_study_plan(plan: schemas.StudyPlanCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    try:
        existing_plans = db.query(models.StudyPlan).filter(
            models.StudyPlan.faculty_id == plan.faculty_id,
            get_semester_filter(plan.semester),
            models.StudyPlan.academic_year == plan.academic_year,
            models.StudyPlan.is_deleted == False
        ).all()
        
        saved_state = {
            "is_finished": False, "finished_by": None,
            "is_reviewed_1": False, "reviewed_1_by": None,
            "is_reviewed_2": False, "reviewed_2_by": None,
            "is_approved": False, "approved_by": None
        }
        
        for ep in existing_plans:
            if ep.is_finished:
                saved_state["is_finished"] = ep.is_finished
                saved_state["finished_by"] = ep.finished_by
            if ep.is_reviewed_1:
                saved_state["is_reviewed_1"] = ep.is_reviewed_1
                saved_state["reviewed_1_by"] = ep.reviewed_1_by
            if ep.is_reviewed_2:
                saved_state["is_reviewed_2"] = ep.is_reviewed_2
                saved_state["reviewed_2_by"] = ep.reviewed_2_by
            if ep.is_approved:
                saved_state["is_approved"] = ep.is_approved
                saved_state["approved_by"] = ep.approved_by
            db.delete(ep)
        db.commit()
    
        # حساب مجموع الساعات
        total_theory = 0
        total_practical = 0
        total_activity = 0.0
    
        # إنشاء سجل الخطة الدراسية الرئيسي
        db_plan = models.StudyPlan(
            faculty_id=plan.faculty_id,
            semester=plan.semester,
            academic_year=plan.academic_year,
            total_theory_hours=plan.total_theory_hours,
            total_practical_hours=plan.total_practical_hours,
            **saved_state
        )
        db.add(db_plan)
        db.flush()
    
        # إضافة بنود المقررات
        for item_data in plan.items:
            course = db.query(models.Course).filter(models.Course.id == item_data.course_id).first()
            if not course:
                continue
        
            # التحقق من وجود الأستاذ
            if item_data.professor_id:
                prof = db.query(models.Professor).filter(models.Professor.id == item_data.professor_id).first()
                if not prof:
                    raise HTTPException(status_code=404, detail="الأستاذ غير موجود")
                
                # أضف الكلية للأستاذ إذا لم تكن موجودة ضمن كلياته
                faculty = db.query(models.Faculty).filter(models.Faculty.id == plan.faculty_id).first()
                if faculty and faculty not in prof.faculties:
                    prof.faculties.append(faculty)
        
            # حساب الساعات المطلوبة
            req_theory = (course.theory_hours or 0) * (item_data.groups_theory or 0)
            req_practical = ((course.practical_hours or 0) + (course.exercise_hours or 0)) * (item_data.groups_practical or 0)
            req_exercise = (course.exercise_hours or 0) * (item_data.groups_exercise or 0)
            req_activity = (course.activity_hours or 0) * (item_data.groups_activity or 0)
        
            db_item = models.StudyPlanItem(
                study_plan_id=db_plan.id,
                base_course_id=item_data.base_course_id or item_data.course_id,
                course_id=item_data.course_id,
                module_id=item_data.module_id,
                program_id=item_data.program_id,
                professor_id=item_data.professor_id,
                entry_group_id=item_data.entry_group_id,
                level=item_data.level,
                student_count=item_data.student_count,
                groups_theory=item_data.groups_theory,
                groups_practical=item_data.groups_practical,
                groups_exercise=item_data.groups_exercise,
                groups_activity=item_data.groups_activity,
                hours_actual_theory=item_data.hours_actual_theory,
                hours_actual_practical=item_data.hours_actual_practical,
                hours_actual_exercise=item_data.hours_actual_exercise,
                hours_actual_activity=item_data.hours_actual_activity,
                required_hours_theory=req_theory,
                required_hours_practical=req_practical,
                required_hours_exercise=item_data.required_hours_exercise or req_exercise,
                required_hours_activity=item_data.required_hours_activity or req_activity,
                notes=item_data.notes,
                course_notes=item_data.course_notes,
            )
            db.add(db_item)
            total_theory += req_theory
            total_practical += req_practical
            total_activity += (item_data.required_hours_activity or req_activity)
    
        # تحديث المجاميع
        db_plan.total_theory_hours = total_theory
        db_plan.total_practical_hours = total_practical
        db_plan.total_activity_hours = total_activity
    
        # إضافة إشعار
        if current_user.role == 'admin':
            user_role_str = 'مدير عام'
        elif current_user.role == 'faculty_professor':
            user_role_str = 'عضو هيئة تدريس'
        else:
            user_role_str = 'مسؤول كلية'
        action_text = f"قام بتعديل الخطة الدراسية ({plan.semester} - العام الجامعي: {plan.academic_year})"
        create_notification(db, plan.faculty_id, f"{current_user.username} ({user_role_str})", action_text, academic_year=plan.academic_year, semester=plan.semester)
        db.commit()

        return db_plan

    except Exception as e:
        db.rollback()
        import traceback
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail=traceback.format_exc())

@app.post("/api/study-plans/status")
def update_study_plan_status(
    status_data: schemas.StudyPlanStatusUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    plan = db.query(models.StudyPlan).filter(
        models.StudyPlan.faculty_id == status_data.faculty_id,
        get_semester_filter(status_data.semester),
        models.StudyPlan.academic_year == status_data.academic_year,
        models.StudyPlan.is_deleted == False
    ).first()

    if not plan:
        raise HTTPException(status_code=404, detail="الخطة الدراسية غير موجودة، يرجى حفظ الخطة أولاً.")

    role_ar_map = {
        models.UserRole.admin: "مدير عام",
        models.UserRole.manager: "مدير",
        models.UserRole.faculty_admin: "مسؤول كلية",
        models.UserRole.faculty_professor: "مدير برنامج",
        models.UserRole.student_affairs: "مدير شؤون الطلاب",
    }
    role_str = current_user.job_title if current_user.job_title else role_ar_map.get(current_user.role, "مسؤول")
    user_full_name = f"{role_str} ({current_user.username.split('@')[0]})".strip()
    faculty_name = plan.faculty.name if plan.faculty else "الكلية"

    action_text = ""

    is_super_admin = current_user.role in [models.UserRole.admin]
    is_manager = (
        current_user.role == models.UserRole.manager
        or (current_user.job_title and ("مدير" in current_user.job_title and "برنامج" not in current_user.job_title and "شؤون" not in current_user.job_title))
        or (current_user.job_title and "عميد" in current_user.job_title)
    )
    is_program_director = (
        current_user.role == models.UserRole.faculty_professor
        or (current_user.job_title and "مدير برنامج" in current_user.job_title)
    )

    is_matching_faculty = True
    if is_program_director and not is_super_admin and not is_manager:
        if current_user.faculty_id and current_user.faculty_id != plan.faculty_id:
            assigned_ids = [f.id for f in current_user.assigned_faculties] if current_user.assigned_faculties else []
            if plan.faculty_id not in assigned_ids:
                is_matching_faculty = False

    is_finish_allowed = (
        is_super_admin
        or is_manager
        or (is_program_director and is_matching_faculty)
        or bool(current_user.perm_finish_plan)
    )

    if status_data.action == "finish":
        if not is_finish_allowed:
            if is_program_director and not is_matching_faculty:
                raise HTTPException(status_code=403, detail="لا تملك صلاحية إنهاء الخطة لهذه الكلية (خاص بمدير البرنامج التابع لهذه الكلية).")
            raise HTTPException(status_code=403, detail="لا تملك صلاحية إنهاء الخطة (الصلاحية للمدير العام والمدير ومدير البرنامج الخاص بالكلية).")
        plan.is_finished = True
        plan.finished_by = user_full_name
        action_text = f"قام بإنهاء الخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"

    elif status_data.action == "cancel_finish":
        if not is_finish_allowed:
            if is_program_director and not is_matching_faculty:
                raise HTTPException(status_code=403, detail="لا تملك صلاحية إلغاء إنهاء الخطة لهذه الكلية (خاص بمدير البرنامج التابع لهذه الكلية).")
            raise HTTPException(status_code=403, detail="لا تملك صلاحية إلغاء إنهاء الخطة (الصلاحية للمدير العام والمدير ومدير البرنامج الخاص بالكلية).")
        if plan.is_approved:
            raise HTTPException(status_code=400, detail="لا يمكن إلغاء إنهاء الخطة بعد اعتمادها. يمكن ذلك فقط بعد إلغاء اعتماد الخطة بواسطة المسؤول.")
        plan.is_finished = False
        plan.finished_by = None
        plan.is_reviewed_1 = False # Cascading cancel
        plan.reviewed_1_by = None
        plan.is_reviewed_2 = False
        plan.reviewed_2_by = None
        plan.is_approved = False
        plan.approved_by = None
        action_text = f"تم الغاء إنهاء الخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"

    elif status_data.action == "review1":
        if not current_user.perm_review_1 and current_user.role != models.UserRole.admin:
            raise HTTPException(status_code=403, detail="لا تملك صلاحية المراجعة الأولى.")
        plan.is_finished = True
        plan.is_reviewed_1 = True
        plan.reviewed_1_by = user_full_name
        action_text = f"قام بإجراء المراجعة الأولى للخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"

    elif status_data.action == "cancel_review1":
        if not current_user.perm_review_1 and current_user.role != models.UserRole.admin:
            raise HTTPException(status_code=403, detail="لا تملك صلاحية إلغاء المراجعة الأولى.")
        plan.is_finished = False # إلغاء الإنهاء لفتح التعديل وإظهار الأزرار ثانية
        plan.finished_by = None
        plan.is_reviewed_1 = False
        plan.reviewed_1_by = None
        plan.is_reviewed_2 = False # Cascading cancel
        plan.reviewed_2_by = None
        plan.is_approved = False
        plan.approved_by = None
        action_text = f"تم الغاء المراجعة الاولى للخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"

    elif status_data.action == "review2":
        if not current_user.perm_review_2 and current_user.role != models.UserRole.admin:
            raise HTTPException(status_code=403, detail="لا تملك صلاحية المراجعة الثانية.")
        plan.is_finished = True
        plan.is_reviewed_1 = True
        plan.is_reviewed_2 = True
        plan.reviewed_2_by = user_full_name
        action_text = f"قام بإجراء المراجعة الثانية للخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"

    elif status_data.action == "cancel_review2":
        if not current_user.perm_review_2 and current_user.role != models.UserRole.admin:
            raise HTTPException(status_code=403, detail="لا تملك صلاحية إلغاء المراجعة الثانية.")
        plan.is_finished = False # إلغاء الإنهاء لفتح الخطة بالكامل وإظهار جميع الأزرار وعمود الإجراءات مجدداً
        plan.finished_by = None
        plan.is_reviewed_1 = False
        plan.reviewed_1_by = None
        plan.is_reviewed_2 = False
        plan.reviewed_2_by = None
        plan.is_approved = False
        plan.approved_by = None
        action_text = f"تم الغاء المراجعة الثانية للخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"

    elif status_data.action == "approve":
        if not current_user.perm_approve_plan and current_user.role != models.UserRole.admin:
            raise HTTPException(status_code=403, detail="لا تملك صلاحية اعتماد الخطة.")
            
        import json
        signatures = db.query(models.Signature).filter(
            models.Signature.faculty_id == plan.faculty_id,
            models.Signature.report_type.contains("الخطة الدراسية"),
            models.Signature.is_deleted == False
        ).order_by(models.Signature.order_index.asc()).all()
        frozen = []
        for s in signatures:
            frozen.append({
                "id": s.id,
                "signature_title": s.signature_title,
                "official_name": s.official_name,
                "order_index": s.order_index
            })
        plan.frozen_signatures = json.dumps(frozen, ensure_ascii=False)
        
        plan.is_finished = True
        plan.is_reviewed_1 = True
        plan.is_reviewed_2 = True
        plan.is_approved = True
        plan.approved_by = user_full_name
        action_text = f"تم اعتماد الخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"

    elif status_data.action == "cancel_approve":
        if not current_user.perm_approve_plan and current_user.role != models.UserRole.admin and plan.approved_by != user_full_name:
            raise HTTPException(status_code=403, detail="لا تملك صلاحية إلغاء اعتماد الخطة.")
        plan.is_approved = False
        plan.approved_by = None
        plan.frozen_signatures = None
        plan.is_reviewed_2 = False
        plan.reviewed_2_by = None
        plan.is_reviewed_1 = False
        plan.reviewed_1_by = None
        plan.is_finished = False # إلغاء الإنهاء لفتح الخطة بالكامل وإظهار جميع الأزرار وعمود الإجراءات مجدداً
        plan.finished_by = None
        action_text = f"تم الغاء اعتماد الخطة الدراسية ل{faculty_name} - {plan.semester} - العام الجامعي {plan.academic_year}"
    
    else:
        raise HTTPException(status_code=400, detail="إجراء غير معروف.")

    if action_text:
        create_notification(db, plan.faculty_id, f"{current_user.username} ({role_str})", action_text, academic_year=plan.academic_year, semester=plan.semester)

    db.commit()
    return {"message": "تم تحديث حالة الخطة بنجاح"}

@app.get("/api/study-plans/download-template")
def download_study_plan_template(faculty_id: Optional[int] = None, db: Session = Depends(get_db)):
    wb = Workbook()
    
    ws = wb.active
    ws.title = "نموذج الخطة الدراسية"
    ws.views.sheetView[0].rightToLeft = True
    
    header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    header_font = Font(name="Cairo", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "اسم المقرر الرئيسي",
        "البرنامج / القسم",
        "كود المقرر",
        "عدد الطلاب",
        "عدد مجموعات نظري",
        "عدد مجموعات عملي",
        "اسم عضو هيئة التدريس",
        "المستوى",
        "ساعات منفذة (نظري)",
        "ساعات منفذة (عملي)",
        "ملاحظات"
    ]
    
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    ws_lists = wb.create_sheet(title="قوائم_المراجع")
    ws_lists.views.sheetView[0].rightToLeft = True
    
    courses = []
    if faculty_id:
        courses = db.query(models.Course).filter(models.Course.faculty_id == faculty_id).all()
    else:
        courses = db.query(models.Course).all()
        
    unique_course_names = sorted(list(set(c.name_ar for c in courses if c.name_ar)))
    unique_codes = sorted(list(set(c.code for c in courses if c.code)))
    
    if faculty_id:
        programs = db.query(models.Program).filter(models.Program.faculty_id == faculty_id).all()
    else:
        programs = db.query(models.Program).all()
    program_names = sorted([p.name for p in programs if p.name])
    
    professors = db.query(models.Professor).all()
    prof_names = sorted([p.name_ar for p in professors if p.name_ar])
    
    levels = ["تمهيدي", "الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس"]
    
    for idx, name in enumerate(unique_course_names, 1): ws_lists.cell(row=idx, column=1, value=name)
    for idx, name in enumerate(program_names, 1): ws_lists.cell(row=idx, column=2, value=name)
    for idx, code in enumerate(unique_codes, 1): ws_lists.cell(row=idx, column=3, value=code)
    for idx, name in enumerate(prof_names, 1): ws_lists.cell(row=idx, column=4, value=name)
    for idx, lvl in enumerate(levels, 1): ws_lists.cell(row=idx, column=5, value=lvl)
    
    if unique_course_names:
        dv_course = DataValidation(type="list", formula1=f"=قوائم_المراجع!$A$1:$A${len(unique_course_names)}", allow_blank=True)
        ws.add_data_validation(dv_course)
        dv_course.add("A2:A200")
        
    if program_names:
        dv_prog = DataValidation(type="list", formula1=f"=قوائم_المراجع!$B$1:$B${len(program_names)}", allow_blank=True)
        ws.add_data_validation(dv_prog)
        dv_prog.add("B2:B200")
        
    if unique_codes:
        dv_code = DataValidation(type="list", formula1=f"=قوائم_المراجع!$C$1:$C${len(unique_codes)}", allow_blank=True)
        ws.add_data_validation(dv_code)
        dv_code.add("C2:C200")
        
    if prof_names:
        dv_prof = DataValidation(type="list", formula1=f"=قوائم_المراجع!$D$1:$D${len(prof_names)}", allow_blank=True)
        ws.add_data_validation(dv_prof)
        dv_prof.add("G2:G200")
        
    dv_lvl = DataValidation(type="list", formula1=f"=قوائم_المراجع!$E$1:$E${len(levels)}", allow_blank=True)
    ws.add_data_validation(dv_lvl)
    dv_lvl.add("H2:H200")
    
    dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv_num.error = "يجب إدخال رقم موجب (يُسمح بالأرقام العشرية والكسور مثل 1.5)"
    dv_num.errorTitle = "إدخال غير صحيح"
    ws.add_data_validation(dv_num)
    dv_num.add("D2:F200")
    dv_num.add("I2:J200")

    column_widths = {'A': 32, 'B': 25, 'C': 15, 'D': 14, 'E': 16, 'F': 16, 'G': 28, 'H': 12, 'I': 18, 'J': 18, 'K': 25}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=study_plan_template.xlsx"}
    )

@app.get("/api/study-plans/{id}", response_model=schemas.StudyPlanOut)
def get_study_plan(id: int, db: Session = Depends(get_db)):
    plan = db.query(models.StudyPlan).filter(models.StudyPlan.id == id).first()
    if not plan:
        raise HTTPException(status_code=404, detail="الخطة الدراسية غير موجودة")
    return plan

@app.delete("/api/study-plans/{id}")
def delete_study_plan(id: int, db: Session = Depends(get_db)):
    plan = db.query(models.StudyPlan).filter(models.StudyPlan.id == id).first()
    if plan:
        plan.is_deleted = True
        plan.deleted_at = datetime.utcnow()
        db.commit()
    return {"message": "تم حذف الخطة الدراسية"}

@app.delete("/api/study-plans/bulk/all")
def bulk_delete_study_plan(faculty_id: int, semester: str, academic_year: str, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if "الصيفي" in semester:
        semester_filter = models.StudyPlan.semester.like("%الصيفي%")
    elif "الأول" in semester:
        semester_filter = models.StudyPlan.semester.like("%الأول%")
    elif "الثاني" in semester:
        semester_filter = models.StudyPlan.semester.like("%الثاني%")
    else:
        semester_filter = models.StudyPlan.semester == semester

    plans = db.query(models.StudyPlan).filter(
        models.StudyPlan.faculty_id == faculty_id,
        semester_filter,
        models.StudyPlan.academic_year == academic_year,
        models.StudyPlan.is_deleted == False
    ).all()
    
    if not plans:
        return {"message": "لا يوجد خطة دراسية لمسحها"}

    for plan in plans:
        plan.is_deleted = True
        plan.deleted_at = datetime.utcnow()
        
    user_role_str = get_user_role_display(current_user)
        
    action_text = f"قام بمسح الخطة الدراسية بالكامل ({semester} - العام الجامعي: {academic_year})"
    create_notification(db, faculty_id, f"{current_user.username} ({user_role_str})", action_text, academic_year=academic_year, semester=semester)
    db.commit()
    
    return {"message": "تم مسح الخطة الدراسية بالكامل"}

@app.get("/api/professors/by-faculty/{faculty_id}")
def get_professors_by_faculty(faculty_id: int, db: Session = Depends(get_db)):
    """جلب الدكاترة المنتمين لكلية معينة"""
    profs = db.query(models.Professor).filter(
        models.Professor.faculties.any(models.Faculty.id == faculty_id)
    ).all()
    return [{"id": p.id, "name_ar": p.name_ar, "job_title": p.job_title, "original_workplace": p.original_workplace} for p in profs]

@app.get("/api/professors/{id}/assignments")
def get_professor_assignments(id: int, db: Session = Depends(get_db)):
    """جلب التكليفات (المقررات) الخاصة بدكتور معين من الخطط الدراسية"""
    items = db.query(models.StudyPlanItem).join(models.StudyPlan).filter(models.StudyPlanItem.professor_id == id, models.StudyPlan.is_deleted == False).all()
    result = []
    for item in items:
        sp = item.study_plan
        faculty = sp.faculty if sp else None
        course = item.course
        total_hours = round(
            (item.hours_actual_theory or 0.0) +
            (item.hours_actual_practical or 0.0) +
            (item.hours_actual_exercise or 0.0) +
            (item.hours_actual_activity or 0.0),
            2
        )
        
        course = item.course or item.base_course or (item.module.course if item.module else None)
        course_names = []
        base_course_name = ""
        if course:
            base_course_name = (course.name_ar or course.name_en or course.code or "").strip()
            if base_course_name and item.module and item.module.department_name:
                m_dept = item.module.department_name.strip()
                if m_dept and m_dept.lower() not in base_course_name.lower():
                    base_course_name = f"{base_course_name} ({m_dept})"
            elif not base_course_name and item.module and item.module.department_name:
                base_course_name = item.module.department_name.strip()
        elif item.module and item.module.department_name:
            base_course_name = item.module.department_name.strip()

        if base_course_name:
            course_names.append(base_course_name)

        program_names = []
        if item.entry_group_id:
            raw_entry = str(item.entry_group_id).strip()
            if raw_entry.startswith("{"):
                try:
                    import json
                    parsed = json.loads(raw_entry)
                    pids = parsed.get("prog_ids") or []
                    for pid in pids:
                        prg = db.query(models.Program).filter(models.Program.id == int(pid)).first()
                        if prg and prg.name and prg.name not in program_names:
                            program_names.append(prg.name)
                    sc_list = parsed.get("shared_courses") or []
                    for sc in sc_list:
                        sc_name = (sc.get("name_ar") or sc.get("name_en") or "").strip()
                        if sc_name and sc_name not in course_names:
                            course_names.append(sc_name)
                except:
                    pass
            elif "," in raw_entry:
                try:
                    pids = [int(x.strip()) for x in raw_entry.split(",") if x.strip()]
                    for pid in pids:
                        prg = db.query(models.Program).filter(models.Program.id == pid).first()
                        if prg and prg.name and prg.name not in program_names:
                            program_names.append(prg.name)
                except:
                    pass

        if not program_names:
            if item.program and item.program.name:
                program_names.append(item.program.name)
            elif course and course.program and course.program.name:
                program_names.append(course.program.name)

        final_prog_name = " - ".join(program_names) if program_names else ""
        final_course_name = " - ".join(course_names) if course_names else (course.name_ar or (course.name_en if course else ""))
        
        result.append({
            "id": item.id,
            "academic_year": sp.academic_year if sp else "",
            "semester": sp.semester if sp else "",
            "faculty_id": faculty.id if faculty else None,
            "faculty_name": faculty.name if faculty else "",
            "program_name": final_prog_name,
            "course_name": final_course_name,
            "course_code": course.code if course else "",
            "level": item.level if item.level is not None else (course.level if course else ""),
            "course_semester": sp.semester if sp else "",
            "hours": total_hours
        })
    return result

# ==========================================
# مسارات توقيعات المسؤولين (Signatures APIs)
# ==========================================

@app.get("/api/signatures", response_model=list[schemas.SignatureOut])
def get_signatures(faculty_id: Optional[int] = None, report_type: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(models.Signature).filter(models.Signature.is_deleted == False)
    if faculty_id:
        query = query.filter(models.Signature.faculty_id == faculty_id)
    if report_type:
        query = query.filter(models.Signature.report_type.contains(report_type))
    return query.order_by(models.Signature.order_index.asc()).all()

@app.post("/api/signatures", response_model=schemas.SignatureOut)
def create_signature(signature: schemas.SignatureCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    user_role_str = get_user_role_display(current_user)
    max_order = db.query(func.max(models.Signature.order_index)).filter(models.Signature.faculty_id == signature.faculty_id).scalar() or 0
    new_sig = models.Signature(
        faculty_id=signature.faculty_id,
        signature_title=signature.signature_title,
        official_name=signature.official_name,
        order_index=signature.order_index if signature.order_index is not None else max_order + 1,
        report_type=signature.report_type
    )
    db.add(new_sig)
    db.commit()
    db.refresh(new_sig)
    
    create_notification(db, signature.faculty_id, f"{current_user.username} ({user_role_str})", f"قام بإضافة توقيع: {signature.signature_title} - {signature.official_name}")
    db.commit()
    
    return new_sig

@app.put("/api/signatures/{signature_id}", response_model=schemas.SignatureOut)
def update_signature(signature_id: int, signature: schemas.SignatureUpdate, db: Session = Depends(get_db)):
    db_sig = db.query(models.Signature).filter(models.Signature.id == signature_id, models.Signature.is_deleted == False).first()
    if not db_sig:
        raise HTTPException(status_code=404, detail="التوقيع غير موجود")
    
    if signature.signature_title is not None:
        db_sig.signature_title = signature.signature_title
    if signature.official_name is not None:
        db_sig.official_name = signature.official_name
    if signature.faculty_id is not None:
        db_sig.faculty_id = signature.faculty_id
    if signature.order_index is not None:
        db_sig.order_index = signature.order_index
    if signature.report_type is not None:
        db_sig.report_type = signature.report_type
        
    db.commit()
    db.refresh(db_sig)
    return db_sig

@app.post("/api/signatures/reorder")
def reorder_signatures(items: List[schemas.SignatureReorderItem], db: Session = Depends(get_db)):
    for item in items:
        db.query(models.Signature).filter(models.Signature.id == item.id).update(
            {models.Signature.order_index: item.order_index}
        )
    db.commit()
    return {"message": "تم إعادة الترتيب بنجاح"}

@app.delete("/api/signatures/{signature_id}")
def delete_signature(signature_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    user_role_str = get_user_role_display(current_user)
    db_sig = db.query(models.Signature).filter(models.Signature.id == signature_id, models.Signature.is_deleted == False).first()
    if not db_sig:
        raise HTTPException(status_code=404, detail="التوقيع غير موجود")
        
    fid = db_sig.faculty_id
    sig_title = db_sig.signature_title
    off_name = db_sig.official_name
    
    db_sig.is_deleted = True
    db_sig.deleted_at = datetime.utcnow()
    
    create_notification(db, fid, f"{current_user.username} ({user_role_str})", f"قام بحذف توقيع: {sig_title} - {off_name}")
    db.commit()

    return {"message": "تم الحذف بنجاح"}

# ==========================================
# مسار الإحصائيات (Statistics API)
# ==========================================

@app.get("/api/statistics/{faculty_id}")
def get_statistics(faculty_id: int, academic_year: Optional[str] = None, db: Session = Depends(get_db)):
    # 1. Programs stats
    programs_stats = []
    
    fac = db.query(models.Faculty).filter(models.Faculty.id == faculty_id).first()
    is_medicine_fac = (faculty_id == 10) or (fac and "الطب والجراحة" in fac.name)

    if is_medicine_fac:
        course_query = db.query(
            models.Course.id,
            models.Course.name_ar,
            models.Course.name_en,
            models.Course.is_bundle,
            models.Course.duration,
            models.CourseModule.department_name
        ).join(
            models.StudyPlanItem, models.Course.id == models.StudyPlanItem.course_id
        ).join(
            models.StudyPlan, models.StudyPlanItem.study_plan_id == models.StudyPlan.id
        ).outerjoin(
            models.CourseModule, models.StudyPlanItem.module_id == models.CourseModule.id
        ).filter(
            models.StudyPlan.faculty_id == faculty_id,
            models.StudyPlan.is_deleted == False
        )
        if academic_year:
            course_query = course_query.filter(models.StudyPlan.academic_year == academic_year)
            
        courses_data = course_query.all()
        bundles_dict = {} 
        long_dict = {}
        
        for row in courses_data:
            c_id = row[0]
            name_ar = (row[1] or "").strip()
            name_en = (row[2] or "").strip()
            c_name = name_ar if name_ar else name_en
            if not c_name:
                c_name = f"مقرر {c_id}"
            is_bundle = row[3]
            duration = row[4]
            dept = (row[5] or "").strip()
            
            is_longitudinal = (duration == "مقرر طولي")
            target_dict = long_dict if is_longitudinal else bundles_dict
            
            if c_id not in target_dict:
                # Also fetch all departments defined in CourseModule for this course
                all_mods = db.query(models.CourseModule.department_name).filter(models.CourseModule.course_id == c_id).all()
                all_depts = {m[0].strip() for m in all_mods if m[0] and m[0].strip() and m[0].strip() != "-"}
                target_dict[c_id] = {"name": c_name, "depts": all_depts}
            if dept and dept != "-":
                target_dict[c_id]["depts"].add(dept)
                
        def format_course_list(c_dict):
            lst = []
            for c_id, data in c_dict.items():
                depts = [d for d in data["depts"] if d and d != "-"]
                if depts:
                    depts_str = "، ".join(sorted(list(depts)))
                    lst.append(f"{data['name']} ({depts_str})")
                else:
                    lst.append(f"{data['name']}")
            return lst
            
        bundles_list = format_course_list(bundles_dict)
        long_list = format_course_list(long_dict)
        
        programs_stats.append({
            "name": "المقررات الطولية",
            "course_count": len(long_list),
            "courses": long_list
        })
        programs_stats.append({
            "name": "الحزم الدراسية",
            "course_count": len(bundles_list),
            "courses": bundles_list
        })
    else:
        programs = db.query(models.Program).filter(models.Program.faculty_id == faculty_id).all()
        for p in programs:
            course_query = db.query(models.Course.name_ar).join(
                models.StudyPlanItem, models.Course.id == models.StudyPlanItem.course_id
            ).join(
                models.StudyPlan, models.StudyPlanItem.study_plan_id == models.StudyPlan.id
            ).filter(
                models.StudyPlanItem.program_id == p.id,
                models.StudyPlan.is_deleted == False
            )
            if academic_year:
                course_query = course_query.filter(models.StudyPlan.academic_year == academic_year)
            
            unique_courses = {row[0] for row in course_query.all() if row[0]}
            
            programs_stats.append({
                "name": p.name,
                "course_count": len(unique_courses),
                "courses": list(unique_courses)
            })
        
    # 2. total_professors
    prof_query = db.query(models.StudyPlanItem.professor_id).join(models.StudyPlan).filter(models.StudyPlan.faculty_id == faculty_id, models.StudyPlan.is_deleted == False)
    if academic_year:
        prof_query = prof_query.filter(models.StudyPlan.academic_year == academic_year)
        
    unique_profs = {row[0] for row in prof_query.all() if row[0]}
    total_professors = len(unique_profs)
    
    # 3. courses and professors by semester
    semesters = ["الأول", "الثاني", "الصيفي"]
    courses_by_semester = []
    assigned_professors_by_semester = []
    
    for sem in semesters:
        # Courses count by semester (using course id for accurate unique count)
        c_query = db.query(models.Course.id).join(
            models.StudyPlanItem, models.Course.id == models.StudyPlanItem.course_id
        ).join(
            models.StudyPlan, models.StudyPlanItem.study_plan_id == models.StudyPlan.id
        ).filter(
            models.StudyPlan.faculty_id == faculty_id,
            models.StudyPlan.semester.like(f"%{sem}%"),
            models.StudyPlan.is_deleted == False
        )
        if academic_year:
            c_query = c_query.filter(models.StudyPlan.academic_year == academic_year)
        u_courses = {r[0] for r in c_query.all() if r[0]}
        courses_by_semester.append({"semester": sem, "count": len(u_courses)})
        
        # Professors and their names grouped by course and department
        items_query = db.query(models.StudyPlanItem).join(models.StudyPlan).filter(
            models.StudyPlan.faculty_id == faculty_id,
            models.StudyPlan.semester.like(f"%{sem}%"),
            models.StudyPlan.is_deleted == False
        )
        if academic_year:
            items_query = items_query.filter(models.StudyPlan.academic_year == academic_year)
            
        sem_items = items_query.all()
        
        course_professors_dict = {}
        sem_u_profs = set()
        
        for item in sem_items:
            if not item.course or not item.professor:
                continue
            sem_u_profs.add(item.professor_id)
            c_name = (item.course.name_ar or item.course.name_en or "").strip()
            if not c_name:
                c_name = f"مقرر {item.course.id}"
            p_name = (item.professor.name_ar or item.professor.name_en or "").strip()
            
            d_name = item.module.department_name.strip() if item.module and item.module.department_name else ""
            if d_name and d_name != "-":
                prof_entry = f"{p_name} ({d_name})"
            else:
                prof_entry = f"{p_name}"
            
            if c_name not in course_professors_dict:
                course_professors_dict[c_name] = []
                
            if prof_entry not in course_professors_dict[c_name]:
                course_professors_dict[c_name].append(prof_entry)
                
        prof_names = []
        if sem_u_profs:
            prof_objs = db.query(models.Professor).filter(models.Professor.id.in_(sem_u_profs)).all()
            prof_names = [(p.name_ar or p.name_en or "").strip() for p in prof_objs]
        
        assigned_professors_by_semester.append({
            "semester": sem,
            "count": len(sem_u_profs),
            "professors": prof_names,
            "course_professors": course_professors_dict
        })
        
    return {
        "programs_stats": programs_stats,
        "total_professors": total_professors,
        "courses_by_semester": courses_by_semester,
        "assigned_professors_by_semester": assigned_professors_by_semester
    }

@app.get("/api/main-table-report")
def get_main_table_report(
    academic_year: Optional[str] = None, 
    faculty_id: Optional[int] = None, 
    program_id: Optional[int] = None, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    # Query faculties
    faculties_query = db.query(models.Faculty)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            faculties_query = faculties_query.filter(models.Faculty.id == faculty_id)
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id is not None:
            if faculty_id in faculty_ids:
                faculties_query = faculties_query.filter(models.Faculty.id == faculty_id)
            else:
                return []
        else:
            faculties_query = faculties_query.filter(models.Faculty.id.in_(faculty_ids))
    elif current_user.faculty_id:
        faculties_query = faculties_query.filter(models.Faculty.id == current_user.faculty_id)
    elif faculty_id:
        faculties_query = faculties_query.filter(models.Faculty.id == faculty_id)
    faculties = faculties_query.all()
    
    report = []
    
    for f in faculties:
        # Query programs for this faculty
        programs_query = db.query(models.Program).filter(models.Program.faculty_id == f.id)
        if program_id:
            programs_query = programs_query.filter(models.Program.id == program_id)
        programs = programs_query.all()
        
        programs_list = []
        
        for p in programs:
            # Count UNIQUE courses assigned in the study plan for this program per semester
            sp_course_query = db.query(models.StudyPlanItem.course_id, models.StudyPlan.semester, models.Course.is_bundle).join(
                models.StudyPlan, models.StudyPlanItem.study_plan_id == models.StudyPlan.id
            ).join(
                models.Course, models.StudyPlanItem.course_id == models.Course.id
            ).filter(
                models.StudyPlanItem.program_id == p.id, models.StudyPlan.is_deleted == False
            )
            if academic_year:
                sp_course_query = sp_course_query.filter(models.StudyPlan.academic_year == academic_year)
            
            rows_courses = sp_course_query.all()
            unique_courses_t1 = {r[0] for r in rows_courses if r[0] and "الأول" in (r[1] or "")}
            unique_courses_t2 = {r[0] for r in rows_courses if r[0] and "الثاني" in (r[1] or "")}
            unique_courses_t3 = {r[0] for r in rows_courses if r[0] and "الصيفي" in (r[1] or "")}
            
            bundles_t1 = {r[0] for r in rows_courses if r[0] and "الأول" in (r[1] or "") and r[2] == True}
            bundles_t2 = {r[0] for r in rows_courses if r[0] and "الثاني" in (r[1] or "") and r[2] == True}
            bundles_t3 = {r[0] for r in rows_courses if r[0] and "الصيفي" in (r[1] or "") and r[2] == True}
            
            long_t1 = {r[0] for r in rows_courses if r[0] and "الأول" in (r[1] or "") and r[2] == False}
            long_t2 = {r[0] for r in rows_courses if r[0] and "الثاني" in (r[1] or "") and r[2] == False}
            long_t3 = {r[0] for r in rows_courses if r[0] and "الصيفي" in (r[1] or "") and r[2] == False}
            
            courses_count_t1 = len(unique_courses_t1)
            courses_count_t2 = len(unique_courses_t2)
            courses_count_t3 = len(unique_courses_t3)
            
            # Count unique professors assigned to this program in study plans per term
            prof_query = db.query(models.StudyPlanItem.professor_id, models.StudyPlan.semester).join(
                models.StudyPlan, models.StudyPlanItem.study_plan_id == models.StudyPlan.id
            ).filter(
                models.StudyPlanItem.program_id == p.id, models.StudyPlan.is_deleted == False
            )
            if academic_year:
                prof_query = prof_query.filter(models.StudyPlan.academic_year == academic_year)
            
            rows_profs = prof_query.all()
            unique_profs_t1 = {r[0] for r in rows_profs if r[0] and "الأول" in (r[1] or "")}
            unique_profs_t2 = {r[0] for r in rows_profs if r[0] and "الثاني" in (r[1] or "")}
            unique_profs_t3 = {r[0] for r in rows_profs if r[0] and "الصيفي" in (r[1] or "")}
            
            professors_count_t1 = len(unique_profs_t1)
            professors_count_t2 = len(unique_profs_t2)
            professors_count_t3 = len(unique_profs_t3)
            
            programs_list.append({
            "program_id": p.id,
                "program_name": p.name,
                "courses_count_t1": courses_count_t1,
                "courses_count_t2": courses_count_t2,
                "courses_count_t3": courses_count_t3,
                "bundles_count_t1": len(bundles_t1),
                "bundles_count_t2": len(bundles_t2),
                "bundles_count_t3": len(bundles_t3),
                "long_count_t1": len(long_t1),
                "long_count_t2": len(long_t2),
                "long_count_t3": len(long_t3),
                "professors_count_t1": professors_count_t1,
                "professors_count_t2": professors_count_t2,
                "professors_count_t3": professors_count_t3
            })
            
        if program_id and not programs_list:
            continue
            
        # Calculate Faculty totals for courses and professors
        faculty_sp_course_query = db.query(models.StudyPlanItem.course_id, models.StudyPlan.semester).join(models.StudyPlan).filter(
            models.StudyPlan.faculty_id == f.id, models.StudyPlan.is_deleted == False
        )
        if academic_year:
            faculty_sp_course_query = faculty_sp_course_query.filter(models.StudyPlan.academic_year == academic_year)
        
        fac_rows_courses = faculty_sp_course_query.all()
        fac_unique_courses_t1 = {r[0] for r in fac_rows_courses if r[0] and "الأول" in (r[1] or "")}
        fac_unique_courses_t2 = {r[0] for r in fac_rows_courses if r[0] and "الثاني" in (r[1] or "")}
        fac_unique_courses_t3 = {r[0] for r in fac_rows_courses if r[0] and "الصيفي" in (r[1] or "")}
        
        total_courses_t1 = len(fac_unique_courses_t1)
        total_courses_t2 = len(fac_unique_courses_t2)
        total_courses_t3 = len(fac_unique_courses_t3)
        
        faculty_prof_query = db.query(models.StudyPlanItem.professor_id, models.StudyPlan.semester).join(models.StudyPlan).filter(
            models.StudyPlan.faculty_id == f.id, models.StudyPlan.is_deleted == False
        )
        if academic_year:
            faculty_prof_query = faculty_prof_query.filter(models.StudyPlan.academic_year == academic_year)
        
        fac_rows_profs = faculty_prof_query.all()
        fac_unique_profs_t1 = {r[0] for r in fac_rows_profs if r[0] and "الأول" in (r[1] or "")}
        fac_unique_profs_t2 = {r[0] for r in fac_rows_profs if r[0] and "الثاني" in (r[1] or "")}
        fac_unique_profs_t3 = {r[0] for r in fac_rows_profs if r[0] and "الصيفي" in (r[1] or "")}
        
        total_professors_t1 = len(fac_unique_profs_t1)
        total_professors_t2 = len(fac_unique_profs_t2)
        total_professors_t3 = len(fac_unique_profs_t3)
        
        report.append({
            "faculty_id": f.id,
            "faculty_name": f.name,
            "programs": programs_list,
            "total_courses_t1": total_courses_t1,
            "total_courses_t2": total_courses_t2,
            "total_courses_t3": total_courses_t3,
            "total_professors_t1": total_professors_t1,
            "total_professors_t2": total_professors_t2,
            "total_professors_t3": total_professors_t3
        })
        
    return report

# ==========================================
# مسارات استرجاع المحذوف (Recycle Bin APIs)
# ==========================================
@app.get("/api/recycle-bin")
def get_recycle_bin(faculty_id: Optional[int] = None, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    items = []
    
    # 1. أعضاء هيئة التدريس
    prof_query = db.query(models.Professor).filter(models.Professor.is_deleted == True)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            prof_query = prof_query.filter(models.Professor.faculties.any(models.Faculty.id == faculty_id))
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id:
            if faculty_id in faculty_ids:
                prof_query = prof_query.filter(models.Professor.faculties.any(models.Faculty.id == faculty_id))
            else:
                prof_query = prof_query.filter(models.Professor.id == -1)
        else:
            prof_query = prof_query.filter(
                (models.Professor.faculties.any(models.Faculty.id.in_(faculty_ids))) | 
                (~models.Professor.faculties.any())
            )
    elif current_user.faculty_id:
        prof_query = prof_query.filter(
            (models.Professor.faculties.any(models.Faculty.id == current_user.faculty_id)) | 
            (~models.Professor.faculties.any())
        )
    elif faculty_id:
        prof_query = prof_query.filter(models.Professor.faculties.any(models.Faculty.id == faculty_id))
        
    for p in prof_query.all():
        items.append({
            "type": "professor",
            "id": p.id,
            "name": p.name_ar,
            "deleted_at": p.deleted_at.isoformat() if p.deleted_at else None,
            "faculty": "، ".join([f.name for f in p.faculties]) if p.faculties else "غير محدد",
            "academic_year": p.academic_year
        })

    # 2. المقررات الدراسية
    course_query = db.query(models.Course).filter(models.Course.is_deleted == True)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            course_query = course_query.filter(models.Course.faculty_id == faculty_id)
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id:
            if faculty_id in faculty_ids:
                course_query = course_query.filter(models.Course.faculty_id == faculty_id)
            else:
                course_query = course_query.filter(models.Course.id == -1)
        else:
            course_query = course_query.filter(models.Course.faculty_id.in_(faculty_ids))
    elif current_user.faculty_id:
        course_query = course_query.filter(models.Course.faculty_id == current_user.faculty_id)
    elif faculty_id:
        course_query = course_query.filter(models.Course.faculty_id == faculty_id)
        
    for c in course_query.all():
        course_name = (c.name_ar or c.name_en or "").strip()
        items.append({
            "type": "course",
            "id": c.id,
            "name": course_name if course_name else (c.code or "مقرر بدون اسم"),
            "name_ar": c.name_ar or "",
            "name_en": c.name_en or "",
            "code": c.code or "",
            "deleted_at": c.deleted_at.isoformat() if c.deleted_at else None,
            "faculty": c.faculty.name if c.faculty else "غير محدد",
            "academic_year": c.year
        })

    # 3. الخطط الدراسية
    plan_query = db.query(models.StudyPlan).filter(models.StudyPlan.is_deleted == True)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            plan_query = plan_query.filter(models.StudyPlan.faculty_id == faculty_id)
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id:
            if faculty_id in faculty_ids:
                plan_query = plan_query.filter(models.StudyPlan.faculty_id == faculty_id)
            else:
                plan_query = plan_query.filter(models.StudyPlan.id == -1)
        else:
            plan_query = plan_query.filter(models.StudyPlan.faculty_id.in_(faculty_ids))
    elif current_user.faculty_id:
        plan_query = plan_query.filter(models.StudyPlan.faculty_id == current_user.faculty_id)
    elif faculty_id:
        plan_query = plan_query.filter(models.StudyPlan.faculty_id == faculty_id)
        
    for pl in plan_query.all():
        programs = set()
        for item in pl.items:
            if item.program:
                programs.add(item.program.name)
        prog_str = (" لبرنامج " + " و ".join(programs)) if programs else ""
        items.append({
            "type": "study_plan",
            "id": pl.id,
            "name": f"خطة {pl.semester} - {pl.academic_year}{prog_str}",
            "deleted_at": pl.deleted_at.isoformat() if pl.deleted_at else None,
            "faculty": pl.faculty.name if pl.faculty else "غير محدد",
            "academic_year": pl.academic_year
        })
        
    # 4. التوقيعات
    sig_query = db.query(models.Signature).filter(models.Signature.is_deleted == True)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            sig_query = sig_query.filter(models.Signature.faculty_id == faculty_id)
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id:
            if faculty_id in faculty_ids:
                sig_query = sig_query.filter(models.Signature.faculty_id == faculty_id)
            else:
                sig_query = sig_query.filter(models.Signature.id == -1)
        else:
            sig_query = sig_query.filter(models.Signature.faculty_id.in_(faculty_ids))
    elif current_user.faculty_id:
        sig_query = sig_query.filter(models.Signature.faculty_id == current_user.faculty_id)
    elif faculty_id:
        sig_query = sig_query.filter(models.Signature.faculty_id == faculty_id)
        
    for s in sig_query.all():
        items.append({
            "type": "signature",
            "id": s.id,
            "name": f"توقيع: {s.signature_title} ({s.official_name})",
            "deleted_at": s.deleted_at.isoformat() if s.deleted_at else None,
            "faculty": s.faculty.name if s.faculty else "غير محدد",
            "academic_year": s.academic_year
        })
        
    return items

@app.get("/api/professors-report")
def get_professors_report(
    academic_year: Optional[str] = None,
    faculty_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    query = db.query(models.Professor)
    
    base_conditions = []
    if academic_year:
        base_conditions.append(models.Professor.academic_year == academic_year)
    
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            base_conditions.append(models.Professor.faculties.any(models.Faculty.id == faculty_id))
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id is not None:
            if faculty_id in faculty_ids:
                base_conditions.append(models.Professor.faculties.any(models.Faculty.id == faculty_id))
            else:
                return []
        else:
            base_conditions.append(models.Professor.faculties.any(models.Faculty.id.in_(faculty_ids)))
    elif current_user.faculty_id:
        base_conditions.append(models.Professor.faculties.any(models.Faculty.id == current_user.faculty_id))
    elif faculty_id:
        base_conditions.append(models.Professor.faculties.any(models.Faculty.id == faculty_id))
        
    sp_subquery = db.query(models.StudyPlanItem.professor_id).join(models.StudyPlan).filter(models.StudyPlan.is_deleted == False)
    if academic_year:
        sp_subquery = sp_subquery.filter(models.StudyPlan.academic_year == academic_year)
    if current_user.role in [models.UserRole.admin, models.UserRole.student_affairs] or current_user.all_faculties_access:
        if faculty_id:
            sp_subquery = sp_subquery.filter(models.StudyPlan.faculty_id == faculty_id)
    elif current_user.assigned_faculties and len(current_user.assigned_faculties) > 0:
        faculty_ids = [f.id for f in current_user.assigned_faculties]
        if faculty_id is not None:
            sp_subquery = sp_subquery.filter(models.StudyPlan.faculty_id == faculty_id)
        else:
            sp_subquery = sp_subquery.filter(models.StudyPlan.faculty_id.in_(faculty_ids))
    elif current_user.faculty_id:
        sp_subquery = sp_subquery.filter(models.StudyPlan.faculty_id == current_user.faculty_id)
    elif faculty_id:
        sp_subquery = sp_subquery.filter(models.StudyPlan.faculty_id == faculty_id)
        
    from sqlalchemy import or_, and_
    
    if base_conditions:
        query = query.filter(
            or_(
                and_(*base_conditions),
                models.Professor.id.in_(sp_subquery)
            )
        )
    else:
        query = query.filter(models.Professor.id.in_(sp_subquery))
        
    professors = query.all()
    result = []
    
    for prof in professors:
        # Get course assignments in study plans
        items_query = db.query(models.StudyPlanItem).join(models.StudyPlan).filter(
            models.StudyPlanItem.professor_id == prof.id, models.StudyPlan.is_deleted == False
        )
        if academic_year:
            items_query = items_query.filter(models.StudyPlan.academic_year == academic_year)
        if current_user.role in [models.UserRole.faculty_admin, models.UserRole.faculty_professor]:
            items_query = items_query.filter(models.StudyPlan.faculty_id == current_user.faculty_id)
        elif current_user.role == models.UserRole.reviewer:
            faculty_ids = [f.id for f in current_user.assigned_faculties]
            if faculty_id is not None:
                items_query = items_query.filter(models.StudyPlan.faculty_id == faculty_id)
            else:
                items_query = items_query.filter(models.StudyPlan.faculty_id.in_(faculty_ids))
        elif faculty_id:
            items_query = items_query.filter(models.StudyPlan.faculty_id == faculty_id)
            
        items = items_query.all()
        
        grouped_t1 = {}
        grouped_t2 = {}
        grouped_t3 = {}
        
        hours_t1 = 0.0
        hours_t2 = 0.0
        hours_t3 = 0.0
        
        for item in items:
            if not item.course:
                continue
            semester_name = item.study_plan.semester or ""
            act_hours = round((item.hours_actual_theory or 0.0) + (item.hours_actual_practical or 0.0), 2)
            
            if "الأول" in semester_name:
                grouped_t1.setdefault(item.course.id, []).append((item, act_hours))
                hours_t1 += act_hours
            elif "الثاني" in semester_name:
                grouped_t2.setdefault(item.course.id, []).append((item, act_hours))
                hours_t2 += act_hours
            elif "الصيفي" in semester_name:
                grouped_t3.setdefault(item.course.id, []).append((item, act_hours))
                hours_t3 += act_hours
                
        def format_course_group(course_items):
            if not course_items:
                return None
            c = course_items[0][0].course
            has_module = any(itm[0].module_id for itm in course_items)
            total_hrs = sum(itm[1] for itm in course_items)
            
            c_name = (c.name_ar or "").strip() or (c.name_en or "").strip() or (c.code or "").strip() or "مقرر"
            
            valid_dept_items = [itm for itm in course_items if itm[0].module and itm[0].module.department_name and itm[0].module.department_name.strip() not in ['-', '']]
            
            if not has_module or not valid_dept_items:
                return {
                    "course_name": c_name,
                    "course_code": c.code or "",
                    "hours": total_hrs,
                    "is_module": False,
                    "dept_details": [],
                    "display_text": f"{c_name} ({c.code}) ({total_hrs:g} س)" if c.code else f"{c_name} ({total_hrs:g} س)"
                }
            
            dept_details = []
            for itm, hrs in course_items:
                d_name = itm.module.department_name.strip() if (itm.module and itm.module.department_name) else ""
                if d_name and d_name not in ['-', '']:
                    dept_details.append({
                        "dept_name": d_name,
                        "hours": hrs,
                        "display": f"{d_name}: {hrs:g} س"
                    })
                    
            if not dept_details:
                return {
                    "course_name": c_name,
                    "course_code": c.code or "",
                    "hours": total_hrs,
                    "is_module": False,
                    "dept_details": [],
                    "display_text": f"{c_name} ({c.code}) ({total_hrs:g} س)" if c.code else f"{c_name} ({total_hrs:g} س)"
                }
                    
            return {
                "course_name": c_name,
                "course_code": c.code or "",
                "hours": total_hrs,
                "is_module": True,
                "dept_details": dept_details,
                "display_text": f"{c_name} ({c.code})" if c.code else c_name
            }
            
        courses_t1 = [format_course_group(group) for group in grouped_t1.values() if format_course_group(group)]
        courses_t2 = [format_course_group(group) for group in grouped_t2.values() if format_course_group(group)]
        courses_t3 = [format_course_group(group) for group in grouped_t3.values() if format_course_group(group)]
                
        if courses_t1 or courses_t2 or courses_t3:
            result.append({
                "professor_id": prof.id,
                "professor_name": prof.name_ar,
                "job_title": prof.job_title,
                "contract_type": prof.contract_type,
                "work_days": prof.work_days,
                "original_workplace": prof.original_workplace,
                "faculties": [fac.name for fac in prof.faculties],
                "courses_t1": courses_t1,
                "courses_t2": courses_t2,
                "courses_t3": courses_t3,
                "hours_t1": round(hours_t1, 2),
                "hours_t2": round(hours_t2, 2),
                "hours_t3": round(hours_t3, 2),
                "academic_year": prof.academic_year,
                "semester1_weeks": prof.semester1_weeks,
                "semester2_weeks": prof.semester2_weeks,
                "summer_weeks": prof.summer_weeks,
                "academic_year_weeks": prof.academic_year_weeks
            })
        
    return result
# ==========================================
# مسارات الإشعارات (Notifications)
# ==========================================

@app.post("/api/notifications/log")
def log_notification(data: schemas.NotificationCreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    user_action_by_str = get_user_action_by(current_user, db)
    action_text = data.action_text.replace('[ADMIN_ONLY] ', '').replace('[ADMIN_ONLY]', '').strip()
    if getattr(data, 'admin_only', False):
        action_text = f"[ADMIN_ONLY] {action_text}"
    
    if data.faculty_ids:
        for fid in data.faculty_ids:
            notif = models.Notification(
                faculty_id=fid,
                action_by=user_action_by_str,
                action_text=action_text,
                created_at=datetime.utcnow(),
                academic_year=data.academic_year,
                semester=data.semester
            )
            db.add(notif)
    else:
        notif = models.Notification(
            faculty_id=None,
            action_by=user_action_by_str,
            action_text=action_text,
            created_at=datetime.utcnow(),
            academic_year=data.academic_year,
            semester=data.semester
        )
        db.add(notif)
        
    db.commit()

    # توثيق الحدث أيضاً في سجل العمليات (Audit Logs)
    try:
        log_activity(
            db=db,
            username=current_user.username,
            action_type="EXPORT" if ("تصدير" in data.action_text or "تنزيل" in data.action_text) else "PRINT",
            description=data.action_text,
            entity_type="MAIN_TABLE" if "الجدول الرئيسي" in data.action_text else ("PROFESSORS" if "تدريس" in data.action_text else "STUDY_PLAN"),
            user_id=current_user.id,
            user_role=get_user_role_display(current_user, db),
            faculty_id=data.faculty_ids[0] if (data.faculty_ids and len(data.faculty_ids) == 1) else None,
            academic_year=data.academic_year,
            semester=data.semester,
            status="success"
        )
    except Exception as e:
        logger.error(f"Error logging activity for notification: {e}")

    return {"message": "تم تسجيل الإشعار بنجاح"}


from datetime import timedelta

@app.get("/api/notifications", response_model=list[schemas.NotificationOut])
def get_notifications(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    # مسح الإشعارات القديمة (أكثر من 20 يوم)
    cutoff_date = datetime.utcnow() - timedelta(days=20)
    db.query(models.Notification).filter(models.Notification.created_at < cutoff_date).delete()
    db.commit()

    query = db.query(models.Notification)
    if current_user.role == models.UserRole.student_affairs:
        query = query.filter(~models.Notification.action_text.startswith('[ADMIN_ONLY]'))
    elif current_user.role in [models.UserRole.faculty_admin, models.UserRole.faculty_professor]:
        if current_user.faculty_id:
            query = query.filter(models.Notification.faculty_id == current_user.faculty_id)
            query = query.filter(~models.Notification.action_text.startswith('[ADMIN_ONLY]'))
        else:
            return []
    elif current_user.role == models.UserRole.manager:
        if not current_user.all_faculties_access and current_user.assigned_faculties:
            fac_ids = [f.id for f in current_user.assigned_faculties]
            query = query.filter(
                (models.Notification.faculty_id.in_(fac_ids)) | 
                (models.Notification.faculty_id == None)
            )

    notifications = query.order_by(models.Notification.id.desc()).all()
    
    # خريطة المسميات الوظيفية والكليات الحالية للمستخدمين لتحديث حتى الإشعارات السابقة بدقة
    user_job_map = {u.username: get_user_role_display(u, db) for u in db.query(models.User).all()}

    for n in notifications:
        if n.action_text.startswith('[ADMIN_ONLY] '):
            n.action_text = n.action_text.replace('[ADMIN_ONLY] ', '')
        if n.action_by:
            m = re.match(r"^([^\(]+)(?:\s*\((.+)\))?$", n.action_by.strip())
            if m:
                raw_uname = m.group(1).strip()
                clean_uname = raw_uname.replace('@gmail.com', '').strip()
                matched_role = user_job_map.get(clean_uname) or user_job_map.get(raw_uname)
                if matched_role:
                    n.action_by = f"{clean_uname} ({matched_role})"
            
    return notifications

@app.delete("/api/notifications/{id}")
def delete_notification(id: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role in [models.UserRole.faculty_admin, models.UserRole.student_affairs] and not current_user.perm_delete_notif_btn:
        raise HTTPException(status_code=403, detail="غير مصرح لك بمسح الإشعارات")

    notif = db.query(models.Notification).filter(models.Notification.id == id).first()
    if notif:
        db.delete(notif)
        db.commit()
    return {"status": "deleted"}

@app.delete("/api/notifications/bulk/all")
def bulk_delete_notifications(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    if current_user.role in [models.UserRole.faculty_admin, models.UserRole.student_affairs] and not current_user.perm_delete_notif_btn:
        raise HTTPException(status_code=403, detail="غير مصرح لك بمسح الإشعارات")
    
    query = db.query(models.Notification)
    if current_user.role == models.UserRole.faculty_professor:
        query = query.filter(models.Notification.faculty_id == current_user.faculty_id)
    
    query.delete()
    db.commit()
    return {"status": "deleted"}



@app.post("/api/recycle-bin/restore")
def restore_recycle_bin(action: schemas.RecycleBinAction, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    user_role_str = get_user_role_display(current_user)
    
    restored_by_faculty = {}
    
    for item in action.items:
        obj = None
        fids = []
        name = ""
        
        if item.type == "professor":
            obj = db.query(models.Professor).filter(models.Professor.id == item.id).first()
            if obj:
                name = f"عضو هيئة تدريس: {obj.name_ar}"
                fids = [f.id for f in obj.faculties] if obj.faculties else [None]
        elif item.type == "course":
            obj = db.query(models.Course).filter(models.Course.id == item.id).first()
            if obj:
                c_title = obj.name_ar or obj.name_en or obj.code or ""
                name = f"مقرر: {c_title}"
                fids = [obj.faculty_id]
        elif item.type == "study_plan":
            obj = db.query(models.StudyPlan).filter(models.StudyPlan.id == item.id).first()
            if obj:
                name = f"الخطة الدراسية {obj.semester} - {obj.academic_year}"
                fids = [obj.faculty_id]
        elif item.type == "signature":
            obj = db.query(models.Signature).filter(models.Signature.id == item.id).first()
            if obj:
                name = f"توقيع: {obj.signature_title}"
                fids = [obj.faculty_id]
        else:
            continue
            
        if obj:
            obj.is_deleted = False
            obj.deleted_at = None
            
            for fid in fids:
                if fid not in restored_by_faculty:
                    restored_by_faculty[fid] = []
                if name not in restored_by_faculty[fid]:
                    restored_by_faculty[fid].append(name)
            
    for fid, names in restored_by_faculty.items():
        if len(names) == 1:
            action_text = f"قام باسترجاع {names[0]}"
        else:
            action_text = f"قام باسترجاع {len(names)} عناصر من سلة المحذوفات"
            
        create_notification(db, fid, f"{current_user.username} ({user_role_str})", action_text)
        
    db.commit()
    return {"message": "تم الاسترجاع بنجاح"}

@app.delete("/api/recycle-bin/hard-delete")
def hard_delete_recycle_bin(action: schemas.RecycleBinAction, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    user_role_str = get_user_role_display(current_user)
    
    deleted_by_faculty = {}
    
    for item in action.items:
        obj = None
        fids = []
        name = ""
        
        if item.type == "professor":
            obj = db.query(models.Professor).filter(models.Professor.id == item.id).first()
            if obj:
                name = f"عضو هيئة تدريس: {obj.name_ar}"
                fids = [f.id for f in obj.faculties] if obj.faculties else [None]
        elif item.type == "course":
            obj = db.query(models.Course).filter(models.Course.id == item.id).first()
            if obj:
                c_title = obj.name_ar or obj.name_en or obj.code or ""
                name = f"مقرر: {c_title}"
                fids = [obj.faculty_id]
        elif item.type == "study_plan":
            obj = db.query(models.StudyPlan).filter(models.StudyPlan.id == item.id).first()
            if obj:
                name = f"الخطة الدراسية {obj.semester} - {obj.academic_year}"
                fids = [obj.faculty_id]
        elif item.type == "signature":
            obj = db.query(models.Signature).filter(models.Signature.id == item.id).first()
            if obj:
                name = f"توقيع: {obj.signature_title}"
                fids = [obj.faculty_id]
        else:
            continue
            
        if obj:
            for fid in fids:
                if fid not in deleted_by_faculty:
                    deleted_by_faculty[fid] = []
                if name not in deleted_by_faculty[fid]:
                    deleted_by_faculty[fid].append(name)
            
            db.delete(obj)
            
    for fid, names in deleted_by_faculty.items():
        if len(names) == 1:
            action_text = f"قام بالحذف النهائي لـ {names[0]}"
        else:
            action_text = f"قام بالحذف النهائي لـ {len(names)} عناصر من سلة المحذوفات"
            
        create_notification(db, fid, f"{current_user.username} ({user_role_str})", action_text)
        
    db.commit()
    return {"message": "تم الحذف النهائي بنجاح"}

# ==========================================
# 17. API الأعوام الجامعية (Academic Years)
# ==========================================

@app.get("/api/academic-years", response_model=List[schemas.AcademicYearOut])
def get_academic_years(db: Session = Depends(get_db)):
    return db.query(models.AcademicYear).order_by(models.AcademicYear.order_index.asc(), models.AcademicYear.id.asc()).all()

@app.post("/api/academic-years/reorder")
def reorder_academic_years(
    reorder_data: List[schemas.AcademicYearReorderItem],
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="غير مصرح لك بتعديل ترتيب الأعوام الجامعية")
    
    for item in reorder_data:
        db.query(models.AcademicYear).filter(models.AcademicYear.id == item.id).update({"order_index": item.order_index})
    
    db.commit()
    return {"message": "تم تحديث الترتيب بنجاح"}

@app.post("/api/academic-years", response_model=schemas.AcademicYearOut)
def create_academic_year(
    year: schemas.AcademicYearCreate, 
    db: Session = Depends(get_db), 
    current_user: models.User = Depends(auth.get_current_user)
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="غير مصرح لك بإضافة عام جامعي")
        
    db_year = db.query(models.AcademicYear).filter(models.AcademicYear.name == year.name).first()
    if db_year:
        raise HTTPException(status_code=400, detail="العام الجامعي موجود بالفعل")
        
    new_year = models.AcademicYear(
        name=year.name,
        semester1_weeks=year.semester1_weeks if year.semester1_weeks is not None else 15,
        semester2_weeks=year.semester2_weeks if year.semester2_weeks is not None else 14,
        summer_weeks=year.summer_weeks if year.summer_weeks is not None else 7,
        med_semester1_weeks=year.med_semester1_weeks if year.med_semester1_weeks is not None else (year.semester1_weeks if year.semester1_weeks is not None else 15),
        med_semester2_weeks=year.med_semester2_weeks if year.med_semester2_weeks is not None else (year.semester2_weeks if year.semester2_weeks is not None else 14),
        med_summer_weeks=year.med_summer_weeks if year.med_summer_weeks is not None else (year.summer_weeks if year.summer_weeks is not None else 7),
    )
    db.add(new_year)
    db.commit()
    db.refresh(new_year)
    return new_year

@app.put("/api/academic-years/{year_id}", response_model=schemas.AcademicYearOut)
def update_academic_year(
    year_id: int,
    year_update: schemas.AcademicYearUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user)
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="غير مصرح لك بتعديل عام جامعي")
        
    db_year = db.query(models.AcademicYear).filter(models.AcademicYear.id == year_id).first()
    if not db_year:
        raise HTTPException(status_code=404, detail="العام الجامعي غير موجود")
        
    if year_update.name is not None and year_update.name != db_year.name:
        existing_year = db.query(models.AcademicYear).filter(models.AcademicYear.name == year_update.name, models.AcademicYear.id != year_id).first()
        if existing_year:
            raise HTTPException(status_code=400, detail="يوجد عام جامعي آخر بنفس الاسم")
        db_year.name = year_update.name
        
    if year_update.semester1_weeks is not None:
        db_year.semester1_weeks = year_update.semester1_weeks
    if year_update.semester2_weeks is not None:
        db_year.semester2_weeks = year_update.semester2_weeks
    if year_update.summer_weeks is not None:
        db_year.summer_weeks = year_update.summer_weeks

    if year_update.med_semester1_weeks is not None:
        db_year.med_semester1_weeks = year_update.med_semester1_weeks
    if year_update.med_semester2_weeks is not None:
        db_year.med_semester2_weeks = year_update.med_semester2_weeks
    if year_update.med_summer_weeks is not None:
        db_year.med_summer_weeks = year_update.med_summer_weeks

    db.commit()
    db.refresh(db_year)
    return db_year

@app.delete("/api/academic-years/{year_id}")
def delete_academic_year(
    year_id: int, 
    db: Session = Depends(get_db), 
    current_user: models.User = Depends(auth.get_current_user)
):
    if current_user.role != models.UserRole.admin:
        raise HTTPException(status_code=403, detail="غير مصرح لك بحذف عام جامعي")
        
    db_year = db.query(models.AcademicYear).filter(models.AcademicYear.id == year_id).first()
    if not db_year:
        raise HTTPException(status_code=404, detail="العام الجامعي غير موجود")
        
    db.delete(db_year)
    db.commit()
    return {"message": "تم الحذف بنجاح"}
